from pathlib import Path
import io
import pandas as pd


def supported(path):
    return Path(path).suffix.lower() in {".csv", ".txt", ".pdf", ".docx", ".png", ".jpg", ".jpeg", ".webp", ".xlsx"}


def _ocr_image_bytes(data, filename, language="auto"):
    from ai_engine.nlp.document_intelligence import ocr_image
    return ocr_image(data, filename, language)


def read_text(path_or_bytes, filename="upload.txt", language="auto", return_meta=False):
    suffix = Path(filename).suffix.lower()
    data = path_or_bytes if isinstance(path_or_bytes, bytes) else Path(path_or_bytes).read_bytes()
    meta = {"method": "text", "language": "English", "ocr_confidence": None, "warnings": []}
    if suffix in {".txt", ".csv"}:
        text = data.decode("utf-8", errors="ignore")
        from ai_engine.nlp.document_intelligence import detect_language
        li = detect_language(text, language)
        meta.update({"method":"text-decode", "language":li["primary"], "language_candidates":li["candidates"]})
    elif suffix == ".xlsx":
        df = pd.read_excel(io.BytesIO(data))
        text = df.to_csv(index=False)
        meta.update({"method":"excel-table"})
    elif suffix == ".pdf":
        text = ""
        try:
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        except Exception as exc:
            meta["warnings"].append(f"PDF text extraction failed: {exc}")
        # Scanned PDFs often have no text layer. Render pages and OCR them locally.
        if len(text.strip()) < 80:
            try:
                import fitz
                doc = fitz.open(stream=data, filetype="pdf")
                chunks=[]; confs=[]
                for page in doc:
                    pix = page.get_pixmap(matrix=fitz.Matrix(1.7,1.7), alpha=False)
                    result = _ocr_image_bytes(pix.tobytes("png"), f"{filename}-page.png", language)
                    chunks.append(result["text"]); confs.append(result.get("confidence",0))
                    meta["warnings"].extend(result.get("warnings",[]))
                text = "\n\n".join(chunks)
                meta.update({"method":"pdf-render+ocr", "ocr_confidence": round(sum(confs)/len(confs),3) if confs else 0})
            except Exception as exc:
                meta["warnings"].append(f"Scanned PDF OCR unavailable: {exc}")
        else:
            from ai_engine.nlp.document_intelligence import detect_language
            li=detect_language(text,language); meta.update({"method":"pdf-text-layer","language":li["primary"],"language_candidates":li["candidates"]})
    elif suffix == ".docx":
        try:
            from docx import Document
            doc = Document(io.BytesIO(data))
            text = "\n".join(p.text for p in doc.paragraphs)
            # Include table cells, common in FIR templates.
            for table in doc.tables:
                for row in table.rows:
                    text += "\n" + " | ".join(cell.text for cell in row.cells)
        except Exception as exc:
            text = f"[DOCX extraction unavailable: {exc}]"
            meta["warnings"].append(text)
    elif suffix in {".png", ".jpg", ".jpeg", ".webp"}:
        result = _ocr_image_bytes(data, filename, language)
        text = result["text"]
        meta.update({"method":result.get("method"), "language":result.get("language"), "language_code":result.get("language_code"), "language_candidates":result.get("language_candidates",[]), "ocr_confidence":result.get("confidence",0), "warnings":result.get("warnings",[])})
    else:
        text = ""
    if return_meta:
        return text, meta
    return text
