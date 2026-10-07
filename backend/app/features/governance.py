"""Governance and oversight: management dashboard, data protection, backup, access history, legal packs, import connectors.

Everything here is derived from real rows, is permission-gated (see each route) and every state-changing action is
written to the hash-chained audit log. Nothing in this module rewrites or deletes audit rows.

Operational limits (documented in the UI as well):
* BACKUP_MAX_MB (env, default 200): maximum combined size of the database snapshot + evidence files in one backup.
* Backups are Fernet-encrypted with a key derived by scrypt from the passphrase and a random salt stored in the file
  header. The passphrase is never stored; without it a backup cannot be opened by anyone (including the server).
* INBOX_WATCH=true enables the folder inbox watcher (off by default); INBOX_DIR overrides the folder.
"""
from __future__ import annotations

import base64
import csv
import hashlib
import hmac
import io
import json
import logging
import os
import re
import secrets
import shutil
import sqlite3
import statistics
import tempfile
import threading
import time
import zipfile
from collections import Counter, defaultdict, deque
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any

import pandas as pd
import sqlalchemy as sa
from fastapi import APIRouter, Depends, File, HTTPException, Query, Request, UploadFile
from fastapi.responses import FileResponse, Response
from pydantic import BaseModel, Field
from sqlalchemy import Boolean, Column, DateTime, Integer, String, Text, inspect as sa_inspect
from sqlalchemy.orm import Session
from starlette.concurrency import run_in_threadpool

from backend.app.db.database import Base, engine
from backend.app.models.models import (Alert, AuditEvent, Case, CaseMember, DocumentRecord, Entity, Evidence, Event,
                                       Note, Relationship, Task, User)
from backend.app.services import app_service as A

log = logging.getLogger("dcn.governance")

CASE_RE = re.compile(r"[A-Za-z0-9._-]{3,100}")
DOCX_MT = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
CLASSIFICATIONS = ("UNCLASSIFIED", "RESTRICTED", "CONFIDENTIAL", "SECRET")


def _now() -> datetime:
    return datetime.now(timezone.utc).replace(tzinfo=None)


def _iso(d) -> str | None:
    return d.isoformat() if d else None


def _naive(d):
    if d is None:
        return None
    return d.replace(tzinfo=None) if getattr(d, "tzinfo", None) else d


def _loads(text, default=None):
    try:
        return json.loads(text) if text else (default if default is not None else {})
    except (ValueError, TypeError):
        return default if default is not None else {}


def _clean_case(value: str) -> str:
    cn = (value or "").strip()
    if not CASE_RE.fullmatch(cn) or set(cn) <= {"."}:
        raise HTTPException(400, "Invalid case number")
    return cn


def _cd(name: str) -> dict:
    safe = re.sub(r"[^A-Za-z0-9._-]+", "_", name)[:120] or "download"
    return {"Content-Disposition": f'attachment; filename="{safe}"'}


# ---------------------------------------------------------------------------------------------- models
class CaseGovernance(Base):
    __tablename__ = "case_governance"
    case_number = Column(String(100), primary_key=True)
    legal_hold = Column(Boolean, default=False)
    hold_reason = Column(Text, default="")
    hold_by = Column(String(100), default="")
    hold_at = Column(DateTime, nullable=True)
    retention_days = Column(Integer, nullable=True)
    classification = Column(String(30), default="RESTRICTED")


class ImportProfile(Base):
    __tablename__ = "import_profiles"
    id = Column(Integer, primary_key=True)
    name = Column(String(120), unique=True, index=True)
    mapping = Column(Text, default="{}")
    created_by = Column(String(100), default="")
    created_at = Column(DateTime, default=_now)


# ---------------------------------------------------------------------------------------------- request models
class RedactIn(BaseModel):
    phones: bool = True
    accounts: bool = True
    emails: bool = True
    vehicles: bool = True
    national_ids: bool = True
    names: bool = False
    addresses: bool = False
    free_text_patterns: list[str] = Field(default_factory=list, max_length=10)


class RedactedExportIn(BaseModel):
    case_number: str = Field(min_length=3, max_length=100)
    format: str = Field(default="json", pattern="^(json|csv|docx)$")
    redact: RedactIn = Field(default_factory=RedactIn)
    mode: str = Field(default="pseudonymise", pattern="^(mask|pseudonymise|remove)$")
    preview: bool = False


class HoldIn(BaseModel):
    legal_hold: bool
    reason: str = Field(min_length=5, max_length=500)


class RetentionIn(BaseModel):
    retention_days: int | None = Field(default=None, ge=1, le=36500)
    classification: str | None = Field(default=None, max_length=30)


class PurgeIn(BaseModel):
    confirm_case_number: str = Field(max_length=100)
    reason: str = Field(min_length=15, max_length=500)
    dry_run: bool = True


class BackupIn(BaseModel):
    passphrase: str = Field(min_length=12, max_length=256)


class BackupVerifyIn(BaseModel):
    passphrase: str = Field(min_length=1, max_length=256)


class ImportCommitIn(BaseModel):
    upload_token: str = Field(min_length=8, max_length=80)
    case_number: str = Field(min_length=3, max_length=100)
    mapping: dict[str, str] = Field(default_factory=dict)
    profile_name: str = Field(default="", max_length=120)


class ProfileIn(BaseModel):
    name: str = Field(min_length=2, max_length=120)
    mapping: dict[str, str] = Field(default_factory=dict)


class ScanIn(BaseModel):
    min_age_seconds: int = Field(default=0, ge=0, le=3600)


# ---------------------------------------------------------------------------------------------- dashboard
def _fetch(db: Session, cols, filter_col, values, *extra):
    """SELECT cols WHERE filter_col IN values (chunked so large deployments stay under SQLite's variable limit)."""
    out = []
    values = list(values)
    for i in range(0, len(values), 500):
        q = db.query(*cols).filter(filter_col.in_(values[i:i + 500]))
        for f in extra:
            q = q.filter(f)
        out.extend(q.all())
    return out


def _pct(values: list[float], q: float) -> float | None:
    if not values:
        return None
    s = sorted(values)
    k = (len(s) - 1) * q
    lo, hi = int(k), min(int(k) + 1, len(s) - 1)
    return round(s[lo] + (s[hi] - s[lo]) * (k - lo), 2)


def _bucket(days: float | None, edges: list[tuple[str, float]]) -> str:
    for label, upper in edges:
        if days is not None and days <= upper:
            return label
    return edges[-1][0]


_CASE_AGE_EDGES = [("0-30 days", 30), ("31-90 days", 90), ("91-180 days", 180), ("181-365 days", 365), (">1 year", 1e9)]
_BACKLOG_EDGES = [("<1 day", 1), ("1-3 days", 3), ("3-7 days", 7), ("7-14 days", 14), (">14 days", 1e9)]
_DONE_TASK = {"done", "closed", "completed", "complete", "cancelled", "canceled"}


def _workflow_snapshot(db: Session, cnset: set[str], now: datetime) -> dict:
    """Pending approvals / overdue deadlines from the workflow feature's tables, IF they exist (found by name only)."""
    out = {"available": False, "approvals_pending": None, "deadlines_overdue": None, "deadlines_open": None, "source_tables": []}
    try:
        insp = sa_inspect(db.get_bind())
        names = insp.get_table_names()
        for t in names:
            low = t.lower()
            kind = "approval" if "approval" in low else ("deadline" if "deadline" in low else None)
            if not kind or not re.fullmatch(r"[A-Za-z0-9_]+", t):
                continue
            cols = {c["name"] for c in insp.get_columns(t)}
            if "status" not in cols:
                continue
            due_col = next((c for c in ("due_at", "due_date", "due", "deadline", "due_on", "due_time", "deadline_at") if c in cols), None)
            has_case = "case_number" in cols
            select = ["status"] + (["case_number"] if has_case else []) + ([due_col] if due_col else [])
            rows = db.execute(sa.text(f'SELECT {", ".join(select)} FROM "{t}" LIMIT 50000')).fetchall()
            out["available"] = True
            out["source_tables"].append(t)
            n_pending = n_overdue = n_open = 0
            for r in rows:
                m = dict(zip(select, r))
                if has_case and m.get("case_number") not in cnset:
                    continue
                st = str(m.get("status") or "").lower()
                if kind == "approval":
                    if st in {"pending", "requested", "submitted", "awaiting", "open"}:
                        n_pending += 1
                else:
                    if st in {"done", "completed", "complete", "closed", "met", "cancelled", "canceled", "approved", "rejected"}:
                        continue
                    n_open += 1
                    due = m.get(due_col) if due_col else None
                    if isinstance(due, str):
                        try:
                            due = datetime.fromisoformat(due)
                        except ValueError:
                            due = None
                    if due and _naive(due) < now:
                        n_overdue += 1
            if kind == "approval":
                out["approvals_pending"] = (out["approvals_pending"] or 0) + n_pending
            else:
                out["deadlines_overdue"] = (out["deadlines_overdue"] or 0) + n_overdue
                out["deadlines_open"] = (out["deadlines_open"] or 0) + n_open
    except Exception:  # noqa: BLE001 - optional data, must never break the dashboard
        log.exception("workflow snapshot failed")
    return out


def _dashboard(ctx, db: Session, user: dict, days: int) -> dict:
    now = _now()
    since = now - timedelta(days=days)
    cases = ctx.svc.cases(db, user)
    cns = [c["case_number"] for c in cases]
    cnset = set(cns)
    title_of = {c["case_number"]: c["title"] for c in cases}

    # ---- cases
    by_status = Counter(c["status"] or "Unknown" for c in cases)
    open_cases = [c for c in cases if (c["status"] or "") not in {"Closed", "Archived"}]
    age_buckets = Counter()
    oldest_open = []
    for c in open_cases:
        created = _naive(datetime.fromisoformat(c["created_at"])) if c.get("created_at") else None
        age = (now - created).total_seconds() / 86400 if created else None
        age_buckets[_bucket(age, _CASE_AGE_EDGES) if age is not None else "Unknown"] += 1
        oldest_open.append((age or 0, c["case_number"], c["title"], c["status"]))
    oldest_open.sort(reverse=True)

    # ---- documents
    D = DocumentRecord
    docs = _fetch(db, [D.document_id, D.case_number, D.filename, D.status, D.created_at, D.reviewed_at, D.review_actor], D.case_number, cns)
    pending = [d for d in docs if (d.status or "") == "Pending Review"]
    backlog_ages = [max(0.0, (now - _naive(d.created_at)).total_seconds() / 86400) for d in pending if d.created_at]
    backlog_buckets = Counter(_bucket(a, _BACKLOG_EDGES) for a in backlog_ages)
    oldest_pending = sorted(({"document_id": d.document_id, "case_number": d.case_number, "filename": d.filename,
                              "age_days": round((now - _naive(d.created_at)).total_seconds() / 86400, 1)}
                             for d in pending if d.created_at), key=lambda x: -x["age_days"])[:10]
    reviewed = [d for d in docs if d.reviewed_at and d.created_at and _naive(d.reviewed_at) >= since and d.status in {"Verified", "Rejected"}]
    turnaround_h = [max(0.0, (_naive(d.reviewed_at) - _naive(d.created_at)).total_seconds() / 3600) for d in reviewed]
    verification = {
        "reviewed_in_window": len(reviewed), "median_hours": _pct(turnaround_h, 0.5), "p90_hours": _pct(turnaround_h, 0.9),
        "mean_hours": round(statistics.fmean(turnaround_h), 2) if turnaround_h else None,
        "basis": "hours from document creation to human review (Verified or Rejected), for reviews completed in the window",
    }

    # ---- evidence
    E = Evidence
    ev = _fetch(db, [E.evidence_id, E.case_number, E.integrity_status, E.stored_path, E.created_at], E.case_number, cns)
    integrity = Counter((e.integrity_status or "unknown") for e in ev)
    missing = 0
    for e in ev[:5000]:
        try:
            if not A._abs_path(e.stored_path).exists():
                missing += 1
        except Exception:  # noqa: BLE001
            missing += 1

    # ---- weekly ingestion
    def week(d):
        d = _naive(d)
        return (d - timedelta(days=d.weekday())).date() if d else None
    weeks = Counter(); ev_weeks = Counter()
    for d in docs:
        if d.created_at and _naive(d.created_at) >= since:
            weeks[week(d.created_at)] += 1
    for e in ev:
        if e.created_at and _naive(e.created_at) >= since:
            ev_weeks[week(e.created_at)] += 1
    ingestion = []
    w = week(since)
    while w and w <= week(now):
        ingestion.append({"week": w.isoformat(), "documents": weeks.get(w, 0), "evidence": ev_weeks.get(w, 0)})
        w += timedelta(days=7)

    # ---- alerts
    al = _fetch(db, [Alert.status, Alert.severity], Alert.case_number, cns)
    alerts_by_status = Counter(a.status or "Unknown" for a in al)
    alerts_by_sev = Counter((a.severity or "Unknown") for a in al)
    open_alerts = sum(1 for a in al if (a.status or "New") in {"New", "Reviewing", "Under Review", "Escalated"})

    # ---- officers
    workload: dict[str, Counter] = defaultdict(Counter)
    for c in cases:
        if c.get("created_by"):
            workload[c["created_by"]]["cases_created"] += 1
    for d in reviewed:
        if d.review_actor:
            workload[d.review_actor]["documents_verified" if d.status == "Verified" else "documents_rejected"] += 1
    for t in _fetch(db, [Task.assignee, Task.status], Task.case_number, cns):
        if t.assignee and (t.status or "").lower() not in _DONE_TASK:
            workload[t.assignee]["tasks_open"] += 1
    for n in _fetch(db, [Note.author, Note.created_at], Note.case_number, cns):
        if n.author and n.created_at and _naive(n.created_at) >= since:
            workload[n.author]["notes"] += 1
    roles = {u.username: (u.role, bool(u.active)) for u in db.query(User.username, User.role, User.active).all()}
    officers = [{"username": u, "role": roles.get(u, ("", True))[0], "active": roles.get(u, ("", True))[1],
                 "cases_created": c["cases_created"], "documents_verified": c["documents_verified"],
                 "documents_rejected": c["documents_rejected"], "tasks_open": c["tasks_open"], "notes": c["notes"]}
                for u, c in workload.items()]
    officers.sort(key=lambda o: -(o["documents_verified"] + o["cases_created"] + o["tasks_open"] + o["notes"]))

    # ---- most active cases (audited actions that cite the case)
    activity = Counter()
    rows = (db.query(AuditEvent.details).filter(AuditEvent.created_at >= since, ~AuditEvent.action.like("auth.%"))
            .order_by(AuditEvent.id.desc()).limit(20000).all())
    for (details,) in rows:
        cn = _loads(details, {}).get("case_number") if details and details.startswith("{") else None
        if cn in cnset:
            activity[cn] += 1
    top_cases = [{"case_number": cn, "title": title_of.get(cn, ""), "audited_actions": n} for cn, n in activity.most_common(8)]

    return {
        "window_days": days, "generated_at": now.isoformat(), "scope": {"cases": len(cns), "all_cases": user.get("role") == "admin"},
        "cases": {"total": len(cases), "open": len(open_cases), "by_status": dict(by_status),
                  "age_buckets_open": {k: age_buckets.get(k, 0) for k, _ in _CASE_AGE_EDGES if age_buckets.get(k)},
                  "oldest_open": [{"case_number": cn, "title": t, "status": s, "age_days": round(a, 1)} for a, cn, t, s in oldest_open[:8]]},
        "officers": officers,
        "verification": verification,
        "backlog": {"unverified_documents": len(pending), "age_buckets": dict(backlog_buckets),
                    "oldest_days": max(backlog_ages) if backlog_ages else None, "oldest": oldest_pending},
        "ingestion_weekly": ingestion,
        "alerts": {"total": len(al), "open": open_alerts, "by_status": dict(alerts_by_status), "by_severity": dict(alerts_by_sev)},
        "evidence": {"total": len(ev), "by_integrity": dict(integrity), "files_missing": missing},
        "workflow": _workflow_snapshot(db, cnset, now),
        "top_cases": top_cases,
        "note": "All figures are counts of stored records for the cases you can access; nothing is estimated.",
    }


# ---------------------------------------------------------------------------------------------- redaction engine
_PATTERNS: list[tuple[str, str]] = [
    # order = priority when two patterns could match at the same position
    ("emails", r"[A-Za-z0-9._%+-]+@[A-Za-z0-9-]+(?:\.[A-Za-z0-9-]+)+"),
    ("national_ids", r"(?<![\dA-Za-z])[A-Z]{5}\d{4}[A-Z](?![\dA-Za-z])"),                 # PAN
    ("national_ids", r"(?<!\d)[2-9]\d{3}[ -]\d{4}[ -]\d{4}(?!\d)"),                        # Aadhaar, grouped
    ("national_ids", r"(?<![\d.])[2-9]\d{11}(?![\d])"),                                    # Aadhaar, 12 contiguous digits
    ("phones", r"(?<![\dA-Za-z])(?:(?:\+|00)?91[ -]?|0)?[6-9]\d{4}[ -]?\d{5}(?!\d)"),
    ("vehicles", r"(?<![A-Za-z0-9])[A-Z]{2}[ -]?\d{1,2}[ -]?[A-Z]{1,3}[ -]?\d{3,4}(?![A-Za-z0-9])"),
    ("accounts", r"(?<![A-Za-z0-9])(?:A/C|ACCT?|AC)[ .:#_-]*\d{4,}(?!\d)"),
    ("accounts", r"(?<![\d.])\d{9,18}(?!\d)"),
]
_PREFIX = {"phones": "PHONE", "accounts": "ACCOUNT", "emails": "EMAIL", "vehicles": "VEHICLE", "national_ids": "NATID",
           "names": "PERSON", "addresses": "ADDRESS", "free_text_patterns": "TEXT", "ENTID": "ENT"}
_TYPE_CAT = {"PHONE": "phones", "ACCOUNT": "accounts", "EMAIL": "emails", "VEHICLE": "vehicles", "PERSON": "names",
             "LOCATION": "addresses", "ADDRESS": "addresses"}
_ATTR_CAT = {"phone": "phones", "mobile": "phones", "msisdn": "phones", "imei": "phones", "imsi": "phones", "contact": "phones",
             "account": "accounts", "account_number": "accounts", "acct": "accounts", "iban": "accounts",
             "email": "emails", "e-mail": "emails",
             "plate": "vehicles", "vehicle": "vehicles", "registration": "vehicles", "reg_no": "vehicles",
             "aadhaar": "national_ids", "pan": "national_ids", "national_id": "national_ids", "passport": "national_ids",
             "address": "addresses", "addr": "addresses",
             "name": "names", "alias": "names", "owner": "names", "holder": "names", "father_name": "names", "complainant": "names"}
_DOC_CAT = {"complainant": "names", "father_name": "names", "complainant_address": "addresses", "place_of_occurrence": "addresses",
            "phones": "phones", "vehicles": "vehicles", "accounts": "accounts", "emails": "emails"}
_KEEP_KEYS = {"case_number", "case_numbers", "source_file", "document_id", "evidence_id", "fir_number", "year", "sha256"}
_NESTED_QUANT = re.compile(r"\((?:[^()\\]|\\.)*[+*](?:[^()\\]|\\.)*\)\s*[+*{]")


def _norm_value(cat: str, raw: str) -> str:
    s = str(raw or "")
    if cat == "phones":
        d = re.sub(r"\D", "", s)
        return d[-10:] if len(d) >= 10 else d
    if cat in {"accounts", "national_ids", "vehicles", "ENTID"}:
        return re.sub(r"[^A-Za-z0-9]", "", s).upper() if cat != "ENTID" else s
    if cat == "emails":
        return s.strip().lower()
    return re.sub(r"\s+", " ", s).strip().casefold()


def _mask(cat: str, raw: str) -> str:
    s = str(raw)
    if cat == "emails":
        user, _, dom = s.partition("@")
        head, _, tld = dom.rpartition(".")
        return f"{user[:1]}***@{head[:1]}***.{tld}" if dom else "***"
    if cat == "names":
        return " ".join((w[:1] + "***") if w else w for w in s.split()) or "***"
    if cat == "addresses":
        return "[ADDRESS MASKED]"
    if cat == "free_text_patterns":
        return "[REDACTED]"
    alnum = [i for i, ch in enumerate(s) if ch.isalnum()]
    keep = set(alnum[-4:]) if cat in {"phones", "accounts", "vehicles", "national_ids"} else set()
    return "".join(ch if (not ch.isalnum() or i in keep) else "X" for i, ch in enumerate(s))


class Pseudonymizer:
    """Consistent labels within one export. Numbering follows a per-export salted HMAC order, so the same value gets
    a different label in a different export (labels cannot be joined across exports) while links inside one export
    are preserved. Two passes: collect, then finalise."""

    def __init__(self, salt: bytes):
        self.salt = salt
        self.seen: dict[str, set[str]] = defaultdict(set)
        self.final: dict[tuple[str, str], str] | None = None

    def label(self, cat: str, norm: str) -> str:
        if self.final is None:
            self.seen[cat].add(norm)
            return f"{_PREFIX.get(cat, 'X')}-000"
        return self.final[(cat, norm)]

    def finalize(self):
        self.final = {}
        for cat, values in self.seen.items():
            ordered = sorted(values, key=lambda v: hmac.new(self.salt, f"{cat}|{v}".encode(), hashlib.sha256).digest())
            for i, v in enumerate(ordered, 1):
                self.final[(cat, v)] = f"{_PREFIX.get(cat, 'X')}-{i:03d}"


class Redactor:
    def __init__(self, cats: set[str], mode: str, pseudo: Pseudonymizer, names=(), addresses=(), extra_patterns=()):
        self.cats, self.mode, self.pseudo = set(cats), mode, pseudo
        self.counts: Counter = Counter()
        self.distinct: dict[str, set] = defaultdict(set)
        self.fields_touched = 0
        self.muted = False
        self.extra_regex: list[re.Pattern] = []
        alts, self._alt_cat = [], {}
        for i, (cat, pat) in enumerate(_PATTERNS):
            if cat in self.cats:
                alts.append(f"(?P<p{i}>{pat})"); self._alt_cat[f"p{i}"] = cat
        n_list = sorted({re.sub(r"\s+", " ", n).strip() for n in names if n and len(str(n).strip()) >= 3}, key=lambda x: (-len(x), x))
        if "names" in self.cats:
            for i, n in enumerate(n_list):
                alts.append(f"(?P<n{i}>(?i:(?<!\\w){re.escape(n).replace(chr(92) + ' ', chr(92) + 's+')}(?!\\w)))"); self._alt_cat[f"n{i}"] = "names"
        a_list = sorted({re.sub(r"\s+", " ", a).strip() for a in addresses if a and len(str(a).strip()) >= 6}, key=lambda x: (-len(x), x))
        if "addresses" in self.cats:
            for i, a in enumerate(a_list):
                alts.append(f"(?P<a{i}>(?i:{re.escape(a).replace(chr(92) + ' ', chr(92) + 's+')}))"); self._alt_cat[f"a{i}"] = "addresses"
        self.literal_terms = []
        if "free_text_patterns" in self.cats:
            for j, raw in enumerate(extra_patterns):
                if raw.startswith("re:"):
                    self.extra_regex.append(re.compile(raw[3:], re.I))
                else:
                    alts.append(f"(?P<t{j}>(?i:{re.escape(raw)}))"); self._alt_cat[f"t{j}"] = "free_text_patterns"
        self.rx = re.compile("|".join(alts)) if alts else None

    # -- replacement of one detected value
    def token(self, cat: str, raw: str) -> str:
        raw = str(raw)
        norm = _norm_value(cat, raw)
        if not self.muted:
            self.counts[cat] += 1
            self.distinct[cat].add(norm)
        if self.mode == "remove":
            return "[REMOVED]"
        if self.mode == "mask":
            return _mask(cat, raw)
        return self.pseudo.label(cat, norm)

    def opaque_id(self, raw: str) -> str:
        return self.pseudo.label("ENTID", _norm_value("ENTID", raw))

    def text(self, s):
        if not isinstance(s, str) or not s:
            return s
        out = s
        if self.rx:
            out = self.rx.sub(lambda m: self.token(self._alt_cat[m.lastgroup], m.group(0)), out)
        for rx in self.extra_regex:
            out = rx.sub(lambda m: self.token("free_text_patterns", m.group(0)), out)
        if out != s and not self.muted:
            self.fields_touched += 1
        return out

    def field(self, cat: str, value):
        """Redact a whole field of a known category (e.g. an entity's name or a phone attribute)."""
        if value in (None, ""):
            return value
        if cat in self.cats:
            if not self.muted:
                self.fields_touched += 1
            return self.token(cat, value)
        return self.text(value)

    def deep(self, value, key: str = ""):
        k = (key or "").lower()
        if k in _KEEP_KEYS:
            return value
        if isinstance(value, dict):
            return {kk: self.deep(vv, str(kk)) for kk, vv in value.items()}
        if isinstance(value, list):
            return [self.deep(v, key) for v in value]
        if isinstance(value, str):
            cat = _ATTR_CAT.get(k)
            return self.field(cat, value) if cat else self.text(value)
        return value


def _compile_free_patterns(raw: list[str]) -> list[str]:
    out = []
    for p in raw:
        p = (p or "").strip()
        if not p:
            continue
        if len(p) > 100:
            raise HTTPException(400, "Custom redaction patterns are limited to 100 characters")
        if p.startswith("re:"):
            body = p[3:]
            if _NESTED_QUANT.search(body):
                raise HTTPException(400, f"Pattern '{p}' contains a nested repetition that could hang the server; simplify it")
            try:
                re.compile(body)
            except re.error as exc:
                raise HTTPException(400, f"Invalid regular expression: {exc}")
        out.append(p)
    return out


def redact_text(text: str, categories=("phones", "accounts", "emails", "vehicles", "national_ids"), mode: str = "mask") -> str:
    """Convenience for callers/tests: redact PII patterns in free text (no name / address knowledge)."""
    pz = Pseudonymizer(secrets.token_bytes(16))
    r = Redactor(set(categories), mode, pz)
    if mode == "pseudonymise":
        r.text(text); pz.finalize(); r = Redactor(set(categories), mode, pz)
    return r.text(text)


def _collect_export_data(ctx, db: Session, user: dict, cn: str) -> dict:
    svc = ctx.svc
    ents = svc.entities(db, cn, role=user["role"])
    rels = svc.relationships(db, cn)
    evs = svc.events(db, cn)
    docs = svc.documents(db, cn)
    return {"entities": ents, "relationships": rels, "events": evs, "documents": docs}


def _known_names_and_addresses(data: dict) -> tuple[set, set]:
    names, addrs = set(), set()
    for e in data["entities"]:
        t = str(e.get("type", "")).upper()
        if t == "PERSON" and e.get("name"):
            names.add(str(e["name"]))
        if t in {"LOCATION", "ADDRESS"} and e.get("name"):
            addrs.add(str(e["name"]))
        for k, v in (e.get("attributes") or {}).items():
            if isinstance(v, str) and _ATTR_CAT.get(str(k).lower()) == "names":
                names.add(v)
            if isinstance(v, str) and _ATTR_CAT.get(str(k).lower()) == "addresses":
                addrs.add(v)
    for d in data["documents"]:
        st = d.get("structured") or {}
        for k in ("complainant", "father_name"):
            if st.get(k):
                names.add(str(st[k]))
        for k in ("complainant_address", "place_of_occurrence"):
            if st.get(k):
                addrs.add(str(st[k]))
    return names, addrs


def _redact_bundle(data: dict, R: Redactor) -> dict:
    """Apply the redactor to every structured field. Entity ids that embed PII are replaced by opaque ids consistently
    in entities, relationships and events so the graph structure survives."""
    id_cache: dict[str, str] = {}

    def map_id(raw):
        if raw in (None, ""):
            return raw
        if raw not in id_cache:
            R.muted = True
            try:
                changed = R.text(raw) != raw
            finally:
                R.muted = False
            id_cache[raw] = R.opaque_id(raw) if changed else raw
        return id_cache[raw]

    ents = []
    for e in data["entities"]:
        t = str(e.get("type", "")).upper()
        cat = _TYPE_CAT.get(t)
        raw_id = e.get("external_id")
        if cat and cat in R.cats:
            name = R.field(cat, e.get("name"))
        else:
            name = R.text(e.get("name"))
        ent = {"external_id": map_id(raw_id), "name": name, "type": e.get("type"), "confidence": e.get("confidence"),
               "attributes": R.deep(e.get("attributes") or {}),
               "latitude": None if "addresses" in R.cats else e.get("latitude"),
               "longitude": None if "addresses" in R.cats else e.get("longitude")}
        ents.append(ent)
    rels = [{"id": r["id"], "source": map_id(r["source"]), "target": map_id(r["target"]), "relation": r["relation"],
             "confidence": r["confidence"], "verification_state": r["verification_state"], "event_time": r["event_time"],
             "source_ref": R.text(r.get("source_ref")), "metadata": R.deep(r.get("metadata") or {})} for r in data["relationships"]]
    evs = [{"event_id": e["event_id"], "event_type": e["event_type"], "entity_id": map_id(e["entity_id"]),
            "related_entity_id": map_id(e.get("related_entity_id")), "event_time": e["event_time"], "amount": e["amount"],
            "duration_seconds": e["duration_seconds"], "source_ref": R.text(e.get("source_ref")),
            "latitude": None if "addresses" in R.cats else e.get("latitude"),
            "longitude": None if "addresses" in R.cats else e.get("longitude"),
            "metadata": R.deep(e.get("metadata") or {})} for e in data["events"]]
    docs = []
    for d in data["documents"]:
        st = d.get("structured") or {}
        new_st = {}
        for k, v in st.items():
            cat = _DOC_CAT.get(k)
            if k in _KEEP_KEYS:
                new_st[k] = v
            elif cat and isinstance(v, str):
                new_st[k] = R.field(cat, v)
            elif cat and isinstance(v, list):
                new_st[k] = [R.field(cat, x) if isinstance(x, str) else x for x in v]
            else:
                new_st[k] = R.deep(v, k)
        docs.append({"document_id": d["document_id"], "filename": R.text(d.get("filename")), "document_type": d.get("document_type"),
                     "status": d.get("status"), "fir_number": d.get("fir_number"), "sha256": d.get("sha256"),
                     "evidence_id": d.get("evidence_id"), "created_at": d.get("created_at"), "reviewed_at": d.get("reviewed_at"),
                     "structured": new_st})
    return {"entities": ents, "relationships": rels, "events": evs, "documents": docs}


def _run_redaction(data: dict, cats: set[str], mode: str, extra: list[str]) -> tuple[dict, dict]:
    names, addrs = _known_names_and_addresses(data)
    salt = secrets.token_bytes(16)
    pz = Pseudonymizer(salt)
    # pass 1 collects the distinct values, pass 2 emits the final labels
    R1 = Redactor(cats, mode, pz, names, addrs, extra)
    _redact_bundle(data, R1)
    pz.finalize()
    R = Redactor(cats, mode, pz, names, addrs, extra)
    out = _redact_bundle(data, R)
    log_ = {
        "mode": mode, "categories": sorted(cats),
        "replacements_by_category": {c: R.counts.get(c, 0) for c in sorted(cats)},
        "distinct_values_by_category": {c: len(R.distinct.get(c, ())) for c in sorted(cats)},
        "fields_changed": R.fields_touched,
        "records": {k: len(v) for k, v in out.items()},
        "pseudonyms_consistent_within_export": mode == "pseudonymise",
        "limitations": ["Names are matched by exact known-name (entities and FIR fields), not by name recognition; "
                        "free-text names that never appear as an entity or FIR field can remain.",
                        "Regex detection is a heuristic (Aadhaar, PAN, Indian mobile numbers, e-mails, account numbers, vehicle plates); "
                        "review the output before sharing."],
    }
    return out, log_


def _export_csv(out: dict, log_: dict) -> bytes:
    buf = io.StringIO(newline="")
    w = csv.writer(buf)
    w.writerow(["section", "id", "name_or_relation", "type", "source", "target", "time", "amount", "detail_json"])
    for e in out["entities"]:
        w.writerow(["entity", e["external_id"], e["name"], e["type"], "", "", "", "", json.dumps(e["attributes"], ensure_ascii=False)])
    for r in out["relationships"]:
        w.writerow(["relationship", r["id"], r["relation"], r["verification_state"], r["source"], r["target"], r["event_time"] or "", "",
                    json.dumps(r["metadata"], ensure_ascii=False)])
    for e in out["events"]:
        w.writerow(["event", e["event_id"], e["event_type"], "", e["entity_id"], e["related_entity_id"], e["event_time"] or "",
                    e["amount"] if e["amount"] is not None else "", json.dumps(e["metadata"], ensure_ascii=False)])
    for d in out["documents"]:
        w.writerow(["document", d["document_id"], d["filename"], d["document_type"], "", "", d["created_at"] or "", "",
                    json.dumps({"status": d["status"], "structured": d["structured"]}, ensure_ascii=False)])
    w.writerow([])
    w.writerow(["redaction_log", "", "", "", "", "", "", "", json.dumps(log_, ensure_ascii=False)])
    return ("﻿" + buf.getvalue()).encode("utf-8")


# ---------------------------------------------------------------------------------------------- retention / hold / purge
def _gov_row(db: Session, cn: str, create: bool = False) -> CaseGovernance | None:
    row = db.query(CaseGovernance).filter_by(case_number=cn).first()
    if row is None and create:
        row = CaseGovernance(case_number=cn, legal_hold=False, classification="RESTRICTED")
        db.add(row)
    return row


def _retention_rows(db: Session) -> list[dict]:
    now = _now()
    gov = {g.case_number: g for g in db.query(CaseGovernance).all()}
    last: dict[str, datetime] = {}

    def fold(rows):
        for cn, ts in rows:
            if cn and ts and (cn not in last or _naive(ts) > last[cn]):
                last[cn] = _naive(ts)
    fold(db.query(DocumentRecord.case_number, sa.func.max(DocumentRecord.updated_at)).group_by(DocumentRecord.case_number).all())
    fold(db.query(Evidence.case_number, sa.func.max(Evidence.created_at)).group_by(Evidence.case_number).all())
    fold(db.query(Alert.case_number, sa.func.max(Alert.created_at)).group_by(Alert.case_number).all())
    fold(db.query(Note.case_number, sa.func.max(Note.created_at)).group_by(Note.case_number).all())
    out = []
    for c in db.query(Case).order_by(Case.case_number).all():
        g = gov.get(c.case_number)
        created = _naive(c.created_at)
        ref = max([x for x in (created, last.get(c.case_number)) if x], default=None)
        days = g.retention_days if g else None
        due = ref + timedelta(days=days) if (ref and days) else None
        past = bool(due and now > due)
        hold = bool(g and g.legal_hold)
        out.append({
            "case_number": c.case_number, "title": c.title, "status": c.status, "created_at": _iso(created),
            "last_activity": _iso(ref), "retention_days": days, "retention_due": _iso(due), "past_retention": past,
            "legal_hold": hold, "hold_reason": g.hold_reason if g and hold else "", "hold_by": g.hold_by if g and hold else "",
            "hold_at": _iso(g.hold_at) if g and hold else None,
            "classification": (g.classification if g and g.classification else "RESTRICTED"),
            "purge_review_candidate": past and not hold,
        })
    out.sort(key=lambda r: (not r["legal_hold"], not r["past_retention"], r["case_number"]))
    return out


_PROTECTED_TABLES = {"audit_events", "users", "ledger_blocks", "ledger_rows"}  # the ledger must outlive purges and demo-off


def _case_tables(db: Session) -> list[str]:
    insp = sa_inspect(db.get_bind())
    out = []
    for t in insp.get_table_names():
        if t in _PROTECTED_TABLES or not re.fullmatch(r"[A-Za-z0-9_]+", t):
            continue
        if "case_number" in {c["name"] for c in insp.get_columns(t)}:
            out.append(t)
    return sorted(out, key=lambda t: (t == "cases", t))


def _sha256_file(path: Path) -> str:
    from security import evidence_crypto
    if evidence_crypto.is_encrypted(path):  # hash of the plaintext, like everywhere else
        try:
            return evidence_crypto.sha256_plain(path)
        except evidence_crypto.EvidenceCryptoError as exc:
            raise OSError(str(exc)) from exc
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def _purge_inventory(ctx, db: Session, cn: str) -> dict:
    svc = ctx.svc
    root = A.EVIDENCE_DIR.resolve()
    tables = {}
    for t in _case_tables(db):
        n = db.execute(sa.text(f'SELECT COUNT(*) FROM "{t}" WHERE case_number = :cn'), {"cn": cn}).scalar() or 0
        if n:
            tables[t] = int(n)
    # ---- files (only ever paths that resolve inside the evidence directory)
    files, seen = [], set()
    others = set()
    for stored, in db.query(Evidence.stored_path).filter(Evidence.case_number != cn).all():
        others.add(stored)
    for stored, in db.query(DocumentRecord.stored_path).filter(DocumentRecord.case_number != cn).all():
        others.add(stored)
    refs = [(e.evidence_id, e.filename, e.stored_path, e.sha256) for e in db.query(Evidence).filter_by(case_number=cn).all()]
    refs += [(d.document_id, d.filename, d.stored_path, d.sha256) for d in
             db.query(DocumentRecord.document_id, DocumentRecord.filename, DocumentRecord.stored_path, DocumentRecord.sha256).filter_by(case_number=cn).all()]
    for ref_id, filename, stored, recorded in refs:
        if not stored:
            continue
        item = {"ref": ref_id, "filename": filename, "recorded_sha256": recorded, "sha256": None, "bytes": 0, "action": "delete", "path": None}
        try:
            p = A._abs_path(stored)
            resolved = p.resolve()
            item["path"] = str(resolved.relative_to(root)) if resolved.is_relative_to(root) else "(outside evidence directory)"
            if not resolved.is_relative_to(root):
                item["action"] = "skipped: outside the evidence directory, never touched"
            elif p.is_symlink():
                item["action"] = "skipped: symbolic link"
            elif stored in others:
                item["action"] = "skipped: file is also referenced by another case"
            elif not resolved.is_file():
                item["action"] = "skipped: file already missing"
            else:
                item["bytes"] = resolved.stat().st_size
                item["sha256"] = _sha256_file(resolved)
                if resolved in seen:
                    item["action"] = "duplicate reference"
                seen.add(resolved)
                item["_abs"] = resolved
        except OSError as exc:
            item["action"] = f"skipped: {exc.__class__.__name__}"
        files.append(item)
    # ---- graph rows
    ent_delete, ent_detach = [], []
    untagged = 0
    for e in svc._case_entities(db, cn):
        attrs = A.load_json(e.attributes, {})
        owners = set(attrs.get("case_numbers") or []) | ({attrs["case_number"]} if attrs.get("case_number") else set())
        if cn not in owners:
            untagged += 1
        elif owners == {cn}:
            ent_delete.append(e.id)
        else:
            ent_detach.append(e.id)
    rel_ids = [r.id for r in svc._case_relationships(db, cn, include_candidates=True)]
    ev_ids = [e.id for e in svc._case_events(db, cn)]
    return {"tables": tables, "files": files, "entity_ids": ent_delete, "detach_entity_ids": ent_detach, "untagged": untagged,
            "relationship_ids": rel_ids, "event_ids": ev_ids}


def _inventory_summary(inv: dict) -> dict:
    deletable = [f for f in inv["files"] if f["action"] == "delete"]
    return {
        "database_rows_by_table": inv["tables"],
        "entities_deleted": len(inv["entity_ids"]), "entities_shared_with_other_cases_detached": len(inv["detach_entity_ids"]),
        "entities_left_untouched_not_tagged_to_case": inv["untagged"],
        "relationships_deleted": len(inv["relationship_ids"]), "events_deleted": len(inv["event_ids"]),
        "files_deleted": len(deletable), "bytes_deleted": sum(f["bytes"] for f in deletable),
        "files_skipped": len(inv["files"]) - len(deletable),
    }


def _public_files(inv: dict) -> list[dict]:
    return [{k: v for k, v in f.items() if not k.startswith("_")} for f in inv["files"]]


def _delete_chunks(db: Session, model, ids: list[int]) -> None:
    for i in range(0, len(ids), 500):
        db.query(model).filter(model.id.in_(ids[i:i + 500])).delete(synchronize_session=False)


def _purge_case(ctx, db: Session, user: dict, cn: str, x: PurgeIn) -> dict:
    svc = ctx.svc
    if not db.query(Case).filter_by(case_number=cn).first():
        raise HTTPException(404, "Case not found")
    if x.confirm_case_number != cn:
        raise HTTPException(400, "confirm_case_number must exactly match the case number in the URL")
    g = _gov_row(db, cn)
    if g and g.legal_hold:
        raise HTTPException(409, "Case is under legal hold and cannot be purged. Release the hold first (with a recorded reason).")
    before = svc.verify_audit_chain(db)
    if not before.get("verified"):
        raise HTTPException(409, "The audit chain does not verify; purging is refused because it could not be proven afterwards. Investigate first.")
    inv = _purge_inventory(ctx, db, cn)
    summary = _inventory_summary(inv)
    base = {"case_number": cn, "dry_run": x.dry_run, "summary": summary, "files": _public_files(inv), "audit_chain_before": before}
    if x.dry_run:
        svc.audit(db, user["sub"], "governance.purge_dry_run", {"case_number": cn, "summary": summary, "reason": x.reason})
        base["message"] = "Dry run: nothing was deleted. Re-submit with dry_run=false to purge."
        return base

    deletable = [f for f in inv["files"] if f["action"] == "delete"]
    try:
        # graph rows first (found through the case tags), then every table that carries case_number
        _delete_chunks(db, Relationship, inv["relationship_ids"])
        _delete_chunks(db, Event, inv["event_ids"])
        _delete_chunks(db, Entity, inv["entity_ids"])
        for ent in db.query(Entity).filter(Entity.id.in_(inv["detach_entity_ids"] or [-1])).all():
            attrs = A.load_json(ent.attributes, {})
            owners = [c for c in (attrs.get("case_numbers") or []) if c != cn]
            attrs["case_numbers"] = owners
            if attrs.get("case_number") == cn:
                attrs["case_number"] = owners[0] if owners else ""
            ent.attributes = A.j(attrs)
        for t in _case_tables(db):
            db.execute(sa.text(f'DELETE FROM "{t}" WHERE case_number = :cn'), {"cn": cn})
        db.commit()
    except Exception:  # noqa: BLE001
        db.rollback()
        log.exception("purge failed for %s", cn)
        raise HTTPException(500, "Purge failed and was rolled back; no data was deleted")
    file_failures = []
    for f in deletable:
        try:
            f["_abs"].unlink()
        except OSError as exc:
            file_failures.append({"ref": f["ref"], "error": exc.__class__.__name__})
    manifest = [{"ref": f["ref"], "sha256": f["sha256"], "bytes": f["bytes"]} for f in deletable]
    tomb = {"case_number": cn, "reason": x.reason, "summary": summary, "purged_files": manifest[:500],
            "files_manifest_sha256": hashlib.sha256(json.dumps(manifest, sort_keys=True).encode()).hexdigest(),
            "file_delete_failures": file_failures, "tombstone": True}
    svc.audit(db, user["sub"], "governance.case_purged", tomb)
    ctx.hooks.emit("case_purged", db=db, user=user, case_number=cn, detail={"files": len(manifest), "reason_sha": tomb["files_manifest_sha256"]})
    after = svc.verify_audit_chain(db)
    base.update({"message": "Case purged. A tombstone event with counts and file hashes was written to the audit log.",
                 "audit_chain_after": after, "file_delete_failures": file_failures,
                 "files_manifest_sha256": tomb["files_manifest_sha256"]})
    return base


# ---------------------------------------------------------------------------------------------- backup
_MAGIC = b"DCNBAK1\n"
_SCRYPT = {"n": 2 ** 15, "r": 8, "p": 1}
_BACKUP_NAME_RE = re.compile(r"dcn-backup-\d{8}T\d{6}Z-[0-9a-f]{6}\.dcnbak")


def backup_max_bytes() -> int:
    try:
        return max(1, int(os.environ.get("BACKUP_MAX_MB", "200"))) * 1024 * 1024
    except ValueError:
        return 200 * 1024 * 1024


def _backup_dir() -> Path:
    d = A.EVIDENCE_DIR.resolve().parent / "backups"
    d.mkdir(parents=True, exist_ok=True)
    try:
        os.chmod(d, 0o700)
    except OSError:
        pass
    return d


def _kdf(passphrase: str, salt: bytes, n: int, r: int, p: int) -> bytes:
    raw = hashlib.scrypt(passphrase.encode("utf-8"), salt=salt, n=n, r=r, p=p, maxmem=256 * 1024 * 1024, dklen=32)
    return base64.urlsafe_b64encode(raw)


def _encrypt_blob(plain: bytes, passphrase: str) -> bytes:
    from cryptography.fernet import Fernet
    salt = secrets.token_bytes(16)
    header = {"v": 1, "kdf": "scrypt", **_SCRYPT, "salt": base64.b64encode(salt).decode(), "cipher": "fernet"}
    return _MAGIC + json.dumps(header).encode() + b"\n" + Fernet(_kdf(passphrase, salt, **_SCRYPT)).encrypt(plain)


def _parse_header(blob: bytes) -> tuple[dict, bytes]:
    if not blob.startswith(_MAGIC):
        raise ValueError("Not a DARK CRIMENET backup file")
    nl = blob.find(b"\n", len(_MAGIC))
    if nl < 0 or nl > 2048:
        raise ValueError("Backup header is corrupt")
    try:
        header = json.loads(blob[len(_MAGIC):nl])
        salt = base64.b64decode(header["salt"])
        n, r, p = int(header["n"]), int(header["r"]), int(header["p"])
    except (ValueError, KeyError, TypeError) as exc:
        raise ValueError("Backup header is corrupt") from exc
    if header.get("kdf") != "scrypt" or n < 2 ** 12 or n > 2 ** 20 or n & (n - 1) or not (1 <= r <= 16) or not (1 <= p <= 4) or len(salt) < 8:
        raise ValueError("Backup header has unsupported or unsafe key-derivation parameters")
    header["_salt"], header["_n"], header["_r"], header["_p"] = salt, n, r, p
    return header, blob[nl + 1:]


def _decrypt_blob(blob: bytes, passphrase: str) -> bytes:
    from cryptography.fernet import Fernet, InvalidToken
    header, token = _parse_header(blob)
    try:
        return Fernet(_kdf(passphrase, header["_salt"], header["_n"], header["_r"], header["_p"])).decrypt(token)
    except InvalidToken as exc:
        raise ValueError("Decryption failed: wrong passphrase, or the backup file was modified or corrupted") from exc


def _sqlite_path() -> Path:
    if engine.url.get_backend_name() != "sqlite" or not engine.url.database or engine.url.database == ":memory:":
        raise HTTPException(400, "Backups are only supported for file-based SQLite databases")
    return Path(engine.url.database).resolve()


def _app_version() -> str:
    try:
        import sys
        return str(getattr(sys.modules.get("backend.app.main"), "APP_VERSION", "unknown"))
    except Exception:  # noqa: BLE001
        return "unknown"


def _safe_zip_name(name: str) -> bool:
    return bool(name) and not name.startswith(("/", "\\")) and ".." not in Path(name.replace("\\", "/")).parts and ":" not in name


def _chain_of_sqlite(path: str, secret_key: str) -> dict:
    """Recompute the audit chain of a SQLite file with the given server secret (read-only)."""
    conn = sqlite3.connect(f"file:{path}?mode=ro", uri=True)
    try:
        prev, n = "GENESIS", 0
        for rid, actor, action, details, phash, ehash, created in conn.execute(
                "SELECT id, actor, action, details, previous_hash, event_hash, created_at FROM audit_events ORDER BY id"):
            if phash != prev:
                return {"verified": False, "checked": n, "first_bad_event": rid, "reason": "chain broken"}
            try:
                ts = datetime.fromisoformat(created).isoformat() if created else ""
            except ValueError:
                ts = created or ""
            body = {"actor": actor, "action": action, "details": details, "timestamp": ts, "previous_hash": phash}
            digest = hmac.new(secret_key.encode(), json.dumps(body, sort_keys=True, ensure_ascii=False).encode(), hashlib.sha256).hexdigest()
            if not hmac.compare_digest(digest, ehash or ""):
                return {"verified": False, "checked": n, "first_bad_event": rid,
                        "reason": "event hash mismatch (altered content, or this server's SECRET_KEY differs from the one that wrote the backup)"}
            prev, n = ehash, n + 1
        return {"verified": True, "checked": n, "reason": "OK"}
    except sqlite3.Error as exc:
        return {"verified": False, "checked": 0, "reason": f"audit table unreadable: {exc}"}
    finally:
        conn.close()


def _create_backup(ctx, db: Session, actor: str, passphrase: str) -> dict:
    src = _sqlite_path()
    ev_root = A.EVIDENCE_DIR.resolve()
    files = []
    for p in sorted(ev_root.rglob("*")):
        if p.is_symlink() or not p.is_file():
            continue
        files.append(p)
    total = src.stat().st_size + sum(p.stat().st_size for p in files)
    cap = backup_max_bytes()
    if total > cap:
        raise HTTPException(413, f"Backup would be {total // (1024 * 1024)} MB which exceeds the BACKUP_MAX_MB limit of {cap // (1024 * 1024)} MB")
    chain = ctx.svc.verify_audit_chain(db)
    bdir = _backup_dir()
    tmp = Path(tempfile.mkdtemp(prefix=".snap-", dir=bdir))
    try:
        snap = tmp / "snapshot.sqlite"
        s_conn = sqlite3.connect(str(src)); d_conn = sqlite3.connect(str(snap))
        try:
            s_conn.backup(d_conn)   # consistent online snapshot
        finally:
            d_conn.close(); s_conn.close()
        entries, buf = [], io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
            data = snap.read_bytes()
            zf.writestr("database/dcn.sqlite", data)
            entries.append({"path": "database/dcn.sqlite", "sha256": hashlib.sha256(data).hexdigest(), "bytes": len(data)})
            for p in files:
                rel = "evidence/" + p.relative_to(ev_root).as_posix()
                data = p.read_bytes()
                zf.writestr(rel, data)
                entries.append({"path": rel, "sha256": hashlib.sha256(data).hexdigest(), "bytes": len(data)})
            manifest = {"format": 1, "app_version": _app_version(), "created_at": _now().isoformat() + "Z", "created_by": actor,
                        "audit_chain_at_snapshot": {"verified": chain.get("verified"), "checked": chain.get("checked")},
                        "evidence_files": len(files), "files": entries,
                        "note": "Restore offline with scripts/restore_backup.py; the server has no restore endpoint by design."}
            zf.writestr("manifest.json", json.dumps(manifest, indent=1))
        blob = _encrypt_blob(buf.getvalue(), passphrase)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    name = f"dcn-backup-{_now().strftime('%Y%m%dT%H%M%SZ')}-{secrets.token_hex(3)}.dcnbak"
    final = bdir / name
    part = bdir / (name + ".part")
    part.write_bytes(blob)
    try:
        os.chmod(part, 0o600)
    except OSError:
        pass
    os.replace(part, final)
    return {"name": name, "bytes": len(blob), "files": len(entries), "evidence_files": len(files),
            "sha256": hashlib.sha256(blob).hexdigest(), "created_at": manifest["created_at"], "app_version": manifest["app_version"]}


def _backup_path(name: str) -> Path:
    if not _BACKUP_NAME_RE.fullmatch(name or ""):
        raise HTTPException(400, "Invalid backup name")
    bdir = _backup_dir().resolve()
    p = (bdir / name).resolve()
    if p.parent != bdir or p.is_symlink():
        raise HTTPException(400, "Invalid backup name")
    if not p.is_file():
        raise HTTPException(404, "Backup not found")
    return p


def _verify_backup(blob: bytes, passphrase: str) -> dict:
    from backend.app.core.config import settings
    plain = _decrypt_blob(blob, passphrase)
    result = {"ok": False, "checks": [], "files_checked": 0, "mismatches": [], "missing": [], "unexpected": []}
    try:
        zf = zipfile.ZipFile(io.BytesIO(plain))
    except zipfile.BadZipFile:
        raise ValueError("Backup archive is corrupt")
    with zf:
        names = zf.namelist()
        if any(not _safe_zip_name(n) for n in names):
            raise ValueError("Backup archive contains unsafe paths")
        if "manifest.json" not in names:
            raise ValueError("Backup has no manifest")
        manifest = json.loads(zf.read("manifest.json"))
        listed = {}
        for f in manifest.get("files", []):
            if not _safe_zip_name(f.get("path", "")):
                raise ValueError("Manifest contains unsafe paths")
            listed[f["path"]] = f
        for path, f in listed.items():
            if path not in names:
                result["missing"].append(path); continue
            data = zf.read(path)
            result["files_checked"] += 1
            if hashlib.sha256(data).hexdigest() != f.get("sha256") or len(data) != f.get("bytes"):
                result["mismatches"].append(path)
        result["unexpected"] = [n for n in names if n != "manifest.json" and n not in listed]
        result["checks"].append({"name": "all file hashes match the manifest", "ok": not result["mismatches"] and not result["missing"] and not result["unexpected"]})
        db_ok, integrity, chain, tables = False, "not run", {"verified": False, "reason": "database missing"}, 0
        if "database/dcn.sqlite" in names:
            tmpd = tempfile.mkdtemp(prefix="dcn-verify-")
            try:
                dbf = Path(tmpd) / "v.sqlite"
                dbf.write_bytes(zf.read("database/dcn.sqlite"))
                conn = sqlite3.connect(f"file:{dbf}?mode=ro", uri=True)
                try:
                    integrity = conn.execute("PRAGMA integrity_check").fetchone()[0]
                    tables = conn.execute("SELECT COUNT(*) FROM sqlite_master WHERE type='table'").fetchone()[0]
                    db_ok = integrity == "ok"
                except sqlite3.Error as exc:
                    integrity = f"database does not open: {exc}"
                finally:
                    conn.close()
                if db_ok:
                    chain = _chain_of_sqlite(str(dbf), settings.secret_key)
            finally:
                shutil.rmtree(tmpd, ignore_errors=True)
        result["checks"].append({"name": "database opens and passes SQLite integrity_check", "ok": db_ok, "detail": integrity, "tables": tables})
        result["checks"].append({"name": "audit chain verifies (with this server's SECRET_KEY)", "ok": bool(chain.get("verified")), "detail": chain})
        result["manifest"] = {k: manifest.get(k) for k in ("format", "app_version", "created_at", "created_by", "evidence_files", "audit_chain_at_snapshot")}
        result["ok"] = all(c["ok"] for c in result["checks"])
    return result


# ---------------------------------------------------------------------------------------------- document rendering
# A tiny block model rendered to DOCX (python-docx) or PDF (PyMuPDF, as in AppService.report_pdf):
#   ("title", text) ("h", text) ("p", text) ("note", text) ("kv", [(k, v), ...]) ("table", header, rows) ("sig", [labels])
def _render_docx(blocks: list, landscape: bool = False) -> bytes:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.shared import Cm, Pt, RGBColor
    doc = Document()
    sec = doc.sections[0]
    if landscape:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(1.8))
    doc.styles["Normal"].font.size = Pt(10)
    for b in blocks:
        kind = b[0]
        if kind == "title":
            doc.add_heading(b[1], 1)
        elif kind == "h":
            doc.add_heading(b[1], 2)
        elif kind == "p":
            doc.add_paragraph(b[1])
        elif kind == "note":
            para = doc.add_paragraph()
            run = para.add_run(b[1]); run.bold = True; run.font.color.rgb = RGBColor(0x9A, 0x34, 0x12)
        elif kind == "kv":
            t = doc.add_table(rows=0, cols=2); t.style = "Table Grid"
            for k, v in b[1]:
                cells = t.add_row().cells
                cells[0].text = str(k); cells[1].text = "" if v is None else str(v)
                for r in cells[0].paragraphs[0].runs:
                    r.bold = True
            doc.add_paragraph()
        elif kind == "table":
            header, rows = b[1], b[2]
            t = doc.add_table(rows=1, cols=len(header)); t.style = "Table Grid"
            for i, hd in enumerate(header):
                t.rows[0].cells[i].text = str(hd)
                for r in t.rows[0].cells[i].paragraphs[0].runs:
                    r.bold = True
            for row in rows:
                cells = t.add_row().cells
                for i, v in enumerate(row):
                    cells[i].text = "" if v is None else str(v)
                    for para in cells[i].paragraphs:
                        for r in para.runs:
                            r.font.size = Pt(8)
            doc.add_paragraph()
        elif kind == "sig":
            t = doc.add_table(rows=0, cols=2); t.style = "Table Grid"
            for lab in b[1]:
                cells = t.add_row().cells
                cells[0].text = lab; cells[1].text = ""
            doc.add_paragraph()
    buf = io.BytesIO(); doc.save(buf)
    return buf.getvalue()


def _render_pdf(blocks: list) -> bytes:
    try:
        import pymupdf as fitz
    except ImportError:  # older PyMuPDF
        import fitz  # type: ignore
    W, H, M = 595, 842, 46
    pdf = fitz.open()
    state = {"page": None, "y": M}

    def newpage():
        state["page"] = pdf.new_page(width=W, height=H); state["y"] = M

    def wrap(text, font, size, width):
        lines = []
        for para in str(text).split("\n"):
            words, cur = para.split(" "), ""
            for w in words:
                while fitz.get_text_length(w, fontname=font, fontsize=size) > width and len(w) > 1:  # very long tokens (hashes)
                    cut = len(w)
                    while cut > 1 and fitz.get_text_length(w[:cut], fontname=font, fontsize=size) > width:
                        cut -= 1
                    if cur:
                        lines.append(cur); cur = ""
                    lines.append(w[:cut]); w = w[cut:]
                trial = (cur + " " + w).strip()
                if fitz.get_text_length(trial, fontname=font, fontsize=size) <= width:
                    cur = trial
                else:
                    lines.append(cur); cur = w
            lines.append(cur)
        return lines

    def put(text, font="helv", size=9.5, indent=0, gap=3, color=(0.08, 0.1, 0.14)):
        for line in wrap(text, font, size, W - 2 * M - indent):
            if state["page"] is None or state["y"] > H - M:
                newpage()
            state["y"] += size + 1.5
            state["page"].insert_text((M + indent, state["y"]), line, fontname=font, fontsize=size, color=color)
        state["y"] += gap

    newpage()
    for b in blocks:
        kind = b[0]
        if kind == "title":
            put(b[1], "hebo", 14, gap=8)
        elif kind == "h":
            put(b[1], "hebo", 11, gap=4)
        elif kind == "p":
            put(b[1])
        elif kind == "note":
            put(b[1], "hebo", 9, color=(0.6, 0.2, 0.07), gap=6)
        elif kind == "kv":
            for k, v in b[1]:
                put(f"{k}: {'' if v is None else v}", gap=1)
            state["y"] += 4
        elif kind == "table":
            header, rows = b[1], b[2]
            if not rows:
                put("(none)")
            for i, row in enumerate(rows, 1):
                put(f"{i}. {row[0]}" + (f"   {row[1]}" if len(row) > 1 else ""), "hebo", 9, gap=1)
                for lab, v in zip(header[2:], row[2:]):
                    put(f"{lab}: {'' if v is None else v}", "cour" if lab.lower().startswith("sha") else "helv", 8, indent=14, gap=0.5)
                state["y"] += 4
        elif kind == "sig":
            for lab in b[1]:
                put(f"{lab}: " + "_" * 48, gap=7)
    out = pdf.tobytes(); pdf.close()
    return out


def _render(blocks: list, fmt: str, landscape: bool = False) -> tuple[bytes, str]:
    if fmt == "pdf":
        return _render_pdf(blocks), "application/pdf"
    return _render_docx(blocks, landscape), DOCX_MT


def _redacted_docx(cn: str, out: dict, log_: dict) -> bytes:
    blocks = [("title", f"Redacted data export - {cn}"),
              ("note", "REDACTED EXPORT: personal identifiers have been masked, pseudonymised or removed according to the log below. "
                       "Pseudonyms are consistent inside this export only. Review before sharing; regex-based detection is a heuristic."),
              ("h", "Entities"),
              ("table", ["ID", "Name", "Type", "Confidence"], [[e["external_id"], e["name"], e["type"], f'{float(e["confidence"] or 0):.0%}'] for e in out["entities"][:500]]),
              ("h", "Relationships"),
              ("table", ["Source", "Relation", "Target", "State", "Time"], [[r["source"], r["relation"], r["target"], r["verification_state"], r["event_time"]] for r in out["relationships"][:800]]),
              ("h", "Events"),
              ("table", ["Event", "Type", "Entity", "Related", "Time", "Amount"], [[e["event_id"], e["event_type"], e["entity_id"], e["related_entity_id"], e["event_time"], e["amount"]] for e in out["events"][:800]]),
              ("h", "Documents (structured fields)"),
              ("table", ["Document", "Filename", "Status", "FIR", "Complainant", "Place"],
               [[d["document_id"], d["filename"], d["status"], d["fir_number"], d["structured"].get("complainant"), d["structured"].get("place_of_occurrence")] for d in out["documents"][:300]]),
              ("h", "Redaction log"),
              ("kv", [("Mode", log_["mode"]), ("Categories", ", ".join(log_["categories"])), ("Fields changed", log_["fields_changed"])]
                     + [(f"Replacements: {c}", f'{n} ({log_["distinct_values_by_category"][c]} distinct)') for c, n in log_["replacements_by_category"].items()]),
              ("p", "Limitations: " + " ".join(log_["limitations"]))]
    return _render_docx(blocks, landscape=True)


# ---------------------------------------------------------------------------------------------- legal packs
def _custody_events(db: Session, cn: str, ev_ids: set[str]) -> dict[str, list[dict]]:
    out: dict[str, list[dict]] = defaultdict(list)
    q = db.query(AuditEvent).filter(AuditEvent.action.in_(["evidence.uploaded", "evidence.verified", "evidence.downloaded", "evidence.duplicate_upload"]),
                                    A._like_case(AuditEvent.details, cn)).order_by(AuditEvent.id.asc())
    for a in q.all():
        d = _loads(a.details, {})
        if d.get("evidence_id") in ev_ids and d.get("case_number") == cn:
            out[d["evidence_id"]].append({"time": _iso(a.created_at), "action": a.action, "actor": a.actor,
                                          "detail": d.get("status") or d.get("stage") or ""})
    return out


def _case_header(db: Session, cn: str, user: dict, what: str) -> list:
    c = db.query(Case).filter_by(case_number=cn).first()
    return [("kv", [("Case number", cn), ("Case title", c.title if c else ""), ("Generated (UTC)", _now().strftime("%Y-%m-%d %H:%M:%S")),
                    ("Generated by", user["sub"]), ("Purpose", what)])]


def _bsa63_blocks(db: Session, user: dict, cn: str, ev: Evidence) -> list:
    p = None
    try:
        p = A._abs_path(ev.stored_path)
        current = _sha256_file(p) if p.is_file() else None
    except OSError:
        current = None
    if current is None:
        recheck = "stored file not available at generation time - hash could not be re-computed"
    else:
        recheck = f"re-computed at generation: {current} ({'MATCHES' if current == ev.sha256 else 'DOES NOT MATCH'} the registered hash)"
    size = None
    try:
        size = p.stat().st_size if p and p.is_file() else None
    except OSError:
        pass
    blocks = [
        ("title", "CERTIFICATE UNDER SECTION 63 (4) OF THE BHARATIYA SAKSHYA ADHINIYAM, 2023"),
        ("note", "TEMPLATE AID - to be reviewed by the legal officer; format per current schedule of the Act. "
                 "Pre-filled fields come from this system's evidence register. The certificate has no legal effect until the "
                 "person(s) named in Part A and Part B complete, verify and sign it."),
        *_case_header(db, cn, user, "Certificate for electronic record produced as evidence"),
        ("h", "PART A - to be completed and signed by the person in charge of / producing the electronic record"),
        ("p", "I, ______________________ (name), son/daughter/spouse of ______________________, aged ____ years, "
              "designation ______________________, resident of / posted at ______________________, do solemnly affirm and state as follows:"),
        ("p", "1. I produce the electronic record described below, obtained from the following device / digital source (tick as applicable):"),
        ("p", "[ ] Computer / storage media    [ ] Mobile phone    [ ] DVR / CCTV    [ ] Server / cloud account    [ ] Other: ______________"),
        ("kv", [("Make and model", ""), ("Serial number / IMEI / UDI", ""), ("Owner / lawful custodian", ""), ("Date and place of extraction", "")]),
        ("p", "2. Details of the electronic record (pre-filled from the evidence register):"),
        ("kv", [("Evidence ID", ev.evidence_id), ("File name", ev.filename), ("Media type", ev.media_type), ("Size on disk (bytes)", size if size is not None else "not available"),
                ("Source / description", ev.source), ("Registered by", ev.uploaded_by), ("Registered on (UTC)", ev.created_at.strftime("%Y-%m-%d %H:%M:%S") if ev.created_at else ""),
                ("Linked document", ev.document_id or "-"), ("Integrity status in register", ev.integrity_status)]),
        ("p", "3. Hash value of the electronic record:"),
        ("kv", [("Hash algorithm", "SHA-256"), ("Hash value (registered)", ev.sha256), ("Integrity re-check", recheck)]),
        ("p", "4. The statements below relate to the conditions in Section 63(2) (legal officer to confirm the exact wording against the schedule):"),
        ("p", "[ ] The device/computer was used regularly to store or process information for the purposes of the activities regularly carried on by the person having lawful control over it during the relevant period.\n"
              "[ ] During the relevant period, information of the kind contained in the record was regularly fed into the device in the ordinary course of those activities.\n"
              "[ ] Throughout the material part of the relevant period the device was operating properly (or, if not, the respect in which it was not operating properly did not affect the record or its accuracy).\n"
              "[ ] The information contained in the electronic record reproduces or is derived from the information fed into the device in the ordinary course of those activities."),
        ("sig", ["Signature of the person", "Name", "Designation", "Date", "Place"]),
        ("h", "PART B - to be completed and signed by the Expert (where applicable)"),
        ("p", "I, ______________________ (name), qualification(s) ______________________, of ______________________ (organisation), "
              "have examined the electronic record identified in Part A using ______________________ (tool and version) and state:"),
        ("kv", [("Hash values computed by the expert", "SHA-256: ______________________   Other (algorithm): ______________________"),
                ("Match with Part A hash", "[ ] Yes    [ ] No"), ("Observations / limitations", "")]),
        ("sig", ["Signature of the expert", "Name", "Designation / registration no.", "Date", "Place", "Seal"]),
        ("p", f"Generated by DARK CRIMENET for {user['sub']}. AI outputs are investigative aids; this certificate is a document template, not legal advice."),
    ]
    return blocks


def _evidence_index_blocks(db: Session, user: dict, cn: str) -> list:
    evs = db.query(Evidence).filter_by(case_number=cn).order_by(Evidence.created_at, Evidence.id).all()
    custody = _custody_events(db, cn, {e.evidence_id for e in evs})
    rows = []
    for i, e in enumerate(evs, 1):
        try:
            present = "yes" if A._abs_path(e.stored_path).is_file() else "NO"
        except OSError:
            present = "unknown"
        cust = custody.get(e.evidence_id, [])
        rows.append([i, e.evidence_id, e.filename, e.sha256, e.uploaded_by, e.created_at.strftime("%Y-%m-%d %H:%M:%S") if e.created_at else "",
                     e.integrity_status, present, f"{sum(1 for c in cust if c['action'] == 'evidence.downloaded')} download(s), {sum(1 for c in cust if c['action'] == 'evidence.verified')} verification(s)"])
    blocks = [("title", "ANNEXURE - EVIDENCE INDEX AND CHAIN OF CUSTODY"),
              ("note", "TEMPLATE AID - to be reviewed by the legal officer. Generated from the evidence register and the tamper-evident audit log."),
              *_case_header(db, cn, user, "Evidence index / chain-of-custody annexure"),
              ("h", f"Evidence items ({len(evs)})"),
              ("table", ["S.No", "Evidence ID", "File name", "SHA-256", "Uploaded by", "Uploaded (UTC)", "Integrity status", "File present", "Custody summary"], rows)]
    blocks.append(("h", "Custody events (from the audit log)"))
    blocks.append(("p", "Only events recorded by this system are listed. The absence of a 'downloaded' event means none was recorded here, not that none occurred elsewhere."))
    for e in evs:
        ce = custody.get(e.evidence_id, [])
        blocks.append(("p", f"{e.evidence_id} - {e.filename}"))
        blocks.append(("table", ["No.", "Time (UTC)", "Action", "Actor", "Detail"],
                       [[i, c["time"], c["action"], c["actor"], c["detail"]] for i, c in enumerate(ce, 1)]))
    blocks.append(("sig", ["Prepared by", "Verified by (legal officer)", "Date"]))
    return blocks


def _document_index_blocks(db: Session, user: dict, cn: str) -> list:
    D = DocumentRecord
    docs = db.query(D.document_id, D.filename, D.document_type, D.status, D.sha256, D.review_actor, D.reviewed_at, D.created_at, D.structured_json) \
             .filter_by(case_number=cn).order_by(D.id).all()
    rows = []
    for i, d in enumerate(docs, 1):
        st = _loads(d.structured_json, {})
        rows.append([i, d.document_id, st.get("fir_number") or "", st.get("fir_date") or "", st.get("police_station") or "", d.filename,
                     d.document_type, d.status, d.review_actor or "", d.reviewed_at.strftime("%Y-%m-%d %H:%M") if d.reviewed_at else "", d.sha256])
    return [("title", "ANNEXURE - INDEX OF FIRs / DOCUMENTS"),
            ("note", "TEMPLATE AID - to be reviewed by the legal officer. Status shows the human review state of each document."),
            *_case_header(db, cn, user, "Document / FIR index annexure"),
            ("h", f"Documents ({len(rows)})"),
            ("table", ["S.No", "Document ID", "FIR no.", "FIR date", "Police station", "File name", "Type", "Review status", "Reviewed by", "Reviewed (UTC)", "SHA-256"], rows),
            ("sig", ["Prepared by", "Verified by (legal officer)", "Date"])]


# ---------------------------------------------------------------------------------------------- login history
_AUTH_ACTIONS = ["auth.login", "auth.login_failed", "auth.password_changed", "auth.username_changed", "governance.force_logout"]


def _login_history(db: Session, username: str | None, days: int) -> dict:
    since = _now() - timedelta(days=days)
    rows = (db.query(AuditEvent).filter(AuditEvent.created_at >= since,
                                        sa.or_(AuditEvent.action.in_(_AUTH_ACTIONS), AuditEvent.action.like("admin.%")))
            .order_by(AuditEvent.id.asc()).limit(50000).all())
    users: dict[str, dict] = {}

    def u(name):
        return users.setdefault(name, {"username": name, "logins": 0, "failed_logins": 0, "last_login": None, "last_failure": None,
                                       "current_failed_streak": 0, "max_failed_streak": 0, "ips": set(), "password_changes": 0,
                                       "forced_logouts": 0})
    events = []
    summary = Counter()
    for a in rows:
        d = _loads(a.details, {}) if a.details and a.details.startswith("{") else {}
        subject = d.get("username") or d.get("old_username") or a.actor
        names = {a.actor, d.get("username"), d.get("old_username"), d.get("new_username")} - {None}
        if username and username not in names:
            continue
        ip = d.get("ip")
        if a.action == "auth.login":
            s = u(a.actor); s["logins"] += 1; s["last_login"] = _iso(a.created_at); s["current_failed_streak"] = 0; summary["logins"] += 1
        elif a.action == "auth.login_failed":
            s = u(subject); s["failed_logins"] += 1; s["last_failure"] = _iso(a.created_at); s["current_failed_streak"] += 1
            s["max_failed_streak"] = max(s["max_failed_streak"], s["current_failed_streak"]); summary["failed_logins"] += 1
            if ip:
                s["ips"].add(ip)
        elif a.action == "auth.password_changed":
            u(a.actor)["password_changes"] += 1; summary["password_changes"] += 1
        elif a.action == "auth.username_changed":
            summary["username_changes"] += 1
        elif a.action == "governance.force_logout":
            u(d.get("target") or subject)["forced_logouts"] += 1; summary["forced_logouts"] += 1
        elif a.action.startswith("admin."):
            summary["admin_actions"] += 1
        events.append({"id": a.id, "time": _iso(a.created_at), "actor": a.actor, "action": a.action, "ip": ip,
                       "target": d.get("username") or d.get("target") or d.get("new_username") or "",
                       "detail": {k: v for k, v in d.items() if k not in {"ip", "username"}}})
    for s in users.values():
        s["ips"] = sorted(s["ips"])
    return {"window_days": days, "username": username, "summary": dict(summary),
            "users": sorted(users.values(), key=lambda s: (-s["failed_logins"], s["username"])),
            "events": events[::-1][:300],
            "note": "Derived from the audit log. IP addresses are recorded on failed sign-ins only; successful sign-in events carry no address. "
                    "Failed sign-ins are logged for existing accounts only."}


# ---------------------------------------------------------------------------------------------- import connectors
FIR_FIELDS = [
    ("fir_number", "FIR / crime number", True, ["fir no", "fir number", "firno", "fir_no", "fir no.", "crime no", "crime number", "cr no", "crime no.", "fir/crime no", "fir crime number", "fir #"]),
    ("fir_date", "FIR date", False, ["fir date", "date of fir", "date of registration", "registration date", "reg date", "date registered", "fir reg date"]),
    ("police_station", "Police station", False, ["police station", "ps", "p.s.", "ps name", "police station name", "station", "thana", "police stn"]),
    ("district", "District", False, ["district", "dist", "district name"]),
    ("year", "Year", False, ["year", "fir year", "reg year"]),
    ("complainant", "Complainant name", False, ["complainant name", "complainant", "informant name", "informant", "name of complainant", "victim name", "name of informant"]),
    ("father_name", "Father's / husband's name", False, ["father name", "father's name", "complainant father name", "father/husband name", "father or husband name", "fathers name"]),
    ("complainant_address", "Complainant address", False, ["complainant address", "address of complainant", "informant address", "address", "residential address"]),
    ("place_of_occurrence", "Place of occurrence", False, ["place of occurrence", "place of offence", "crime place", "scene of crime", "location of occurrence", "occurrence place", "place of crime"]),
    ("occurrence_date", "Date of occurrence", False, ["date of occurrence", "occurrence date", "date of offence", "date of incident", "incident date", "doo"]),
    ("sections", "Sections / acts", False, ["sections", "acts & sections", "acts and sections", "sections of law", "ipc sections", "bns sections", "act/section", "law and sections", "section", "acts sections", "act section"]),
    ("accused", "Accused name / description", False, ["accused name", "accused", "suspect", "accused details", "accused name(s)", "name of accused", "accused description"]),
    ("phones", "Mobile / phone", False, ["mobile no", "mobile number", "phone", "contact no", "complainant mobile", "mobile", "phone no", "contact number", "mobile no."]),
    ("vehicles", "Vehicle number", False, ["vehicle no", "vehicle number", "registration no", "veh no", "vehicle reg no", "vehicle registration"]),
    ("accounts", "Bank account", False, ["account no", "account number", "bank account", "a/c no", "bank account no"]),
    ("emails", "E-mail", False, ["email", "e-mail", "email id", "email address"]),
    ("description", "Brief facts / description", False, ["brief facts", "incident description", "description", "gist", "details of incident", "brief description", "gist of case", "facts"]),
    ("investigating_officer", "Investigating officer", False, ["io", "io name", "investigating officer", "investigating officer name", "name of io"]),
]
FIELD_KEYS = {f[0] for f in FIR_FIELDS}
_ALIAS_NORM: dict[str, str] = {}
for _key, _label, _req, _aliases in FIR_FIELDS:
    for _a in _aliases + [_label]:
        _ALIAS_NORM.setdefault(re.sub(r"[^a-z0-9]", "", _a.lower()), _key)
_MAX_IMPORT_ROWS = 5000


def _n(s) -> str:
    return re.sub(r"[^a-z0-9]", "", str(s).lower())


def suggest_mapping(columns: list[str]) -> dict:
    """Map source columns to internal FIR fields using alias dictionaries (exact, contains, then fuzzy). Heuristic."""
    import difflib
    cand: list[tuple[float, str, str, str]] = []
    aliases = list(_ALIAS_NORM)
    for col in columns:
        nc = _n(col)
        if not nc:
            continue
        if nc in _ALIAS_NORM:
            cand.append((1.0, col, _ALIAS_NORM[nc], "alias"))
            continue
        contains = [a for a in aliases if len(a) >= 6 and a in nc]
        if contains:
            best = max(contains, key=len)
            cand.append((0.8, col, _ALIAS_NORM[best], "contains"))
            continue
        close = difflib.get_close_matches(nc, aliases, n=1, cutoff=0.84)
        if close:
            cand.append((round(difflib.SequenceMatcher(None, nc, close[0]).ratio() * 0.9, 2), col, _ALIAS_NORM[close[0]], "fuzzy"))
    cand.sort(key=lambda x: (-x[0], x[1]))
    used_fields, used_cols, out = set(), set(), {}
    for conf, col, fld, how in cand:
        if fld in used_fields or col in used_cols:
            continue
        used_fields.add(fld); used_cols.add(col)
        out[col] = {"field": fld, "confidence": conf, "method": how}
    return out


def _read_table(data: bytes, filename: str) -> pd.DataFrame:
    suffix = A.validate_upload(filename, data)
    if suffix not in {".csv", ".xlsx"}:
        raise ValueError("Import supports CSV and XLSX files only")
    try:
        if suffix == ".csv":
            try:
                text = data.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = data.decode("cp1252", errors="replace")
            df = pd.read_csv(io.StringIO(text), dtype=str, keep_default_na=False, nrows=_MAX_IMPORT_ROWS + 1)
        else:
            df = pd.read_excel(io.BytesIO(data), dtype=str, nrows=_MAX_IMPORT_ROWS + 1).fillna("")
    except Exception as exc:  # noqa: BLE001
        raise ValueError(f"Could not read the table: {exc.__class__.__name__}") from exc
    if len(df) > _MAX_IMPORT_ROWS:
        raise ValueError(f"Table has more than {_MAX_IMPORT_ROWS} rows; split it into smaller files")
    cols, seen = [], Counter()
    for c in df.columns:
        name = str(c).strip() or "column"
        seen[name] += 1
        cols.append(name if seen[name] == 1 else f"{name} ({seen[name]})")
    df.columns = cols
    df = df[[c for c in df.columns if not (c.lower().startswith("unnamed") and not df[c].astype(str).str.strip().any())]]
    return df.fillna("").astype(str)


_UPLOADS: dict[str, dict] = {}
_UPLOAD_LOCK = threading.Lock()


def _stash_upload(owner: str, filename: str, data: bytes) -> str:
    with _UPLOAD_LOCK:
        now = time.time()
        for k in [k for k, v in _UPLOADS.items() if now - v["t"] > 1800]:
            _UPLOADS.pop(k, None)
        mine = sorted((k for k, v in _UPLOADS.items() if v["owner"] == owner), key=lambda k: _UPLOADS[k]["t"])
        for k in mine[:-4]:
            _UPLOADS.pop(k, None)
        while len(_UPLOADS) >= 20:
            _UPLOADS.pop(min(_UPLOADS, key=lambda k: _UPLOADS[k]["t"]), None)
        token = secrets.token_urlsafe(18)
        _UPLOADS[token] = {"owner": owner, "filename": filename, "data": data, "t": now}
        return token


def _get_upload(owner: str, token: str) -> dict:
    with _UPLOAD_LOCK:
        v = _UPLOADS.get(token)
    if not v or v["owner"] != owner or time.time() - v["t"] > 1800:
        raise HTTPException(404, "The uploaded file is no longer available (it is kept for 30 minutes). Upload it again.")
    return v


def _clean_cell(v, limit=500) -> str:
    return re.sub(r"\s+", " ", str(v or "")).strip()[:limit]


def _row_to_fir(vals: dict[str, str]) -> tuple[str, dict]:
    g = lambda k: vals.get(k, "")  # noqa: E731
    lines = ["FIRST INFORMATION REPORT", "(structured-register import - awaiting human review)"]
    for label, key in (("District", "district"), ("Police Station", "police_station"), ("Year", "year"), ("FIR No", "fir_number"),
                       ("Date of FIR", "fir_date"), ("Acts & Sections", "sections")):
        if g(key):
            lines.append(f"{label}: {g(key)}")
    if g("place_of_occurrence"):
        lines.append(f"Place of Occurrence: {g('place_of_occurrence')}")
    if g("occurrence_date"):
        lines.append(f"Date of Occurrence: {g('occurrence_date')}")
    if g("complainant") or g("father_name") or g("complainant_address"):
        lines.append("Complainant / Informant:")
        for label, key in (("Name", "complainant"), ("Father's Name", "father_name"), ("Address", "complainant_address")):
            if g(key):
                lines.append(f"{label}: {g(key)}")
    if g("accused"):
        lines.append(f"Accused: {g('accused')}")
    for label, key in (("Phone", "phones"), ("Vehicle used", "vehicles"), ("Account", "accounts"), ("Email", "emails")):
        if g(key):
            lines.append(f"{label}: {g(key)}")
    if g("description"):
        lines.append(f"Short Description of Incident: {g('description')}")
    if g("investigating_officer"):
        lines.append(f"Investigating Officer: {g('investigating_officer')}")
    overlay = {k: v for k, v in {
        "fir_number": g("fir_number"), "fir_date": g("fir_date"), "police_station": g("police_station"), "district": g("district"),
        "year": g("year") or (re.search(r"(\d{4})$", g("fir_number") or "") or [None, ""])[1],
        "complainant": g("complainant"), "father_name": g("father_name"), "complainant_address": g("complainant_address"),
        "place_of_occurrence": g("place_of_occurrence"), "occurrence_date": g("occurrence_date"),
        "accused_description": g("accused"), "investigating_officer": g("investigating_officer"), "incident_description": g("description"),
    }.items() if v}
    if g("sections"):
        overlay["sections"] = list(dict.fromkeys(re.findall(r"\d{1,5}(?:\([0-9A-Za-z]+\))?", g("sections"))))
    return "\n".join(lines) + "\n", overlay


def _commit_import(ctx, db: Session, user: dict, x: ImportCommitIn) -> dict:
    cn = ctx.scope(db, user, x.case_number)
    up = _get_upload(user["sub"], x.upload_token)
    unknown = [f for f in x.mapping if f not in FIELD_KEYS]
    if unknown:
        raise HTTPException(400, f"Unknown target field(s): {', '.join(unknown)}")
    if not x.mapping.get("fir_number"):
        raise HTTPException(400, "Map a column to 'FIR / crime number' before importing")
    try:
        df = _read_table(up["data"], up["filename"])
    except ValueError as exc:
        raise HTTPException(400, str(exc))
    missing = [c for c in x.mapping.values() if c not in df.columns]
    if missing:
        raise HTTPException(400, f"Mapped column(s) not found in the file: {', '.join(missing)}")
    stem = re.sub(r"[^A-Za-z0-9_-]+", "-", Path(up["filename"]).stem)[:40] or "import"
    imported, duplicates, skipped, ids = 0, 0, [], []
    for idx, row in df.iterrows():
        n = int(idx) + 2  # spreadsheet-style row number (header is row 1)
        vals = {f: _clean_cell(row[c]) for f, c in x.mapping.items() if c}
        if not any(vals.values()):
            continue
        if not vals.get("fir_number"):
            skipped.append({"row": n, "reason": "no FIR number in this row"}); continue
        text, overlay = _row_to_fir(vals)
        payload = text.encode("utf-8")
        if db.query(DocumentRecord.id).filter_by(sha256=hashlib.sha256(payload).hexdigest(), case_number=cn).first():
            duplicates += 1; continue
        try:
            res = ctx.svc.extract_document(db, payload, f"import-{stem}-row{n}.txt", case_number=cn, media_type="text/plain",
                                           actor=user["sub"], user=user)
        except ValueError as exc:
            skipped.append({"row": n, "reason": str(exc)[:200]}); continue
        doc = db.query(DocumentRecord).filter_by(document_id=res["document_id"]).first()
        st = _loads(doc.structured_json, {})
        st.update(overlay)
        st["import"] = {"source_file": up["filename"], "row": n, "imported_by": user["sub"], "profile": x.profile_name or None,
                        "mapping": x.mapping, "note": "Values copied from a structured register; a human must review before verification."}
        doc.structured_json = A.j(st)
        q = _loads(doc.quality_json, {}); q["imported_from_table"] = True; doc.quality_json = A.j(q)
        db.commit()
        imported += 1
        ids.append(doc.document_id)
        ctx.hooks.emit("document_extracted", db=db, user=user, result=res, case_number=cn)
    ctx.audit(db, user["sub"], "import.committed", {"case_number": cn, "file": up["filename"], "imported": imported, "duplicates": duplicates,
                                                    "skipped": len(skipped), "profile": x.profile_name or ""})
    return {"case_number": cn, "imported": imported, "duplicates": duplicates, "skipped": skipped[:50], "skipped_count": len(skipped),
            "document_ids": ids[:200], "status": "Pending Review",
            "note": "Every row is a FIR document in 'Pending Review'. Nothing is verified automatically; open Documents to review each one."}


# ---------------------------------------------------------------------------------------------- folder inbox
class _InboxState:
    lock = threading.Lock()
    recent: deque = deque(maxlen=40)
    totals: Counter = Counter()
    last_scan_at: str | None = None
    last_result: dict | None = None
    thread: threading.Thread | None = None
    enabled = False


INBOX_ACTOR = {"sub": "inbox-watcher", "role": "investigator", "uid": 0}


def _inbox_dir() -> Path:
    raw = os.environ.get("INBOX_DIR", "").strip()
    return Path(raw).resolve() if raw else A.EVIDENCE_DIR.resolve().parent / "inbox"


def _inbox_move(root: Path, src: Path, bucket: str, cn: str, reason: str | None = None) -> str:
    dest_dir = root / bucket / cn
    dest_dir.mkdir(parents=True, exist_ok=True)
    stamp = _now().strftime("%Y%m%dT%H%M%S")
    safe = re.sub(r"[^A-Za-z0-9._-]+", "_", src.name)[:120] or "file"
    dest = dest_dir / f"{stamp}-{safe}"
    if not dest.resolve().is_relative_to(root.resolve()):
        raise OSError("refusing to move outside the inbox")
    shutil.move(str(src), str(dest))
    if reason:
        Path(str(dest) + ".reason.txt").write_text(reason + "\n", encoding="utf-8")
    return str(dest.relative_to(root))


def _inbox_scan(ctx, min_age: float = 3.0) -> dict:
    import mimetypes
    import stat as statmod
    with _InboxState.lock:
        root = _inbox_dir()
        root.mkdir(parents=True, exist_ok=True)
        result = {"processed": [], "failed": [], "skipped_unknown_case": [], "pending_recent": 0}
        max_bytes = ctx.settings.max_upload_bytes
        db = ctx.SessionLocal()
        try:
            for case_dir in sorted(root.iterdir()):
                name = case_dir.name
                if name.startswith(("_", ".")) or case_dir.is_symlink() or not case_dir.is_dir():
                    continue
                if not CASE_RE.fullmatch(name) or set(name) <= {"."}:
                    continue
                case = db.query(Case).filter_by(case_number=name).first() or db.query(Case).filter_by(case_number=name.upper()).first()
                files = [f for f in sorted(case_dir.iterdir()) if not f.name.startswith(".")]
                if case is None:
                    if files:
                        result["skipped_unknown_case"].append(name)
                    continue
                cn = case.case_number
                for f in files:
                    try:
                        st = f.lstat()
                    except OSError:
                        continue
                    if f.is_symlink() or not statmod.S_ISREG(st.st_mode) or f.suffix.lower() in {".part", ".tmp", ".crdownload"} or f.name.endswith("~"):
                        continue
                    if time.time() - st.st_mtime < min_age:
                        result["pending_recent"] += 1
                        continue
                    entry = {"case_number": cn, "file": f.name}
                    try:
                        if st.st_size > max_bytes:
                            raise ValueError(f"File exceeds the {ctx.settings.max_upload_mb} MB upload limit")
                        data = f.read_bytes()
                        suffix = A.validate_upload(f.name, data)
                        digest = hashlib.sha256(data).hexdigest()
                        if suffix in {".csv", ".xlsx"}:
                            res = ctx.svc.ingest_table(db, data, f.name, INBOX_ACTOR["sub"], cn)
                            ctx.hooks.emit("table_ingested", db=db, user=INBOX_ACTOR, result=res, case_number=cn)
                            entry["result"] = {"kind": "table", "rows": res.get("rows")}
                        else:
                            res = ctx.svc.extract_document(db, data, f.name, case_number=cn, media_type=mimetypes.guess_type(f.name)[0] or "application/octet-stream",
                                                           actor=INBOX_ACTOR["sub"], user=INBOX_ACTOR)
                            ctx.hooks.emit("document_extracted", db=db, user=INBOX_ACTOR, result=res, case_number=cn)
                            entry["result"] = {"kind": "document", "document_id": res.get("document_id"), "status": res.get("status")}
                        entry["moved_to"] = _inbox_move(root, f, "_processed", cn)
                        ctx.audit(db, INBOX_ACTOR["sub"], "inbox.ingested", {"case_number": cn, "filename": f.name, "sha256": digest, **entry["result"]})
                        result["processed"].append(entry); _InboxState.totals["processed"] += 1
                    except Exception as exc:  # noqa: BLE001 - one bad file must never stop the pass
                        db.rollback()
                        reason = str(exc)[:300] if isinstance(exc, ValueError) else f"{exc.__class__.__name__}: could not be processed"
                        entry["reason"] = reason
                        try:
                            entry["moved_to"] = _inbox_move(root, f, "_failed", cn, reason)
                        except OSError:
                            entry["moved_to"] = None
                        ctx.audit(db, INBOX_ACTOR["sub"], "inbox.failed", {"case_number": cn, "filename": f.name, "reason": reason})
                        result["failed"].append(entry); _InboxState.totals["failed"] += 1
        finally:
            db.close()
        _InboxState.last_scan_at = _now().isoformat()
        _InboxState.last_result = {k: (len(v) if isinstance(v, list) else v) for k, v in result.items()}
        for e in result["processed"]:
            _InboxState.recent.appendleft({"time": _InboxState.last_scan_at, "outcome": "processed", **e})
        for e in result["failed"]:
            _InboxState.recent.appendleft({"time": _InboxState.last_scan_at, "outcome": "failed", **e})
        return result


def _inbox_loop(ctx):
    while True:
        try:
            _inbox_scan(ctx, min_age=3.0)
        except Exception:  # noqa: BLE001
            log.exception("inbox scan failed")
        time.sleep(10)


def _inbox_pending(root: Path) -> list[dict]:
    out = []
    try:
        if root.is_dir():
            for d in sorted(root.iterdir()):
                if d.name.startswith(("_", ".")) or d.is_symlink() or not d.is_dir() or not CASE_RE.fullmatch(d.name):
                    continue
                n = sum(1 for f in d.iterdir() if f.is_file() and not f.is_symlink() and not f.name.startswith("."))
                if n:
                    out.append({"case_number": d.name, "files": n})
    except OSError:
        pass
    return out


def on_startup(ctx):
    if os.environ.get("INBOX_WATCH", "").strip().lower() in {"1", "true", "yes", "on"}:
        _InboxState.enabled = True
        _inbox_dir().mkdir(parents=True, exist_ok=True)
        if _InboxState.thread is None or not _InboxState.thread.is_alive():
            t = threading.Thread(target=_inbox_loop, args=(ctx,), name="inbox-watcher", daemon=True)
            _InboxState.thread = t
            t.start()
            log.info("Folder inbox watcher started on %s", _inbox_dir())


# ---------------------------------------------------------------------------------------------- router
async def _read_limited(file: UploadFile, limit: int) -> bytes:
    buf = bytearray()
    while True:
        chunk = await file.read(1024 * 1024)
        if not chunk:
            break
        buf += chunk
        if len(buf) > limit:
            raise HTTPException(413, f"File exceeds the {limit // (1024 * 1024)} MB upload limit")
    return bytes(buf)


def build(ctx):
    from security.auth import bump_token_version
    r = APIRouter(prefix="/governance", tags=["governance"])
    perm = ctx.require_perm
    get_db = ctx.get_db

    # ------------------------------------------------------------ a) dashboard
    @r.get("/dashboard")
    def dashboard(days: int = Query(90, ge=1, le=730), db: Session = Depends(get_db), user=Depends(perm("oversight"))):
        return _dashboard(ctx, db, user, days)

    # ------------------------------------------------------------ b) data protection
    @r.post("/redacted-export")
    def redacted_export(x: RedactedExportIn, db: Session = Depends(get_db), user=Depends(perm("report"))):
        cn = ctx.scope(db, user, x.case_number)
        cats = {k for k, v in x.redact.model_dump().items() if k != "free_text_patterns" and v}
        extra = _compile_free_patterns(x.redact.free_text_patterns)
        if extra:
            cats.add("free_text_patterns")
        if not cats:
            raise HTTPException(400, "Select at least one redaction category")
        data = _collect_export_data(ctx, db, user, cn)
        out, log_ = _run_redaction(data, cats, x.mode, extra)
        if x.preview:
            return {"preview": True, "case_number": cn, "redaction_log": log_}
        payload = {"case_number": cn, "exported_at": _now().isoformat() + "Z", "exported_by": user["sub"],
                   "note": "Redacted export. Structured fields only; verified relationships only. Not a finding of guilt.",
                   **out, "redaction_log": log_}
        if x.format == "json":
            body, mt, ext = json.dumps(payload, indent=1, ensure_ascii=False, default=str).encode("utf-8"), "application/json", "json"
        elif x.format == "csv":
            body, mt, ext = _export_csv(out, log_), "text/csv", "csv"
        else:
            body, mt, ext = _redacted_docx(cn, out, log_), DOCX_MT, "docx"
        body = ctx.stamp(db, user, cn, "redacted_export", ext, body)  # ledger record + visible watermark
        ctx.audit(db, user["sub"], "governance.redacted_export", {"case_number": cn, "format": x.format, "mode": x.mode, "categories": sorted(cats),
                                                                  "replacements": log_["replacements_by_category"], "records": log_["records"],
                                                                  "output_sha256": hashlib.sha256(body).hexdigest()})
        return Response(content=body, media_type=mt, headers=_cd(f"redacted_export_{cn}.{ext}"))

    @r.get("/retention")
    def retention(db: Session = Depends(get_db), user=Depends(perm("manage_cases"))):
        rows = _retention_rows(db)
        return {"cases": rows, "summary": {"total": len(rows), "on_hold": sum(1 for x in rows if x["legal_hold"]),
                                           "past_retention": sum(1 for x in rows if x["past_retention"]),
                                           "purge_review_candidates": sum(1 for x in rows if x["purge_review_candidate"])},
                "classifications": list(CLASSIFICATIONS),
                "basis": "Retention runs from the case's last recorded activity (creation, document, evidence, alert or note). "
                         "Being past retention only nominates a case for review; purging always needs an administrator's typed confirmation."}

    @r.put("/cases/{case_number}/hold")
    def set_hold(case_number: str, x: HoldIn, db: Session = Depends(get_db), user=Depends(perm("manage_cases"))):
        cn = ctx.scope(db, user, _clean_case(case_number))
        g = _gov_row(db, cn, create=True)
        was = bool(g.legal_hold)
        g.legal_hold = x.legal_hold
        if x.legal_hold:
            g.hold_reason, g.hold_by, g.hold_at = x.reason, user["sub"], _now()
        db.commit()
        ctx.audit(db, user["sub"], "governance.legal_hold_placed" if x.legal_hold else "governance.legal_hold_released",
                  {"case_number": cn, "reason": x.reason, "previously_on_hold": was})
        return {"case_number": cn, "legal_hold": bool(g.legal_hold), "hold_reason": g.hold_reason if g.legal_hold else "", "hold_by": g.hold_by if g.legal_hold else ""}

    @r.put("/cases/{case_number}/retention")
    def set_retention(case_number: str, x: RetentionIn, db: Session = Depends(get_db), user=Depends(perm("manage_cases"))):
        cn = ctx.scope(db, user, _clean_case(case_number))
        cls = None
        if x.classification is not None:
            cls = x.classification.strip().upper()
            if cls not in CLASSIFICATIONS:
                raise HTTPException(400, f"classification must be one of {', '.join(CLASSIFICATIONS)}")
        g = _gov_row(db, cn, create=True)
        g.retention_days = x.retention_days
        if cls:
            g.classification = cls
        db.commit()
        ctx.audit(db, user["sub"], "governance.retention_set", {"case_number": cn, "retention_days": x.retention_days, "classification": g.classification})
        return {"case_number": cn, "retention_days": g.retention_days, "classification": g.classification}

    @r.post("/cases/{case_number}/purge")
    def purge(case_number: str, x: PurgeIn, db: Session = Depends(get_db), user=Depends(perm("manage_users"))):
        cn = _clean_case(case_number)
        return _purge_case(ctx, db, user, cn, x)

    # ------------------------------------------------------------ c) backup
    @r.post("/backup")
    def backup_create(x: BackupIn, db: Session = Depends(get_db), user=Depends(perm("manage_users"))):
        info = _create_backup(ctx, db, user["sub"], x.passphrase)
        ctx.audit(db, user["sub"], "governance.backup_created", {"name": info["name"], "bytes": info["bytes"], "files": info["files"], "sha256": info["sha256"]})
        ctx.hooks.emit("backup_created", db=db, user=user, ref=info["name"], detail={"bytes": info["bytes"], "files": info["files"], "sha256": info["sha256"]})
        return {**info, "message": "Encrypted backup created. Keep the passphrase safe: it cannot be recovered.",
                "size_cap_mb": backup_max_bytes() // (1024 * 1024)}

    @r.get("/backups")
    def backup_list(user=Depends(perm("manage_users"))):
        out = []
        for p in sorted(_backup_dir().glob("*.dcnbak"), reverse=True):
            if not _BACKUP_NAME_RE.fullmatch(p.name) or p.is_symlink():
                continue
            st = p.stat()
            out.append({"name": p.name, "bytes": st.st_size, "modified": datetime.fromtimestamp(st.st_mtime, timezone.utc).replace(tzinfo=None).isoformat()})
        return {"backups": out, "size_cap_mb": backup_max_bytes() // (1024 * 1024),
                "restore": "Offline only: python scripts/restore_backup.py <file> --target <new empty directory>"}

    @r.get("/backups/{name}/download")
    def backup_download(name: str, db: Session = Depends(get_db), user=Depends(perm("manage_users"))):
        p = _backup_path(name)
        ctx.audit(db, user["sub"], "governance.backup_downloaded", {"name": name})
        ctx.hooks.emit("export_generated", db=db, user=user, ref="backup_download", detail={"kind": "backup_download", "format": "dcnbak", "bytes": p.stat().st_size})
        return FileResponse(str(p), media_type="application/octet-stream", filename=name, content_disposition_type="attachment")

    @r.post("/backups/{name}/verify")
    def backup_verify(name: str, x: BackupVerifyIn, db: Session = Depends(get_db), user=Depends(perm("manage_users"))):
        p = _backup_path(name)
        try:
            res = _verify_backup(p.read_bytes(), x.passphrase)
        except ValueError as exc:
            ctx.audit(db, user["sub"], "governance.backup_verify_failed", {"name": name, "reason": str(exc)[:200]})
            raise HTTPException(400, str(exc))
        ctx.audit(db, user["sub"], "governance.backup_verified", {"name": name, "ok": res["ok"], "files_checked": res["files_checked"]})
        return {"name": name, **res}

    # ------------------------------------------------------------ d) sessions & login history
    @r.get("/login-history")
    def login_history(username: str | None = Query(None, max_length=100), days: int = Query(30, ge=1, le=365),
                      db: Session = Depends(get_db), user=Depends(perm("manage_users"))):
        return _login_history(db, (username or "").strip() or None, days)

    @r.get("/my-activity")
    def my_activity(days: int = Query(30, ge=1, le=365), limit: int = Query(100, ge=1, le=500),
                    db: Session = Depends(get_db), user=Depends(perm("read"))):
        since = _now() - timedelta(days=days)
        rows = (db.query(AuditEvent).filter(AuditEvent.actor == user["sub"], AuditEvent.created_at >= since)
                .order_by(AuditEvent.id.desc()).limit(limit).all())
        out = []
        from security.access import retired_demo_cases
        gone = retired_demo_cases(db)
        for a in rows:
            d = _loads(a.details, {}) if a.details and a.details.startswith("{") else {}
            if (d.get("case_number") or d.get("case") or "") in gone:
                continue
            out.append({"id": a.id, "time": _iso(a.created_at), "action": a.action, "case_number": d.get("case_number") or d.get("case") or "",
                        "detail": {k: v for k, v in d.items() if k not in {"case_number", "reason"} and isinstance(v, (str, int, float, bool))}})
        return {"username": user["sub"], "window_days": days, "events": out}

    @r.post("/users/{username}/force-logout")
    def force_logout(username: str, db: Session = Depends(get_db), user=Depends(perm("manage_users"))):
        row = db.query(User).filter_by(username=username).first()
        if not row:
            raise HTTPException(404, "User not found")
        bump_token_version(row)
        db.commit()
        is_self = username == user["sub"]
        ctx.audit(db, user["sub"], "governance.force_logout", {"target": username, "self": is_self})
        return {"status": "ok", "username": username, "self": is_self,
                "warning": "You signed yourself out too; sign in again." if is_self else None}

    # ------------------------------------------------------------ e) legal packs
    def _legal_case(db, user, case_number):
        return ctx.scope(db, user, case_number)

    @r.get("/legal/bsa63-certificate")
    def bsa63(evidence_id: str = Query(..., max_length=100), case_number: str = Query(..., max_length=100), format: str = Query("docx", pattern="^(docx|pdf)$"),
              db: Session = Depends(get_db), user=Depends(perm("report"))):
        cn = _legal_case(db, user, case_number)
        ev = db.query(Evidence).filter_by(evidence_id=evidence_id, case_number=cn).first()
        if not ev:
            raise HTTPException(404, "Evidence not found in this case")
        body, mt = _render(_bsa63_blocks(db, user, cn, ev), format)
        body = ctx.stamp(db, user, cn, "bsa63_certificate", format, body)
        ctx.audit(db, user["sub"], "legal.bsa63_generated", {"case_number": cn, "evidence_id": evidence_id, "format": format, "sha256": ev.sha256})
        return Response(content=body, media_type=mt, headers=_cd(f"bsa63_certificate_{evidence_id}.{format}"))

    @r.get("/legal/evidence-index")
    def evidence_index(case_number: str = Query(..., max_length=100), format: str = Query("docx", pattern="^(docx|pdf)$"),
                       db: Session = Depends(get_db), user=Depends(perm("report"))):
        cn = _legal_case(db, user, case_number)
        body, mt = _render(_evidence_index_blocks(db, user, cn), format, landscape=True)
        body = ctx.stamp(db, user, cn, "evidence_index", format, body)
        ctx.audit(db, user["sub"], "legal.evidence_index_generated", {"case_number": cn, "format": format})
        return Response(content=body, media_type=mt, headers=_cd(f"evidence_index_{cn}.{format}"))

    @r.get("/legal/document-index")
    def document_index(case_number: str = Query(..., max_length=100), format: str = Query("docx", pattern="^(docx|pdf)$"),
                       db: Session = Depends(get_db), user=Depends(perm("report"))):
        cn = _legal_case(db, user, case_number)
        body, mt = _render(_document_index_blocks(db, user, cn), format, landscape=True)
        body = ctx.stamp(db, user, cn, "document_index", format, body)
        ctx.audit(db, user["sub"], "legal.document_index_generated", {"case_number": cn, "format": format})
        return Response(content=body, media_type=mt, headers=_cd(f"document_index_{cn}.{format}"))

    # ------------------------------------------------------------ f) import connectors
    @r.post("/import/preview")
    async def import_preview(file: UploadFile = File(...), db: Session = Depends(get_db), user=Depends(perm("write"))):
        data = await _read_limited(file, ctx.settings.max_upload_bytes)
        filename = file.filename or "upload.csv"
        try:
            df = await run_in_threadpool(_read_table, data, filename)
        except ValueError as exc:
            raise HTTPException(400, str(exc))
        cols = list(df.columns)
        sug = suggest_mapping(cols)
        token = _stash_upload(user["sub"], filename, data)
        return {"upload_token": token, "filename": filename, "rows": len(df), "columns": [
                    {"name": c, "samples": [_clean_cell(v, 60) for v in df[c].head(4).tolist()], "suggested": sug.get(c)} for c in cols],
                "preview_rows": [{c: _clean_cell(v, 80) for c, v in row.items()} for row in df.head(8).to_dict("records")],
                "fields": [{"key": k, "label": lab, "required": req} for k, lab, req, _ in FIR_FIELDS],
                "note": "Suggestions are heuristics from header names. Check every mapping; imported rows become 'Pending Review' FIR documents."}

    @r.post("/import/commit")
    def import_commit(x: ImportCommitIn, db: Session = Depends(get_db), user=Depends(perm("write"))):
        return _commit_import(ctx, db, user, x)

    @r.get("/import/profiles")
    def profiles(db: Session = Depends(get_db), user=Depends(perm("write"))):
        return {"profiles": [{"id": p.id, "name": p.name, "mapping": _loads(p.mapping, {}), "created_by": p.created_by, "created_at": _iso(p.created_at)}
                             for p in db.query(ImportProfile).order_by(ImportProfile.name).all()]}

    def _check_mapping(mapping: dict):
        bad = [f for f in mapping if f not in FIELD_KEYS]
        if bad:
            raise HTTPException(400, f"Unknown target field(s): {', '.join(bad)}")
        if any(len(str(v)) > 200 for v in mapping.values()):
            raise HTTPException(400, "Column names are limited to 200 characters")

    @r.post("/import/profiles")
    def profile_create(x: ProfileIn, db: Session = Depends(get_db), user=Depends(perm("write"))):
        _check_mapping(x.mapping)
        if db.query(ImportProfile).filter_by(name=x.name.strip()).first():
            raise HTTPException(409, "A profile with this name already exists")
        row = ImportProfile(name=x.name.strip(), mapping=json.dumps(x.mapping), created_by=user["sub"])
        db.add(row); db.commit()
        ctx.audit(db, user["sub"], "import.profile_created", {"name": row.name, "fields": len(x.mapping)})
        return {"id": row.id, "name": row.name}

    @r.put("/import/profiles/{profile_id}")
    def profile_update(profile_id: int, x: ProfileIn, db: Session = Depends(get_db), user=Depends(perm("write"))):
        row = db.query(ImportProfile).filter_by(id=profile_id).first()
        if not row:
            raise HTTPException(404, "Profile not found")
        if row.created_by != user["sub"] and user["role"] != "admin":
            raise HTTPException(403, "Only the creator or an administrator can change this profile")
        _check_mapping(x.mapping)
        clash = db.query(ImportProfile).filter(ImportProfile.name == x.name.strip(), ImportProfile.id != profile_id).first()
        if clash:
            raise HTTPException(409, "A profile with this name already exists")
        row.name, row.mapping = x.name.strip(), json.dumps(x.mapping)
        db.commit()
        ctx.audit(db, user["sub"], "import.profile_updated", {"id": profile_id, "name": row.name})
        return {"id": row.id, "name": row.name}

    @r.delete("/import/profiles/{profile_id}")
    def profile_delete(profile_id: int, db: Session = Depends(get_db), user=Depends(perm("write"))):
        row = db.query(ImportProfile).filter_by(id=profile_id).first()
        if not row:
            raise HTTPException(404, "Profile not found")
        if row.created_by != user["sub"] and user["role"] != "admin":
            raise HTTPException(403, "Only the creator or an administrator can delete this profile")
        name = row.name
        db.delete(row); db.commit()
        ctx.audit(db, user["sub"], "import.profile_deleted", {"id": profile_id, "name": name})
        return {"status": "ok"}

    @r.get("/inbox/status")
    def inbox_status(user=Depends(perm("write"))):
        root = _inbox_dir()
        return {"enabled": _InboxState.enabled, "watching": bool(_InboxState.thread and _InboxState.thread.is_alive()),
                "poll_seconds": 10, "directory": str(root) if user["role"] == "admin" else None, "exists": root.is_dir(),
                "last_scan_at": _InboxState.last_scan_at, "last_result": _InboxState.last_result, "totals": dict(_InboxState.totals),
                "pending": _inbox_pending(root), "recent": list(_InboxState.recent)[:20],
                "how_it_works": "Drop files into <inbox>/<CASE-NUMBER>/. Documents become 'Pending Review'; CSV/XLSX tables use the standard table ingestion. "
                                "Processed files move to _processed, rejected ones to _failed with a reason. Enable with INBOX_WATCH=true."}

    @r.post("/inbox/scan-now")
    def inbox_scan(x: ScanIn | None = None, user=Depends(perm("manage_users"))):
        x = x or ScanIn()
        res = _inbox_scan(ctx, min_age=float(x.min_age_seconds))
        return {"processed": res["processed"], "failed": res["failed"], "skipped_unknown_case": res["skipped_unknown_case"],
                "pending_recent": res["pending_recent"]}

    return r
