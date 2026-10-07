"""Data integrity & leak prevention: a signed, hash-chained, append-only ledger.

Every important event (evidence stored / downloaded / verified, documents extracted and reviewed, tables ingested,
cases created, exports and reports generated, social and surveillance entries ...) becomes a *block*:

    block_hash = SHA-256(index, timestamp, event type, actor, case, ref, payload_sha256, prev_hash, key_id)
    signature  = Ed25519(server key, block_hash)

* Each block contains the hash of the block before it -> editing, deleting or re-ordering any block breaks every later link.
* Every N blocks an ANCHOR block records the Merkle root of the batch, so a single event has a small inclusion proof.
* Blocks carry hashes / ids / counts only, never personal data or file content.
* The signing key lives in ``.ledger_ed25519`` (0600, git-ignored). Key rotation is a block signed by the OLD key.
* A witness file (``.ledger_heads.log``, next to the key, outside the database) receives a signed copy of the head at
  every anchor, so that cutting off the newest blocks of the database is detected as well.
* ``POST /integrity/verify`` re-checks the chain, re-hashes evidence files and live database rows against what the
  ledger recorded, and reports exactly which blocks / rows / files differ.
* ``GET /integrity/bundle`` produces a signed JSON file that ``scripts/verify_ledger.py`` verifies offline.

What it cannot do: nothing stops a determined insider who can read a screen from photographing it; the controls make
leaks attributable (watermark + ledger reference) and tampering detectable, not impossible. Removing the newest blocks
is only detectable against a head hash that was published or witnessed elsewhere.
"""
from __future__ import annotations

import hashlib
import io
import json
import logging
import os
import re
import secrets
import threading
import time
from datetime import datetime, timedelta, timezone
from pathlib import Path

from cryptography.hazmat.primitives import serialization
from cryptography.hazmat.primitives.asymmetric.ed25519 import Ed25519PrivateKey, Ed25519PublicKey
from fastapi import APIRouter, Depends, HTTPException, Query
import sqlalchemy as sa
from sqlalchemy import Column, Integer, String, Text, func
from sqlalchemy.exc import IntegrityError, OperationalError

from backend.app.core.config import ROOT, settings
from backend.app.db.database import Base, SessionLocal
from backend.app.models.models import DocumentRecord, Entity, Evidence, Relationship, User
from security import evidence_crypto
from security.access import retired_demo_cases, visible_cases
from security.auth import require_any_perm
from security.masking import mask_identifier
from security.rbac import allowed

log = logging.getLogger("dcn.integrity")

VERSION = 1
DOMAIN = "DCN-LEDGER-v1"
ZERO = "0" * 64
BUNDLE_FORMAT = "dcn-ledger-bundle/1"
EXPORT_TYPES = ("export_generated", "evidence_downloaded")
CHAIN_NOTE = ("Removing the newest blocks of the ledger can only be detected against a head hash that was recorded somewhere else "
              "(the witness file on the server, or a hash printed / exported earlier). Record the head regularly.")


# ================================================================================================ models
class LedgerBlock(Base):
    __tablename__ = "ledger_blocks"
    idx = Column(Integer, primary_key=True, autoincrement=False)
    ts = Column(String(32), nullable=False)
    event_type = Column(String(60), index=True, nullable=False)
    actor = Column(String(100), index=True, default="")
    case_number = Column(String(100), index=True, default="")
    ref = Column(String(200), index=True, default="")
    payload = Column(Text, default="{}")
    payload_sha256 = Column(String(64), nullable=False)
    prev_hash = Column(String(64), nullable=False)
    block_hash = Column(String(64), unique=True, nullable=False)
    key_id = Column(String(32), nullable=False)
    signature = Column(String(128), nullable=False)


class LedgerRow(Base):
    """Speed-up index of 'which database row was last recorded with which hash'. Never trusted by verification, which
    rebuilds the same facts from the signed blocks."""
    __tablename__ = "ledger_rows"
    table_name = Column(String(80), primary_key=True)
    row_key = Column(String(200), primary_key=True)
    row_hash = Column(String(64), nullable=False)
    block_idx = Column(Integer, index=True)
    case_number = Column(String(100), index=True, default="")


class LedgerError(Exception):
    pass


# ================================================================================================ canonical forms
def canon(obj) -> str:
    return json.dumps(obj, sort_keys=True, separators=(",", ":"), ensure_ascii=True, default=str)


def sha256_hex(data: bytes | str) -> str:
    return hashlib.sha256(data.encode() if isinstance(data, str) else data).hexdigest()


def compute_block_hash(idx, ts, event_type, actor, case_number, ref, payload_sha256, prev_hash, key_id) -> str:
    return sha256_hex(canon([DOMAIN, idx, ts, event_type, actor or "", case_number or "", ref or "", payload_sha256, prev_hash, key_id]))


def signing_message(block_hash: str) -> bytes:
    return f"{DOMAIN}|block|".encode() + bytes.fromhex(block_hash)


def _now_ts() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%S.%fZ")


def fingerprint_of(pub_hex: str) -> str:
    return sha256_hex(bytes.fromhex(pub_hex))


def key_id_of(pub_hex: str) -> str:
    return fingerprint_of(pub_hex)[:16]


# ================================================================================================ Merkle tree
def _leaf(h: str) -> bytes:
    return hashlib.sha256(b"\x00" + bytes.fromhex(h)).digest()


def _node(a: bytes, b: bytes) -> bytes:
    return hashlib.sha256(b"\x01" + a + b).digest()


def merkle_root(leaves: list[str]) -> str:
    if not leaves:
        raise ValueError("no leaves")
    level = [_leaf(h) for h in leaves]
    while len(level) > 1:
        level = [_node(level[i], level[i + 1]) if i + 1 < len(level) else level[i] for i in range(0, len(level), 2)]
    return level[0].hex()


def merkle_proof(leaves: list[str], pos: int) -> list[list[str]]:
    """Path from leaf ``pos`` to the root: [["L"|"R", sibling_hex], ...]; an unpaired node has no sibling at that level."""
    level = [_leaf(h) for h in leaves]
    path = []
    while len(level) > 1:
        if pos % 2 == 0:
            if pos + 1 < len(level):
                path.append(["R", level[pos + 1].hex()])
        else:
            path.append(["L", level[pos - 1].hex()])
        level = [_node(level[i], level[i + 1]) if i + 1 < len(level) else level[i] for i in range(0, len(level), 2)]
        pos //= 2
    return path


def merkle_verify(leaf_hash: str, path: list, root: str) -> bool:
    cur = _leaf(leaf_hash)
    for side, sib in path:
        s = bytes.fromhex(sib)
        cur = _node(s, cur) if side == "L" else _node(cur, s)
    return cur.hex() == root


# ================================================================================================ key handling
def _key_path() -> Path:
    env = os.environ.get("LEDGER_KEY_PATH", "").strip()
    if env:
        return Path(env)
    url = settings.database_url
    if url.startswith("sqlite:///"):
        db = Path(url[len("sqlite:///"):])
        if db.is_absolute() and db.parent.resolve() != ROOT.resolve():
            return db.parent / ".ledger_ed25519"  # isolated deployments / tests keep their key beside their database
    return ROOT / ".ledger_ed25519"


def _witness_path() -> Path:
    p = _key_path()
    return p.with_name(".ledger_heads.log")


def _pub_hex(priv: Ed25519PrivateKey) -> str:
    return priv.public_key().public_bytes(serialization.Encoding.Raw, serialization.PublicFormat.Raw).hex()


def _write_secret(path: Path, priv: Ed25519PrivateKey, exclusive: bool) -> None:
    raw = priv.private_bytes(serialization.Encoding.Raw, serialization.PrivateFormat.Raw, serialization.NoEncryption()).hex()
    path.parent.mkdir(parents=True, exist_ok=True)
    if exclusive:
        fd = os.open(str(path), os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
    else:
        tmp = path.with_name(path.name + ".tmp")
        fd = os.open(str(tmp), os.O_WRONLY | os.O_CREAT | os.O_TRUNC, 0o600)
    with os.fdopen(fd, "w") as fh:
        fh.write(raw + "\n")
    if not exclusive:
        os.replace(str(path.with_name(path.name + ".tmp")), str(path))
    try:
        os.chmod(path, 0o600)
    except OSError:
        pass


def _load_key(path: Path) -> Ed25519PrivateKey | None:
    if not path.exists():
        return None
    return Ed25519PrivateKey.from_private_bytes(bytes.fromhex(path.read_text().strip()))


def _load_or_create_key() -> Ed25519PrivateKey:
    path = _key_path()
    try:
        k = _load_key(path)
    except (ValueError, OSError) as exc:
        raise LedgerError(f"ledger signing key {path} is unreadable: {exc.__class__.__name__}") from exc
    if k is not None:
        return k
    k = Ed25519PrivateKey.generate()
    try:
        _write_secret(path, k, exclusive=True)
    except FileExistsError:  # another process won the race
        return _load_key(path)  # type: ignore[return-value]
    log.warning("ledger: generated a new Ed25519 signing key at %s (key id %s)", path, key_id_of(_pub_hex(k)))
    return k


# ================================================================================================ ledger core
class _State:
    lock = threading.RLock()
    missed = 0
    missed_rows: list = []
    first_error_ts: str | None = None
    last_error: str = ""
    verify_cache: dict = {}
    wq: list = []


def _anchor_every() -> int:
    try:
        return max(2, int(os.environ.get("LEDGER_ANCHOR_EVERY", "50")))
    except ValueError:
        return 50


def _anchor_seconds() -> int:
    try:
        return max(30, int(os.environ.get("LEDGER_ANCHOR_SECONDS", "3600")))
    except ValueError:
        return 3600


def _head(s):
    return s.query(LedgerBlock).order_by(LedgerBlock.idx.desc()).first()


def _active_key_id(s) -> str | None:
    """Key that must sign the next block: the newest genesis / key_rotation block names it."""
    b = (s.query(LedgerBlock).filter(LedgerBlock.event_type.in_(("genesis", "key_rotation")))
         .order_by(LedgerBlock.idx.desc()).first())
    if b is None:
        return None
    return json.loads(b.payload)["key_id" if b.event_type == "genesis" else "new_key_id"]


def _promote_pending_key(active_id: str | None) -> None:
    """A rotation writes the new key to '<key>.next' before the block commits; finish the swap after a crash."""
    path = _key_path()
    nxt = path.with_name(path.name + ".next")
    if nxt.exists():
        try:
            k = _load_key(nxt)
            if k is not None and key_id_of(_pub_hex(k)) == active_id:
                os.replace(str(nxt), str(path))
        except (ValueError, OSError):
            log.exception("ledger: could not promote pending key")


def _make_block(s, priv, idx, prev_hash, event_type, actor, case_number, ref, payload: dict, ts: str | None = None) -> LedgerBlock:
    body = canon(payload)
    pub = _pub_hex(priv)
    kid = key_id_of(pub)
    ts = ts or _now_ts()
    psha = sha256_hex(body)
    bh = compute_block_hash(idx, ts, event_type, actor, case_number, ref, psha, prev_hash, kid)
    sig = priv.sign(signing_message(bh)).hex()
    return LedgerBlock(idx=idx, ts=ts, event_type=event_type, actor=actor or "", case_number=case_number or "", ref=ref or "",
                       payload=body, payload_sha256=psha, prev_hash=prev_hash, block_hash=bh, key_id=kid, signature=sig)


def _index_rows(s, rows: list[dict], idx: int, case_number: str) -> None:
    if not rows:
        return
    from sqlalchemy.dialects.sqlite import insert as sqlite_insert
    vals = [{"table_name": r["t"], "row_key": str(r["k"])[:200], "row_hash": r["h"], "block_idx": idx, "case_number": r.get("c") or case_number or ""}
            for r in rows]
    bind = s.get_bind()
    if bind.dialect.name == "sqlite":
        for i in range(0, len(vals), 200):
            stmt = sqlite_insert(LedgerRow).values(vals[i:i + 200])
            stmt = stmt.on_conflict_do_update(index_elements=["table_name", "row_key"],
                                              set_={"row_hash": stmt.excluded.row_hash, "block_idx": stmt.excluded.block_idx,
                                                    "case_number": stmt.excluded.case_number})
            s.execute(stmt)
    else:  # pragma: no cover - generic fallback
        for v in vals:
            s.merge(LedgerRow(**v))


def _pending_range(s):
    """(from_idx, to_idx) of the blocks not yet covered by an anchor, or None."""
    last = s.query(LedgerBlock).filter(LedgerBlock.event_type == "anchor").order_by(LedgerBlock.idx.desc()).first()
    start = last.idx + 1 if last else 0
    head = _head(s)
    if head is None or head.idx < start:
        return None
    return start, head.idx


def _ensure_genesis(s, priv) -> LedgerBlock:
    head = _head(s)
    if head is None:
        mykid = key_id_of(_pub_hex(priv))
        genesis_payload = {"version": VERSION, "public_key": _pub_hex(priv), "key_id": mykid,
                           "note": "Genesis block: the server's first signing key. Pin the fingerprint outside the server."}
        head = _make_block(s, priv, 0, ZERO, "genesis", "system", "", "", genesis_payload)
        s.add(head)
        s.flush()
    return head


def _append_locked(s, priv, event_type, actor, case_number, ref, payload, *, builder=None, rows=None, allow_anchor=True) -> LedgerBlock:
    """Add one block inside an open session (caller holds the lock and commits). ``builder(idx, ts)`` may compute the
    payload once the block's index is known (used by watermarked exports)."""
    head = _ensure_genesis(s, priv)
    active = _active_key_id(s)
    mykid = key_id_of(_pub_hex(priv))
    if active != mykid:
        raise LedgerError(f"signing key {mykid} is not the ledger's current key {active} (key file replaced or database restored from elsewhere)")
    if allow_anchor:
        pend = _pending_range(s)
        if pend and (pend[1] - pend[0] + 1) >= _anchor_every():
            head = _anchor_locked(s, priv, "system") or head
        elif pend:
            first = s.get(LedgerBlock, pend[0])
            try:
                age = (datetime.now(timezone.utc) - datetime.strptime(first.ts, "%Y-%m-%dT%H:%M:%S.%fZ").replace(tzinfo=timezone.utc)).total_seconds()
            except (ValueError, AttributeError):
                age = 0
            if age > _anchor_seconds():
                head = _anchor_locked(s, priv, "system") or head
    if _State.missed and event_type != "ledger_gap":
        lost = list(_State.missed_rows[:400])
        gap = _make_block(s, priv, head.idx + 1, head.block_hash, "ledger_gap", "system", "", "",
                          {"missed_events": _State.missed, "first_failure_at": _State.first_error_ts, "last_error": _State.last_error[:200],
                           "note": "Some events could not be written to the ledger; the audit log may hold them. Rows they touched are recorded here with their current hashes.",
                           **({"rows": lost} if lost else {})})
        s.add(gap)
        s.flush()
        if lost:
            _index_rows(s, lost, gap.idx, "")
        head = gap
    idx = head.idx + 1
    ts = _now_ts()
    if ts < head.ts:  # never let a clock step backwards make timestamps non-monotonic
        ts = head.ts
    if builder is not None:
        payload = builder(idx, ts)
    blk = _make_block(s, priv, idx, head.block_hash, event_type, actor, case_number, ref, payload, ts=ts)
    s.add(blk)
    s.flush()
    if rows:
        _index_rows(s, rows, idx, case_number)
    return blk


def _anchor_locked(s, priv, actor):
    pend = _pending_range(s)
    if not pend:
        return None
    a, b = pend
    leaves = [x[0] for x in s.query(LedgerBlock.block_hash).filter(LedgerBlock.idx >= a, LedgerBlock.idx <= b,
                                                                     LedgerBlock.event_type != "anchor").order_by(LedgerBlock.idx).all()]
    if not leaves:
        return None
    head = _head(s)
    blk = _make_block(s, priv, head.idx + 1, head.block_hash, "anchor", actor, "", "",
                      {"from": a, "to": b, "count": len(leaves), "merkle_root": merkle_root(leaves), "leaf_rule": "sha256(0x00||block_hash), node sha256(0x01||L||R), odd node promoted"})
    s.add(blk)
    s.flush()
    _State.wq.append((blk.idx, blk.block_hash, blk.ts, blk.key_id))  # written to the witness file only after the commit
    return blk


def _flush_witness(priv) -> None:
    items, _State.wq = list(_State.wq), []
    for idx, bh, ts, kid in items:
        _witness_line(idx, bh, ts, kid, priv)


def _witness(blk: LedgerBlock, priv) -> None:
    _witness_line(blk.idx, blk.block_hash, blk.ts, blk.key_id, priv)


def _witness_line(idx, bh, ts, kid, priv) -> None:
    """Append a signed copy of a head to the witness file (outside the database)."""
    try:
        line = {"index": idx, "hash": bh, "ts": ts, "key_id": kid}
        line["sig"] = priv.sign(f"{DOMAIN}|head|{idx}|{bh}".encode()).hex()
        p = _witness_path()
        p.parent.mkdir(parents=True, exist_ok=True)
        if p.exists():
            tail = p.read_text().strip().splitlines()[-1:] or [""]
            try:
                if json.loads(tail[0]).get("hash") == bh:
                    return
            except ValueError:
                pass
        with p.open("a") as fh:
            fh.write(json.dumps(line, sort_keys=True) + "\n")
        try:
            os.chmod(p, 0o600)
        except OSError:
            pass
    except OSError:
        log.exception("ledger: could not write witness file")


def _with_retry(fn):
    """Run ``fn(session)`` in its own session/transaction; retry on index clashes (other worker) and short DB locks."""
    last = None
    for attempt in range(6):
        with SessionLocal(expire_on_commit=False) as s:
            try:
                out = fn(s)
                s.commit()
                return out
            except IntegrityError as exc:
                s.rollback()
                last = exc
            except OperationalError as exc:
                s.rollback()
                last = exc
                time.sleep(0.15 * (attempt + 1))
            except Exception:
                s.rollback()
                raise
    raise LedgerError(f"ledger write kept failing: {last.__class__.__name__}")


def append(event_type: str, actor: str = "", case_number: str = "", ref: str = "", payload: dict | None = None, *,
           rows: list[dict] | None = None, builder=None) -> LedgerBlock | None:
    """Append one block. Raises LedgerError on failure (use ``record`` from request paths)."""
    payload = payload or {}
    rows = rows or []
    with _State.lock:
        priv = _load_or_create_key()

        def work(s):
            _State.wq.clear()
            _promote_pending_key(_active_key_id(s))
            first_payload = dict(payload)
            chunk = 400
            if rows:
                first_payload = {**first_payload, "rows": rows[:chunk]}
            blk = _append_locked(s, priv, event_type, actor, case_number, ref, first_payload if builder is None else {},
                                 builder=(lambda i, t: {**builder(i, t), **({"rows": rows[:chunk]} if rows else {})}) if builder else None,
                                 rows=rows[:chunk])
            for off in range(chunk, len(rows), chunk):
                _append_locked(s, priv, "rows_recorded", actor, case_number, ref, {"of_block": blk.idx, "rows": rows[off:off + chunk]},
                               rows=rows[off:off + chunk])
            return blk
        blk = _with_retry(work)
        _flush_witness(priv)
        _State.missed = 0
        _State.missed_rows = []
        _State.first_error_ts = None
        return blk


def record(event_type: str, **kw) -> int | None:
    """Failure-proof wrapper for request paths: never raises; a failure is logged loudly and a ledger_gap marker is
    written with the next successful block."""
    try:
        blk = append(event_type, **kw)
        return blk.idx if blk else None
    except Exception as exc:  # noqa: BLE001
        with _State.lock:
            _State.missed += 1
            _State.first_error_ts = _State.first_error_ts or _now_ts()
            _State.last_error = f"{exc.__class__.__name__}: {exc}"
            _State.missed_rows.extend((kw.get("rows") or [])[:400])
            del _State.missed_rows[2000:]
        log.error("LEDGER WRITE FAILED for %s (%s): the event is NOT in the ledger; a ledger_gap marker will follow", event_type, exc, exc_info=True)
        return None


def anchor_now(actor: str = "system") -> LedgerBlock | None:
    with _State.lock:
        priv = _load_or_create_key()

        def work(s):
            _State.wq.clear()
            _ensure_genesis(s, priv)
            return _anchor_locked(s, priv, actor)
        blk = _with_retry(work)
        _flush_witness(priv)
        return blk


def rotate_key(actor: str) -> dict:
    """Introduce a new signing key with a block signed by the OLD key (and a proof-of-possession by the new key)."""
    with _State.lock:
        old = _load_or_create_key()
        path = _key_path()
        new = Ed25519PrivateKey.generate()
        nxt = path.with_name(path.name + ".next")
        _write_secret(nxt, new, exclusive=False)
        old_id, new_id = key_id_of(_pub_hex(old)), key_id_of(_pub_hex(new))

        def work(s):
            _State.wq.clear()
            head = _ensure_genesis(s, old)
            if _active_key_id(s) != old_id:
                raise LedgerError("current key file does not match the ledger's current key; refusing to rotate")
            pop = new.sign(f"{DOMAIN}|rotate|{old_id}|{new_id}|{head.block_hash}".encode()).hex()
            payload = {"old_key_id": old_id, "new_key_id": new_id, "new_public_key": _pub_hex(new), "proof_of_possession": pop,
                       "prev_hash_at_rotation": head.block_hash}
            return _append_locked(s, old, "key_rotation", actor, "", "", payload, allow_anchor=False)
        try:
            blk = _with_retry(work)
        except Exception:
            try:
                nxt.unlink()
            except OSError:
                pass
            raise
        os.replace(str(nxt), str(path))
        try:
            os.chmod(path, 0o600)
        except OSError:
            pass
        with SessionLocal() as s2:
            _witness(s2.get(LedgerBlock, blk.idx), new)  # type: ignore[arg-type]
        return {"rotation_block": blk.idx, "old_key_id": old_id, "new_key_id": new_id}


# ================================================================================================ row hashing
def _iso(v):
    return v.isoformat() if hasattr(v, "isoformat") else v


def _rowhash(table: str, key, values: list) -> str:
    return sha256_hex(canon([table, str(key), [_iso(v) for v in values]]))


def _loads(text, default=None):
    try:
        v = json.loads(text) if text else default
        return v if v is not None else default
    except (ValueError, TypeError):
        return default


def _case_of_entity(e: Entity) -> str:
    a = _loads(e.attributes, {})
    return (a.get("case_number") or (a.get("case_numbers") or [""])[0] or "") if isinstance(a, dict) else ""


def _case_of_rel(r: Relationship) -> str:
    m = _loads(r.metadata_json, {})
    return (m.get("case_number") or "") if isinstance(m, dict) else ""


def _spec_documents():
    return dict(name="documents", model=DocumentRecord, keyf=lambda r: r.document_id, case=lambda r: r.case_number or "", mutable=True,
                values=lambda r: [r.document_id, r.case_number, r.filename, r.sha256, r.evidence_id, r.status, r.review_actor],
                sql_case=DocumentRecord.case_number)


def _spec_evidence():
    return dict(name="evidence", model=Evidence, keyf=lambda r: r.evidence_id, case=lambda r: r.case_number or "", mutable=False,
                values=lambda r: [r.evidence_id, r.filename, r.sha256, r.case_number, r.uploaded_by],
                sql_case=Evidence.case_number, extra=lambda r: {"x": r.sha256})


def _spec_entities():
    return dict(name="entities", model=Entity, keyf=lambda r: r.external_id, case=_case_of_entity, mutable=False,
                values=lambda r: [r.external_id, r.entity_type, r.name], sql_case=None)


def _spec_relationships():
    return dict(name="relationships", model=Relationship, keyf=lambda r: r.id, case=_case_of_rel, mutable=False,
                values=lambda r: [r.id, r.source_id, r.target_id, r.relation_type, r.source_ref], sql_case=None)


def _generic_specs() -> list[dict]:
    """Tables of the social-media / surveillance features, when they exist: every column except the primary key and
    obvious bookkeeping timestamps is hashed. Amendments are recorded by the feature's own ledger events."""
    out = []
    for name, t in sorted(Base.metadata.tables.items()):
        if not (name.startswith(("surveillance", "social", "surv_", "soc_")) and "case_number" in t.c):
            continue
        pk = [c.name for c in t.primary_key.columns]
        if len(pk) != 1:
            continue
        cols = [c.name for c in t.columns if c.name != pk[0] and c.name not in {"updated_at", "modified_at"}]

        def make(t=t, pk=pk[0], cols=cols):
            return dict(name=t.name, table=t, keyf=lambda r: r[pk], case=lambda r: r.get("case_number") or "", mutable=True,
                        values=lambda r: [r.get(c) for c in cols], sql_case=t.c.case_number, generic=True)
        out.append(make())
    return out


def tracked_specs() -> list[dict]:
    return [_spec_documents(), _spec_evidence(), _spec_entities(), _spec_relationships()] + _generic_specs()


def _iter_rows(db, spec, case: str | None = None):
    """Yield (key, hash, case_number, extra) for the rows of one tracked table (optionally one case)."""
    if spec.get("generic"):
        t = spec["table"]
        q = t.select()
        if case:
            q = q.where(t.c.case_number == case)
        for r in db.execute(q).mappings():
            r = dict(r)
            yield str(spec["keyf"](r)), _rowhash(spec["name"], spec["keyf"](r), spec["values"](r)), spec["case"](r), {}
        return
    q = db.query(spec["model"])
    if case and spec.get("sql_case") is not None:
        q = q.filter(spec["sql_case"] == case)
    for r in q.yield_per(500):
        c = spec["case"](r)
        if case and c != case:
            continue
        k = spec["keyf"](r)
        yield str(k), _rowhash(spec["name"], k, spec["values"](r)), c, (spec["extra"](r) if spec.get("extra") else {})


def _row_entry(name, k, h, c, extra) -> dict:
    return {"t": name, "k": k, "h": h, "c": c, **extra}


def rows_for(db, table: str, keys: list, spec_map: dict | None = None) -> list[dict]:
    """Current hash entries for specific rows (used when an event names the row it touched)."""
    specs = spec_map or {s["name"]: s for s in tracked_specs()}
    spec = specs.get(table)
    if not spec:
        return []
    out = []
    if spec.get("generic"):
        t = spec["table"]
        pk = [c.name for c in t.primary_key.columns][0]
        for r in db.execute(t.select().where(t.c[pk].in_(keys))).mappings():
            r = dict(r)
            out.append(_row_entry(table, str(spec["keyf"](r)), _rowhash(table, spec["keyf"](r), spec["values"](r)), spec["case"](r), {}))
        return out
    col = {"documents": DocumentRecord.document_id, "evidence": Evidence.evidence_id, "entities": Entity.external_id,
           "relationships": Relationship.id}[table]
    for r in db.query(spec["model"]).filter(col.in_(keys)):
        k = spec["keyf"](r)
        out.append(_row_entry(table, str(k), _rowhash(table, k, spec["values"](r)), spec["case"](r), spec["extra"](r) if spec.get("extra") else {}))
    return out


def unrecorded_rows(db, case: str | None = None) -> list[dict]:
    """Rows in tracked tables with no ledger record yet (per the speed index)."""
    with SessionLocal() as s:
        known = {(r.table_name, r.row_key) for r in (s.query(LedgerRow.table_name, LedgerRow.row_key).filter(LedgerRow.case_number == case)
                                                    if case else s.query(LedgerRow.table_name, LedgerRow.row_key))}
    out = []
    for spec in tracked_specs():
        try:
            for k, h, c, extra in _iter_rows(db, spec, case):
                if (spec["name"], k) not in known:
                    out.append(_row_entry(spec["name"], k, h, c, extra))
        except Exception:  # noqa: BLE001 - a table that cannot be read never blocks recording
            log.exception("ledger: cannot read tracked table %s", spec["name"])
    return out


def sweep(db, case: str | None, actor: str, event_type: str, ref: str, detail: dict) -> int | None:
    rows = unrecorded_rows(db, case)
    if not rows and event_type == "rows_recorded":
        return None
    return record(event_type, actor=actor, case_number=case or "", ref=ref, payload=detail, rows=rows)


# ================================================================================================ sanitising payloads
_ID_OK = re.compile(r"^[A-Za-z0-9_.:/\-]{1,120}$")
_SAFE_KEY = re.compile(r"(^|_)(id|ids|sha256|hash|count|total|kind|status|format|type|platform|mode|n|rows|size|bytes|version|reason_code|ref)$")


def sanitize(detail) -> dict:
    """Keep numbers, booleans and identifier-like values; replace everything else with a short hash. No personal data
    or free text ever enters the ledger."""
    out: dict = {}
    if not isinstance(detail, dict):
        return out
    for k, v in list(detail.items())[:40]:
        k = str(k)[:60]
        if isinstance(v, bool) or v is None or isinstance(v, (int, float)):
            out[k] = v
        elif isinstance(v, str):
            if _SAFE_KEY.search(k) and _ID_OK.match(v):
                out[k] = v
            elif v:
                out[k + "_sha256"] = sha256_hex(v)[:32]
        elif isinstance(v, (list, tuple, set)):
            out[k + "_count"] = len(v)
        elif isinstance(v, dict):
            out[k + "_keys"] = len(v)
    return out


def _int_fields(d) -> dict:
    return {str(k): v for k, v in (d or {}).items() if isinstance(v, (int, bool)) and not isinstance(v, float)} if isinstance(d, dict) else {}


# ================================================================================================ hook handlers
def _actor(user) -> str:
    return (user or {}).get("sub", "") if isinstance(user, dict) else str(user or "")


def _h_document_extracted(db=None, user=None, result=None, case_number="", **_):
    r = result or {}
    cn = r.get("case_number") or case_number or ""
    rows = rows_for(db, "documents", [r.get("document_id")]) + rows_for(db, "evidence", [r.get("evidence_id")]) if r.get("document_id") else []
    record("document_extracted", actor=_actor(user), case_number=cn, ref=r.get("document_id", ""),
           payload={"document_id": r.get("document_id"), "evidence_id": r.get("evidence_id"), "document_type": r.get("document_type"),
                    "sha256": next((x.get("x") for x in rows if x["t"] == "evidence"), None)}, rows=rows)


def _merge_rows(a: list[dict], b: list[dict]) -> list[dict]:
    seen = {(r["t"], r["k"]) for r in a}
    return a + [r for r in b if (r["t"], r["k"]) not in seen]


def _h_document_reviewed(db=None, user=None, document_id="", payload=None, result=None, **_):
    result = result or {}
    cn = result.get("case_number") or ""
    rows = rows_for(db, "documents", [document_id])
    if not cn and rows:
        cn = rows[0].get("c", "")
    rows = _merge_rows(rows, unrecorded_rows(db, cn or None) if cn else [])
    record("document_reviewed", actor=_actor(user), case_number=cn, ref=document_id,
           payload={"document_id": document_id, "status": str(result.get("status") or (payload or {}).get("status") or "")[:20],
                    "entity_actions": len((payload or {}).get("entities") or []), "relationship_actions": len((payload or {}).get("relationships") or [])},
           rows=rows)


def _h_document_reprocessed(db=None, user=None, document_id="", case_number="", **_):
    rows = rows_for(db, "documents", [document_id])
    record("document_reprocessed", actor=_actor(user), case_number=case_number or (rows[0]["c"] if rows else ""), ref=document_id,
           payload={"document_id": document_id}, rows=rows)


def _h_table_ingested(db=None, user=None, result=None, case_number="", **_):
    cn = case_number or (result or {}).get("case_number") or ""
    rows = unrecorded_rows(db, cn or None)
    record("table_ingested", actor=_actor(user), case_number=cn, ref=cn,
           payload={**_int_fields({k: v for k, v in (result or {}).items() if k in ("rows", "created_entities", "created_relationships", "created_events")})},
           rows=rows)


def _h_evidence_uploaded(db=None, user=None, result=None, case_number="", **_):
    r = result[0] if isinstance(result, list) and result else (result or {})
    eid = r.get("evidence_id")
    rows = rows_for(db, "evidence", [eid]) if eid else []
    record("evidence_uploaded", actor=_actor(user), case_number=r.get("case_number") or case_number, ref=eid or "",
           payload={"evidence_id": eid, "sha256": r.get("sha256"), "document_id": r.get("document_id")}, rows=rows)


def _evidence_sha(db, eid):
    e = db.query(Evidence).filter_by(evidence_id=eid).first() if db is not None else None
    return e.sha256 if e else None


def _h_evidence_downloaded(db=None, user=None, evidence_id="", case_number="", **_):
    idx = record("evidence_downloaded", actor=_actor(user), case_number=case_number, ref=evidence_id,
                 payload={"evidence_id": evidence_id, "sha256": _evidence_sha(db, evidence_id), "purpose": "download of the original file"})
    _check_bulk(_actor(user), case_number)
    return idx


def _h_evidence_verified(db=None, user=None, evidence_id="", result=None, case_number="", **_):
    r = result or {}
    record("evidence_verified", actor=_actor(user), case_number=case_number, ref=evidence_id,
           payload={"evidence_id": evidence_id, "status": r.get("status"), "original_hash": r.get("original_hash"), "current_hash": r.get("current_hash")})


def _h_case_created(db=None, user=None, case_number="", **_):
    record("case_created", actor=_actor(user), case_number=case_number, ref=case_number, payload={"case_number": case_number})


def _h_entity_created(db=None, user=None, case_number="", ref="", **_):
    record("entity_created", actor=_actor(user), case_number=case_number, ref=str(ref), payload={"external_id": str(ref)[:120]},
           rows=rows_for(db, "entities", [ref]))


def _h_relationship_created(db=None, user=None, case_number="", ref="", **_):
    record("relationship_created", actor=_actor(user), case_number=case_number, ref=str(ref), payload={"relationship_id": ref},
           rows=rows_for(db, "relationships", [ref]))


def _generic_event_rows(db, prefixes: tuple, case: str, idvals: set[str]) -> list[dict]:
    """Rows of the social / surveillance tables that an event touched: rows whose primary key or a text column equals one of the
    identifiers the event names (entry uid, operation number, import id ...), plus rows of the case that were never recorded."""
    out: dict = {}
    for spec in _generic_specs():
        if not spec["name"].startswith(prefixes):
            continue
        t = spec["table"]
        pk = [c.name for c in t.primary_key.columns][0]
        textcols = [c.name for c in t.columns if isinstance(c.type, sa.String) and c.name != "case_number"]
        q = t.select()
        if case:
            q = q.where(t.c.case_number == case)
        for r in db.execute(q).mappings():
            r = dict(r)
            k = str(r[pk])
            if k in idvals or any(str(r.get(c)) in idvals for c in textcols):
                out[(spec["name"], k)] = _row_entry(spec["name"], k, _rowhash(spec["name"], r[pk], spec["values"](r)), spec["case"](r), {})
    for r in unrecorded_rows(db, case or None):
        if r["t"].startswith(prefixes):
            out.setdefault((r["t"], r["k"]), r)
    return list(out.values())[:2000]


def _h_generic(event):
    """social_imported, surveillance_* ...: emitted as hooks.emit(name, db=, user=, case_number=, ref=, detail={...})."""
    prefixes = ("surveillance", "surv_") if event.startswith("surveillance") else ("social", "soc_")

    def fn(db=None, user=None, case_number="", ref="", detail=None, **_):
        rows: list[dict] = []
        try:
            ids = {str(ref)} if ref not in (None, "") else set()
            for k, v in (detail or {}).items():
                if isinstance(v, (str, int)) and not isinstance(v, bool) and (k.endswith(("_id", "_uid", "_number", "_ref")) or k in ("supersedes", "uid")):
                    ids.add(str(v))
            rows = _generic_event_rows(db, prefixes, case_number or "", ids)
        except Exception:  # noqa: BLE001
            log.exception("ledger: could not hash rows for %s", event)
        record(event, actor=_actor(user), case_number=case_number or "", ref=str(ref or ""), payload=sanitize(detail), rows=rows)
        if event.endswith(("_generated", "_exported")):
            _check_bulk(_actor(user), case_number)
    return fn


def _h_export_generated(db=None, user=None, case_number="", ref="", detail=None, **_):
    record("export_generated", actor=_actor(user), case_number=case_number or "", ref=str(ref or ""), payload=sanitize(detail))
    _check_bulk(_actor(user), case_number)


def _h_simple(event):
    def fn(db=None, user=None, case_number="", ref="", detail=None, **_):
        record(event, actor=_actor(user), case_number=case_number or "", ref=str(ref or ""), payload=sanitize(detail))
    return fn


def _h_demo_enabled(db=None, user=None, detail=None, **_):
    sweep(db, None, _actor(user), "demo_enabled", "demo", {**sanitize(detail), "baseline": "synthetic demonstration rows recorded"})


def _removed_keys(db, cases: set[str], include_caseless: bool) -> list[dict]:
    """Recorded rows of the given cases that are gone from the database: the attestation that goes into a purge / demo-off block
    (only rows of those cases, so it cannot launder the deletion of anything else)."""
    with SessionLocal() as s:
        idx_rows = [(r.table_name, r.row_key, r.case_number) for r in s.query(LedgerRow.table_name, LedgerRow.row_key, LedgerRow.case_number)]
    wanted = [(t, k) for t, k, c in idx_rows if c in cases or (include_caseless and not c)]
    if not wanted:
        return []
    live: dict[str, set] = {}
    for spec in tracked_specs():
        try:
            live[spec["name"]] = {k for k, _h, _c, _x in _iter_rows(db, spec, None)}
        except Exception:  # noqa: BLE001
            live[spec["name"]] = None  # type: ignore[assignment]
    return [{"t": t, "k": k} for t, k in wanted if live.get(t) is not None and k not in live[t]][:20000]


def _forget_rows(removed: list[dict]) -> None:
    """Drop removed rows from the speed index so a later row that reuses the same key is treated as new."""
    try:
        with SessionLocal() as s:
            for x in removed:
                s.query(LedgerRow).filter(LedgerRow.table_name == x["t"], LedgerRow.row_key == x["k"]).delete()
            s.commit()
    except Exception:  # noqa: BLE001
        log.exception("ledger: could not update the row index after a removal")


def _h_demo_disabled(db=None, user=None, detail=None, cases=None, **_):
    cs = {str(c)[:100] for c in (cases or [])}
    removed = _removed_keys(db, cs, True)
    if record("demo_disabled", actor=_actor(user), payload={**sanitize(detail), "cases": sorted(cs)[:500], "removed": removed}) is not None:
        _forget_rows(removed)


def _h_case_purged(db=None, user=None, case_number="", detail=None, **_):
    removed = _removed_keys(db, {case_number}, False)
    if record("case_purged", actor=_actor(user), case_number=case_number, ref=case_number,
              payload={**sanitize(detail), "cases": [case_number], "removed": removed}) is not None:
        _forget_rows(removed)


# ================================================================================================ exports: watermark + anomaly
def _watermark_text(actor: str, ts: str, idx, wm: str) -> str:
    human = ts.replace("T", " ").rstrip("Z")[:19] + " UTC"
    return f"Exported by {actor} on {human}, ref L-{idx}, watermark {wm}"


def apply_watermark(fmt: str, body: bytes, text: str, meta: dict) -> bytes:
    fmt = (fmt or "").lower()
    if fmt == "csv":
        import csv
        raw = body.decode("utf-8-sig")
        bom = body.startswith(b"\xef\xbb\xbf")
        out = io.StringIO()
        out.write(raw if raw.endswith("\n") or not raw else raw + "\r\n")
        csv.writer(out).writerow(["# WATERMARK", text])
        return out.getvalue().encode("utf-8-sig" if bom else "utf-8")
    if fmt == "json":
        doc = json.loads(body.decode("utf-8"))
        if isinstance(doc, dict):
            doc["watermark"] = {"text": text, **meta}
            return json.dumps(doc, indent=1, ensure_ascii=False, default=str).encode("utf-8")
        return body
    if fmt == "docx":
        from docx import Document
        from docx.shared import Pt, RGBColor
        d = Document(io.BytesIO(body))
        for sec in d.sections:
            fp = sec.footer.paragraphs[0] if sec.footer.paragraphs else sec.footer.add_paragraph()
            run = fp.add_run(text)
            run.font.size = Pt(7)
            run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        p = d.add_paragraph()
        r = p.add_run(text)
        r.italic = True
        r.font.size = Pt(8)
        buf = io.BytesIO()
        d.save(buf)
        return buf.getvalue()
    if fmt == "pdf":
        import fitz
        d = fitz.open(stream=body, filetype="pdf")
        try:
            for page in d:
                page.insert_text((36, page.rect.height - 20), text, fontsize=7, fontname="helv", color=(0.35, 0.35, 0.35))
            return d.tobytes(deflate=True)
        finally:
            d.close()
    return body


def stamp_export(db, user: dict, case_number: str, kind: str, fmt: str, body, purpose: str = ""):
    """Record an export in the ledger and return the body carrying a visible watermark (footer line, CSV comment row,
    JSON ``watermark`` member). ``body`` may be bytes or (for JSON reports) a dict. Never raises: on any failure the
    original body is returned and the failure is logged."""
    actor = _actor(user)
    is_dict = isinstance(body, dict)
    raw = json.dumps(body, sort_keys=True, default=str).encode() if is_dict else bytes(body)
    result: dict = {}

    def build(idx, ts):
        wm = "WM-" + sha256_hex(f"{idx}|{actor}|{ts}|{secrets.token_hex(8)}")[:10].upper()
        text = _watermark_text(actor, ts, idx, wm)
        meta = {"watermark_id": wm, "ledger_ref": f"L-{idx}", "exported_by": actor, "exported_at": ts}
        try:
            if is_dict:
                out = {**body, "watermark": {"text": text, **meta}}
            else:
                out = apply_watermark(fmt, raw, text, meta)
        except Exception:  # noqa: BLE001
            log.exception("watermarking %s/%s failed; delivering the export without it", kind, fmt)
            out, text = (body if is_dict else raw), ""
        result["out"] = out
        final = json.dumps(out, sort_keys=True, default=str).encode() if is_dict else out
        return {"kind": kind, "format": fmt, "bytes": len(final), "output_sha256": sha256_hex(final), "source_sha256": sha256_hex(raw),
                "watermark_id": wm if text else "", "purpose_sha256": sha256_hex(purpose)[:32] if purpose else "",
                "user_role": (user or {}).get("role", "")}
    try:
        blk = append("export_generated", actor=actor, case_number=case_number or "", ref=kind, builder=build)
    except Exception as exc:  # noqa: BLE001
        with _State.lock:
            _State.missed += 1
            _State.first_error_ts = _State.first_error_ts or _now_ts()
            _State.last_error = f"{exc.__class__.__name__}: {exc}"
        log.error("LEDGER WRITE FAILED for export %s: %s", kind, exc, exc_info=True)
        return body
    try:
        _check_bulk(actor, case_number or "")
    except Exception:  # noqa: BLE001
        log.exception("bulk-export check failed")
    return result.get("out", body)


def export_threshold() -> tuple[int, int]:
    try:
        n = int(os.environ.get("EXPORT_ALERT_THRESHOLD", "10"))
    except ValueError:
        n = 10
    try:
        m = int(os.environ.get("EXPORT_ALERT_WINDOW_MINUTES", "10"))
    except ValueError:
        m = 10
    return max(1, n), max(1, m)


def _check_bulk(actor: str, case_number: str) -> None:
    """More than N exports/downloads by one user within M minutes -> ledger event + notification for admins/supervisors."""
    if not actor:
        return
    n, minutes = export_threshold()
    since = (datetime.now(timezone.utc) - timedelta(minutes=minutes)).strftime("%Y-%m-%dT%H:%M:%S.%fZ")
    with SessionLocal() as s:
        cnt = s.query(func.count(LedgerBlock.idx)).filter(LedgerBlock.actor == actor, LedgerBlock.event_type.in_(EXPORT_TYPES),
                                                          LedgerBlock.ts >= since).scalar() or 0
        if cnt <= n:
            return
        already = s.query(LedgerBlock.idx).filter(LedgerBlock.actor == actor, LedgerBlock.event_type == "export_anomaly",
                                                  LedgerBlock.ts >= since).first()
        if already:
            return
        recipients = [u.username for u in s.query(User).filter(User.active.is_(True), User.role.in_(("admin", "supervisor"))).all()]
    record("export_anomaly", actor=actor, case_number=case_number or "", ref=actor,
           payload={"exports_in_window": cnt, "window_minutes": minutes, "threshold": n})
    try:
        from backend.app.features.workflow import notify
        with SessionLocal() as s:
            for name in recipients:
                notify(s, name, "export_anomaly", f"Unusual export volume by {actor}",
                       f"{actor} exported or downloaded {cnt} items in {minutes} minutes (threshold {n}). See Data Integrity > Exports.",
                       case_number or None, None, commit=False)
            s.commit()
    except Exception:  # noqa: BLE001
        log.exception("could not create export-anomaly notifications")


# ================================================================================================ verification
def _load_block_dict(b: LedgerBlock) -> dict:
    return {"index": b.idx, "ts": b.ts, "event_type": b.event_type, "actor": b.actor, "case_number": b.case_number, "ref": b.ref,
            "payload": _loads(b.payload, {}), "payload_sha256": b.payload_sha256, "prev_hash": b.prev_hash, "block_hash": b.block_hash,
            "key_id": b.key_id, "signature": b.signature}


def _verify_sig(pub_hex: str, block_hash: str, sig_hex: str) -> bool:
    try:
        Ed25519PublicKey.from_public_bytes(bytes.fromhex(pub_hex)).verify(bytes.fromhex(sig_hex), signing_message(block_hash))
        return True
    except Exception:  # noqa: BLE001
        return False


def verify_chain(s, *, want_rows: bool = True) -> dict:
    """Walk the whole ledger once: links, hashes, signatures, key rotations, Merkle anchors, witness heads. Also collects
    the per-row hashes and retired cases that the data check needs."""
    problems: list[dict] = []
    recorded: dict = {}
    conflicts: list[dict] = []
    retired: set[str] = set()
    gaps: list[int] = []
    key_pub: str | None = None
    key_id: str | None = None
    keys_seen: list[dict] = []
    prev_hash, prev_idx, checked, anchors, rotations = ZERO, -1, 0, 0, 0
    leaves: list[str] = []
    leaf_from = 0
    head = None
    last_anchor_to = -1
    last_check = None
    chunk = 500
    last_idx = -1
    stop = False
    while not stop:
        blocks = (s.query(LedgerBlock).filter(LedgerBlock.idx > last_idx).order_by(LedgerBlock.idx).limit(chunk).all())
        if not blocks:
            break
        for b in blocks:
            last_idx = b.idx
            checked += 1
            add = lambda kind, msg, b=b: problems.append({"block": b.idx, "kind": kind, "message": msg, "case_number": b.case_number})  # noqa: E731
            if b.idx != prev_idx + 1:
                add("missing_block", f"Expected block {prev_idx + 1} but found {b.idx}: block(s) {prev_idx + 1}..{b.idx - 1} were deleted")
            if b.prev_hash != prev_hash:
                add("broken_link", f"Block {b.idx} does not follow block {prev_idx}: prev_hash mismatch (a block was altered, removed or re-ordered)")
            if sha256_hex(b.payload or "") != b.payload_sha256:
                add("payload_altered", f"Block {b.idx} payload does not match its recorded payload hash")
            expect = compute_block_hash(b.idx, b.ts, b.event_type, b.actor, b.case_number, b.ref, b.payload_sha256, b.prev_hash, b.key_id)
            if expect != b.block_hash:
                add("hash_mismatch", f"Block {b.idx} content does not match its block hash (block was edited)")
            payload = _loads(b.payload, {})
            if b.event_type == "genesis" and b.idx == 0:
                key_pub = payload.get("public_key")
                key_id = key_id_of(key_pub) if key_pub else None
                if key_id != payload.get("key_id"):
                    add("bad_genesis", "Genesis key id does not match its public key")
                keys_seen.append({"key_id": key_id, "from_block": 0, "fingerprint": fingerprint_of(key_pub) if key_pub else None})
            elif key_pub is None:
                add("no_genesis", "The ledger does not start with a genesis block")
            if key_pub:
                if b.key_id != key_id:
                    add("wrong_key", f"Block {b.idx} was signed with key {b.key_id}, expected {key_id} (swapped key)")
                if not _verify_sig(key_pub, b.block_hash, b.signature):
                    add("bad_signature", f"Signature of block {b.idx} is invalid")
            if b.event_type == "key_rotation" and key_pub:
                rotations += 1
                try:
                    np_hex = payload["new_public_key"]
                    ok = (payload["old_key_id"] == key_id and key_id_of(np_hex) == payload["new_key_id"])
                    Ed25519PublicKey.from_public_bytes(bytes.fromhex(np_hex)).verify(
                        bytes.fromhex(payload["proof_of_possession"]), f"{DOMAIN}|rotate|{key_id}|{payload['new_key_id']}|{payload['prev_hash_at_rotation']}".encode())
                    ok = ok and payload["prev_hash_at_rotation"] == b.prev_hash
                except Exception:  # noqa: BLE001
                    ok = False
                if ok:
                    key_pub, key_id = np_hex, payload["new_key_id"]
                    keys_seen.append({"key_id": key_id, "from_block": b.idx + 1, "fingerprint": fingerprint_of(np_hex)})
                else:
                    add("bad_rotation", f"Key rotation block {b.idx} is not valid (old key, new key or proof of possession)")
            if b.event_type == "anchor":
                anchors += 1
                try:
                    exp_root = merkle_root(leaves) if leaves else None
                except ValueError:
                    exp_root = None
                if payload.get("merkle_root") != exp_root or payload.get("from") != leaf_from or payload.get("to") != b.idx - 1 or payload.get("count") != len(leaves):
                    add("anchor_mismatch", f"Anchor {b.idx} does not match the blocks it covers ({leaf_from}..{b.idx - 1}): a covered block was changed or removed")
                last_anchor_to = b.idx - 1
                leaves, leaf_from = [], b.idx + 1
            else:
                leaves.append(b.block_hash)
            if b.event_type == "ledger_gap":
                gaps.append(b.idx)
            if b.event_type == "integrity_check":
                last_check = {"block": b.idx, "ts": b.ts, "actor": b.actor, **{k: payload.get(k) for k in ("verdict", "scope", "problems")}}
            if want_rows:
                for r in (payload.get("rows") or []) if isinstance(payload, dict) else []:
                    try:
                        k = (r["t"], str(r["k"]))
                    except (KeyError, TypeError):
                        continue
                    prev = recorded.get(k)
                    if prev is None:
                        recorded[k] = {"h": r["h"], "block": b.idx, "c": r.get("c") or b.case_number, "x": r.get("x"), "first": r["h"]}
                    else:
                        if prev["first"] != r["h"]:
                            conflicts.append({"t": k[0], "k": k[1], "first": prev["first"], "later": r["h"], "block": b.idx})
                        recorded[k] = {**prev, "h": r["h"], "block": b.idx}
                if b.event_type in ("case_purged", "demo_disabled"):
                    retired.update(str(c) for c in (payload.get("cases") or []))
                    # an attested removal ends the life of those rows: a later row that reuses the same key (SQLite reuses row ids) starts fresh
                    for x in (payload.get("removed") or []):
                        if isinstance(x, dict):
                            recorded.pop((str(x.get("t")), str(x.get("k"))), None)
            prev_hash, prev_idx, head = b.block_hash, b.idx, b
    if head is None:
        problems.append({"block": None, "kind": "empty", "message": "The ledger is empty", "case_number": ""})
    # ---- witness file: signed heads stored outside the database
    witness = {"file": str(_witness_path()), "entries": 0, "ok": True, "problems": []}
    wp = _witness_path()
    if wp.exists() and head is not None:
        for line in wp.read_text().splitlines():
            try:
                w = json.loads(line)
            except ValueError:
                witness["problems"].append("unreadable witness line")
                continue
            witness["entries"] += 1
            blk = s.get(LedgerBlock, w.get("index"))
            if w.get("index", 0) > head.idx:
                witness["problems"].append(f"witnessed head {w['index']} is newer than the database head {head.idx}: the newest blocks were removed")
                problems.append({"block": w["index"], "kind": "truncated", "message": f"Newest block(s) {head.idx + 1}..{w['index']} are missing compared with the witness file", "case_number": ""})
            elif blk is None or blk.block_hash != w.get("hash"):
                witness["problems"].append(f"witnessed block {w.get('index')} differs from the database")
                problems.append({"block": w.get("index"), "kind": "witness_mismatch", "message": f"Block {w.get('index')} differs from the copy in the witness file (chain was rewritten)", "case_number": ""})
    witness["ok"] = not witness["problems"]
    return {"ok": not problems, "blocks_checked": checked, "anchors_checked": anchors, "key_rotations": rotations, "keys": keys_seen,
            "problems": problems, "gaps": gaps, "head": ({"index": head.idx, "block_hash": head.block_hash, "ts": head.ts} if head else None),
            "unanchored_blocks": len(leaves), "witness": witness, "last_check": last_check,
            "_recorded": recorded, "_conflicts": conflicts, "_retired": retired}


def _file_check(db, case: str | None, recorded: dict) -> dict:
    q = db.query(Evidence)
    if case:
        q = q.filter(Evidence.case_number == case)
    problems, checked, ok, unreadable = [], 0, 0, 0
    seen = set()
    from backend.app.services.app_service import _abs_path
    for e in q.yield_per(200):
        checked += 1
        rec = recorded.get(("evidence", e.evidence_id))
        ledger_sha = (rec or {}).get("x")
        p = _abs_path(e.stored_path)
        base = {"evidence_id": e.evidence_id, "case_number": e.case_number, "recorded_sha256": e.sha256, "ledger_sha256": ledger_sha}
        if not p.is_file():
            problems.append({**base, "kind": "file_missing", "message": f"Evidence {e.evidence_id}: the stored file is missing"})
            continue
        try:
            cur = evidence_crypto.sha256_plain(p)
        except evidence_crypto.EvidenceCryptoError as exc:
            unreadable += 1
            problems.append({**base, "kind": "unreadable", "message": f"Evidence {e.evidence_id}: {exc}"})
            continue
        except OSError as exc:
            unreadable += 1
            problems.append({**base, "kind": "unreadable", "message": f"Evidence {e.evidence_id}: cannot read file ({exc.__class__.__name__})"})
            continue
        base["current_sha256"] = cur
        if cur != e.sha256:
            problems.append({**base, "kind": "file_altered", "message": f"Evidence {e.evidence_id}: the file on disk no longer matches its recorded SHA-256"})
        elif ledger_sha and cur != ledger_sha:
            problems.append({**base, "kind": "ledger_mismatch", "message": f"Evidence {e.evidence_id}: file matches the database but not the hash the ledger recorded at upload"})
        else:
            ok += 1
        seen.add(str(p))
    return {"checked": checked, "ok": ok, "problems": problems, "unreadable": unreadable}


def _row_check(db, case: str | None, chain: dict) -> dict:
    recorded, conflicts = chain["_recorded"], chain["_conflicts"]
    tables, problems = {}, []
    live_keys: dict[str, set] = {}
    for spec in tracked_specs():
        name = spec["name"]
        st = {"checked": 0, "mismatch": 0, "unrecorded": 0, "missing": 0, "mutable": spec["mutable"]}
        keys = set()
        try:
            for k, h, c, extra in _iter_rows(db, spec, case):
                st["checked"] += 1
                keys.add(k)
                rec = recorded.get((name, k))
                if rec is None:
                    st["unrecorded"] += 1
                    problems.append({"table": name, "key": k, "case_number": c, "kind": "unrecorded", "block": None,
                                     "message": f"{name} row {k} exists but no ledger event ever recorded it (created outside the normal channels)"})
                    continue
                expect = rec["h"] if spec["mutable"] else rec["first"]
                if h != expect:
                    st["mismatch"] += 1
                    problems.append({"table": name, "key": k, "case_number": c or rec["c"], "kind": "row_modified", "block": rec["block"],
                                     "message": f"{name} row {k} no longer matches the content recorded in block {rec['block']}"})
        except Exception as exc:  # noqa: BLE001
            st["error"] = exc.__class__.__name__
            log.exception("row check of %s failed", name)
        live_keys[name] = keys
        for (t, k), rec in recorded.items():
            if t != name or k in keys:
                continue
            if case and rec["c"] != case:
                continue
            st["missing"] += 1
            problems.append({"table": name, "key": k, "case_number": rec["c"], "kind": "row_missing", "block": rec["block"],
                             "message": f"{name} row {k} was recorded in block {rec['block']} but is no longer in the database (deleted)"})
        tables[name] = st
    mutable_map = {s["name"]: s["mutable"] for s in tracked_specs()}
    for c in conflicts:
        if c["t"] in mutable_map and not mutable_map[c["t"]]:
            problems.append({"table": c["t"], "key": c["k"], "case_number": "", "kind": "conflicting_records", "block": c["block"],
                             "message": f"{c['t']} row {c['k']} was recorded twice with different content (block {c['block']}); the row changed after it was first recorded"})
    return {"tables": tables, "problems": problems}


def run_verify(db, user: dict, case: str | None = None, check_files: bool = True) -> dict:
    """Full integrity check. Read-only. ``case`` narrows the data checks to one case (the chain is always checked whole)."""
    privileged = allowed(user.get("role"), "integrity") or allowed(user.get("role"), "audit")
    t0 = time.time()
    with SessionLocal() as s:
        chain = verify_chain(s)
    files = _file_check(db, case, chain["_recorded"]) if check_files else {"checked": 0, "ok": 0, "problems": [], "skipped": True, "unreadable": 0}
    rows = _row_check(db, case, chain)
    try:
        audit = dict(_CTX["ctx"].svc.verify_audit_chain(db))
    except Exception as exc:  # noqa: BLE001
        audit = {"verified": None, "reason": f"could not run: {exc.__class__.__name__}"}
    partial = []
    notes = []
    if chain["gaps"]:
        notes.append(f"{len(chain['gaps'])} ledger gap marker(s): some events could not be written to the ledger at the time and are only in the audit log")
    if files.get("unreadable"):
        partial.append(f"{files['unreadable']} encrypted or unreadable evidence file(s) could not be re-hashed")
    if not check_files:
        partial.append("evidence files were not re-hashed in this run")
    if audit.get("verified") is None:
        partial.append("the audit log chain could not be checked")
    if chain["unanchored_blocks"] > 0:
        pass  # informational only
    tamper = bool(chain["problems"]) or any(p["kind"] in ("file_missing", "file_altered", "ledger_mismatch") for p in files["problems"]) \
        or bool(rows["problems"]) or audit.get("verified") is False
    verdict = "TAMPERING DETECTED" if tamper else ("PARTIAL" if partial else "INTACT")
    if verdict == "PARTIAL" and not partial:
        verdict = "INTACT"
    n_problems = len(chain["problems"]) + len(files["problems"]) + len(rows["problems"]) + (1 if audit.get("verified") is False else 0)
    if verdict == "INTACT":
        text = ("Everything checked out. The ledger chain, every signature and Merkle anchor are valid; every evidence file still has the exact "
                "content it had when it was stored; and every recorded database row still matches what the ledger recorded.")
    elif verdict == "PARTIAL":
        text = "No tampering was found in what could be checked, but the check was incomplete: " + "; ".join(partial) + "."
    else:
        text = (f"{n_problems} problem(s) found. Something was changed, removed or added outside the normal recorded channels. "
                "Do not rely on the affected items until the cause is explained. The blocks, rows and files involved are listed below.")
    out = {"verdict": verdict, "explanation": text, "scope": case or "all cases", "checked_at": _now_ts(), "duration_ms": int((time.time() - t0) * 1000),
           "problems_total": n_problems, "partial_reasons": partial, "notes": notes,
           "chain": {k: v for k, v in chain.items() if not k.startswith("_")}, "files": files, "rows": rows, "audit_chain": audit,
           "limits": [CHAIN_NOTE, "The ledger proves that recorded data was not silently changed afterwards. It cannot prove that the data was true when it was entered, "
                      "or that nothing was changed before the ledger existed or outside the recorded tables."]}
    if not privileged:
        # A case-scoped view for investigators: only what concerns their case, no global detail.
        others = [p for p in out["chain"]["problems"] if p.get("case_number") != case]
        out["chain"]["problems"] = [p for p in out["chain"]["problems"] if p.get("case_number") == case]
        out["chain"]["other_problems"] = len(others)
        for k in ("witness", "keys", "gaps", "head", "last_check"):
            out["chain"].pop(k, None)
        out["audit_chain"] = {"verified": audit.get("verified")}
    return out


# ================================================================================================ views / masking
def _mask_payload(payload, role):
    """Auditors see structure and ids but not free text (there is none by design; belt and braces)."""
    if not isinstance(payload, dict) or allowed(role, "sensitive_read"):
        return payload
    out = {}
    for k, v in payload.items():
        if isinstance(v, str) and not (re.fullmatch(r"[0-9a-f]{16,64}", v) or _ID_OK.match(v)):
            out[k] = mask_identifier(v)
        else:
            out[k] = v
    return out


def _present(b: LedgerBlock, role: str, with_payload=True) -> dict:
    d = _load_block_dict(b)
    if with_payload:
        d["payload"] = _mask_payload(d["payload"], role)
        if b.event_type == "demo_disabled" and isinstance(d["payload"], dict):
            d["payload"] = {k: v for k, v in d["payload"].items() if k != "cases"}  # names of removed demo cases stay out of view
        rows = d["payload"].get("rows") if isinstance(d["payload"], dict) else None
        if isinstance(rows, list) and len(rows) > 20:
            d["payload"] = {**d["payload"], "rows": rows[:20], "rows_total": len(rows)}
    else:
        d.pop("payload")
    return d


def _scope_filter(q, db, user, hide_retired=True):
    """Blocks of cases the caller can access, plus system blocks that belong to no case.

    hide_retired: leave out blocks about demo cases that were switched off (they stay in the chain, so a full-chain
    bundle for offline verification must pass hide_retired=False or it would show gaps)."""
    gone = retired_demo_cases(db) if hide_retired else set()  # demo data is off: its blocks stay in the chain but are not listed
    if gone:
        q = q.filter(~LedgerBlock.case_number.in_(sorted(gone)))
    if user.get("role") == "admin" or allowed(user.get("role"), "audit"):
        return q
    cns = [c.case_number for c in visible_cases(db, user)]
    return q.filter((LedgerBlock.case_number == "") | (LedgerBlock.case_number.in_(cns or ["\0"])))


def _can_see_block(db, user, b: LedgerBlock) -> bool:
    return _scope_filter(db.query(LedgerBlock).filter(LedgerBlock.idx == b.idx), db, user).first() is not None


def build_bundle(db, user: dict, case: str | None) -> dict:
    """A signed JSON bundle that scripts/verify_ledger.py checks without the server."""
    with _State.lock:
        pend_needed = False
        with SessionLocal() as s0:
            pend_needed = _pending_range(s0) is not None
    if pend_needed:
        anchor_now(_actor(user))
    redact_payload = not allowed(user.get("role"), "sensitive_read") and user.get("role") != "admin"
    with SessionLocal() as s:
        head = _head(s)
        if head is None:
            raise HTTPException(404, "The ledger is empty")
        q = s.query(LedgerBlock).order_by(LedgerBlock.idx)
        if case:
            q = q.filter(LedgerBlock.case_number == case)
        else:
            q = _scope_filter(q, db, user, hide_retired=False)
        slice_blocks = q.all()
        full = not case and (user.get("role") == "admin" or allowed(user.get("role"), "audit"))
        key_blocks = s.query(LedgerBlock).filter(LedgerBlock.event_type.in_(("genesis", "key_rotation"))).order_by(LedgerBlock.idx).all()
        anchors = s.query(LedgerBlock).filter(LedgerBlock.event_type == "anchor").order_by(LedgerBlock.idx).all()
        anchor_payloads = [(a, _loads(a.payload, {})) for a in anchors]
        proofs, need_anchor = [], {}
        cache: dict = {}
        for b in slice_blocks:
            if b.event_type == "anchor":
                continue
            hit = next(((a, p) for a, p in anchor_payloads if p["from"] <= b.idx <= p["to"]), None)
            if not hit:
                continue
            a, p = hit
            if a.idx not in cache:
                cache[a.idx] = [x[0] for x in s.query(LedgerBlock.block_hash).filter(LedgerBlock.idx >= p["from"], LedgerBlock.idx <= p["to"],
                                                                                         LedgerBlock.event_type != "anchor").order_by(LedgerBlock.idx).all()]
                cache[("pos", a.idx)] = {h: i for i, h in enumerate(cache[a.idx])}
            leaves = cache[a.idx]
            pos = cache[("pos", a.idx)][b.block_hash]
            proofs.append({"index": b.idx, "anchor_index": a.idx, "leaf_position": pos, "leaf_count": len(leaves), "path": merkle_proof(leaves, pos)})
            need_anchor[a.idx] = a
        shown_anchors = [a for a in anchors if full or a.idx in need_anchor]
        priv = _load_or_create_key()

        def blk(b, keep_payload=True):
            d = _load_block_dict(b)
            if redact_payload and not (b.event_type in ("genesis", "key_rotation", "anchor")) and keep_payload:
                d["payload"] = None
                d["payload_redacted"] = True
            return d
        bundle = {
            "format": BUNDLE_FORMAT, "generated_at": _now_ts(), "generated_by": _actor(user),
            "scope": {"case_number": case or None, "full_chain": full},
            "public_key": {"key_id": key_id_of(_pub_hex(priv)), "public_key": _pub_hex(priv), "fingerprint": fingerprint_of(_pub_hex(priv))},
            "key_chain": [_load_block_dict(b) for b in key_blocks],
            "blocks": [blk(b) for b in slice_blocks],
            "anchors": [_load_block_dict(a) for a in shown_anchors],
            "proofs": proofs,
            "head": {"index": head.idx, "block_hash": head.block_hash, "ts": head.ts, "chain_length": head.idx + 1},
            "notes": [CHAIN_NOTE, "A case slice can prove that its blocks were not edited and were signed by the server key; only a full-chain bundle can show that no block is missing."],
        }
        digest = sha256_hex(canon(bundle))
        bundle["bundle_signature"] = {"alg": "Ed25519", "key_id": key_id_of(_pub_hex(priv)), "over": "sha256 of canonical JSON of this bundle without bundle_signature",
                                      "digest": digest, "signature": priv.sign(f"{DOMAIN}|bundle|".encode() + bytes.fromhex(digest)).hex()}
        return bundle


def head_statement(s) -> dict:
    head = _head(s)
    if head is None:
        raise HTTPException(404, "The ledger is empty")
    priv = _load_or_create_key()
    stmt = {"index": head.idx, "block_hash": head.block_hash, "ts": head.ts, "key_id": head.key_id, "chain_length": head.idx + 1}
    stmt["signature"] = priv.sign(f"{DOMAIN}|head|{head.idx}|{head.block_hash}".encode()).hex()
    stmt["print_line"] = f"DCN ledger head | block {head.idx} | {head.block_hash} | {head.ts}"
    stmt["note"] = "Print or store this line somewhere the server cannot reach. Later, a chain whose block at this index has a different hash - or is shorter - has been rewritten or truncated."
    return stmt


# ================================================================================================ API
_CTX: dict = {}


def build(ctx):
    _CTX["ctx"] = ctx
    r = APIRouter(prefix="/integrity", tags=["integrity"])
    get_db = ctx.get_db
    priv_dep = require_any_perm("integrity", "audit")
    read_dep = require_any_perm("integrity", "audit", "read")

    def _is_priv(user) -> bool:
        return allowed(user.get("role"), "integrity") or allowed(user.get("role"), "audit")

    def _scope_case(db, user, case_number):
        """Privileged users may look globally; everyone else must name a case they can access."""
        if case_number:
            try:
                return ctx.scope(db, user, case_number)
            except HTTPException as exc:
                # the ledger outlives purged cases: administrators and auditors may still look up what was recorded for them
                if exc.status_code == 404 and _is_priv(user) and (user.get("role") == "admin" or allowed(user.get("role"), "audit")):
                    with SessionLocal() as s:
                        if s.query(LedgerBlock.idx).filter(LedgerBlock.case_number == case_number).first():
                            return case_number
                raise
        if not _is_priv(user):
            raise HTTPException(400, "case_number is required")
        return None

    @r.get("/status")
    def status(case_number: str | None = None, db=Depends(get_db), user=Depends(read_dep)):
        cn = _scope_case(db, user, (case_number or "").strip())
        with SessionLocal() as s:
            head = _head(s)
            n = (head.idx + 1) if head else 0
            last_anchor = s.query(LedgerBlock).filter(LedgerBlock.event_type == "anchor").order_by(LedgerBlock.idx.desc()).first()
            pend = _pending_range(s)
            out = {"chain_length": n, "case_number": cn,
                   "case_blocks": s.query(func.count(LedgerBlock.idx)).filter(LedgerBlock.case_number == cn).scalar() if cn else None,
                   "unanchored_blocks": (pend[1] - pend[0] + 1) if pend else 0,
                   "last_anchor": ({"index": last_anchor.idx, "ts": last_anchor.ts, "merkle_root": _loads(last_anchor.payload, {}).get("merkle_root"),
                                    "covers": [_loads(last_anchor.payload, {}).get("from"), _loads(last_anchor.payload, {}).get("to")]} if last_anchor else None),
                   "encryption_at_rest": {"enabled": evidence_crypto.enabled(), "key_fingerprint": evidence_crypto.key_fingerprint()},
                   "anchor_every": _anchor_every(), "export_alert": {"threshold": export_threshold()[0], "window_minutes": export_threshold()[1]}}
            if _is_priv(user):
                out["head"] = ({"index": head.idx, "block_hash": head.block_hash, "ts": head.ts} if head else None)
                keyblk = (s.query(LedgerBlock).filter(LedgerBlock.event_type.in_(("genesis", "key_rotation"))).order_by(LedgerBlock.idx.desc()).first())
                if keyblk:
                    p = _loads(keyblk.payload, {})
                    pub = p.get("public_key") or p.get("new_public_key")
                    out["key"] = {"key_id": key_id_of(pub), "fingerprint": fingerprint_of(pub), "public_key": pub,
                                  "rotations": s.query(func.count(LedgerBlock.idx)).filter(LedgerBlock.event_type == "key_rotation").scalar()}
                last_chk = s.query(LedgerBlock).filter(LedgerBlock.event_type == "integrity_check", LedgerBlock.case_number == "").order_by(LedgerBlock.idx.desc()).first()
                if last_chk:
                    p = _loads(last_chk.payload, {})
                    out["last_verification"] = {"block": last_chk.idx, "ts": last_chk.ts, "actor": last_chk.actor, "verdict": p.get("verdict"), "problems": p.get("problems")}
                out["ledger_gaps"] = s.query(func.count(LedgerBlock.idx)).filter(LedgerBlock.event_type == "ledger_gap").scalar()
                out["pending_failures"] = _State.missed
                out["witness_file"] = _witness_path().exists()
        if cn and not _is_priv(user):
            lv = _State.verify_cache.get(cn)
            if lv:
                out["last_verification"] = lv
        return out

    @r.get("/head")
    def head(db=Depends(get_db), user=Depends(priv_dep)):
        with SessionLocal() as s:
            stmt = head_statement(s)
            hb = s.get(LedgerBlock, stmt["index"])
            priv = _load_or_create_key()
            if hb is not None:
                _witness(hb, priv)
            return stmt

    @r.post("/verify")
    def verify(case_number: str | None = None, files: bool = True, db=Depends(get_db), user=Depends(read_dep)):
        cn = _scope_case(db, user, (case_number or "").strip())
        res = run_verify(db, user, cn, check_files=files)
        if _is_priv(user) and cn is None:
            record("integrity_check", actor=_actor(user), payload={"verdict": res["verdict"], "scope": "all", "problems": res["problems_total"],
                                                                     "blocks_checked": res["chain"]["blocks_checked"], "files_checked": res["files"]["checked"]})
        elif cn:
            _State.verify_cache[cn] = {"ts": res["checked_at"], "verdict": res["verdict"], "problems": res["problems_total"], "actor": _actor(user)}
        return res

    @r.get("/blocks")
    def blocks(case_number: str | None = None, event_type: str | None = Query(None, max_length=60), limit: int = Query(50, ge=1, le=500),
               after: int | None = Query(None, ge=0), desc: bool = False, db=Depends(get_db), user=Depends(read_dep)):
        cn = _scope_case(db, user, (case_number or "").strip())
        with SessionLocal() as s:
            q = s.query(LedgerBlock)
            if cn:
                q = q.filter(LedgerBlock.case_number == cn)
            else:
                q = _scope_filter(q, db, user)
            if event_type:
                q = q.filter(LedgerBlock.event_type == event_type)
            if after is not None:
                q = q.filter(LedgerBlock.idx < after) if desc else q.filter(LedgerBlock.idx > after)
            rows = q.order_by(LedgerBlock.idx.desc() if desc else LedgerBlock.idx).limit(limit + 1).all()
            more = len(rows) > limit
            rows = rows[:limit]
            return {"blocks": [_present(b, user["role"], with_payload=False) for b in rows], "next": (rows[-1].idx if rows and more else None),
                    "chain_length": (_head(s).idx + 1) if _head(s) else 0}

    @r.get("/blocks/{index}")
    def block(index: int, db=Depends(get_db), user=Depends(read_dep)):
        with SessionLocal() as s:
            b = s.get(LedgerBlock, index)
            if b is None or not _can_see_block(db, user, b):
                raise HTTPException(404, "Block not found")
            d = _present(b, user["role"])
            nxt = s.get(LedgerBlock, index + 1)
            d["next_block_prev_hash_matches"] = (nxt.prev_hash == b.block_hash) if nxt else None
            d["explanation"] = {"block": "One recorded event.", "hash": "SHA-256 fingerprint of the block; any change to the block changes it.",
                                "prev_hash": "Fingerprint of the previous block; this is what chains the blocks together.",
                                "signature": "The server's Ed25519 signature over the block hash."}
            return d

    @r.get("/proof/{index}")
    def proof(index: int, db=Depends(get_db), user=Depends(read_dep)):
        with SessionLocal() as s:
            b = s.get(LedgerBlock, index)
            if b is None or not _can_see_block(db, user, b):
                raise HTTPException(404, "Block not found")
            if b.event_type == "anchor":
                raise HTTPException(400, "Anchor blocks are proofs themselves; they are covered by the hash chain")
            anchors = s.query(LedgerBlock).filter(LedgerBlock.event_type == "anchor", LedgerBlock.idx > index).order_by(LedgerBlock.idx).all()
            for a in anchors:
                p = _loads(a.payload, {})
                if p["from"] <= index <= p["to"]:
                    leaves = [x[0] for x in s.query(LedgerBlock.block_hash).filter(LedgerBlock.idx >= p["from"], LedgerBlock.idx <= p["to"],
                                                                                    LedgerBlock.event_type != "anchor").order_by(LedgerBlock.idx).all()]
                    pos = leaves.index(b.block_hash)
                    path = merkle_proof(leaves, pos)
                    return {"index": index, "block_hash": b.block_hash, "anchor_index": a.idx, "anchor_hash": a.block_hash, "merkle_root": p["merkle_root"],
                            "leaf_position": pos, "leaf_count": len(leaves), "path": path, "valid": merkle_verify(b.block_hash, path, p["merkle_root"]),
                            "explanation": "Hash the block hash as a leaf, then combine it with each sibling in order; the result must equal the anchor's Merkle root."}
            raise HTTPException(404, "This block is not covered by an anchor yet. An administrator or supervisor can create one (POST /integrity/anchor); it is also created automatically.")

    @r.get("/bundle")
    def bundle(case_number: str | None = None, db=Depends(get_db), user=Depends(priv_dep)):
        cn = _scope_case(db, user, (case_number or "").strip())
        b = build_bundle(db, user, cn)
        record("export_generated", actor=_actor(user), case_number=cn or "", ref="ledger_bundle",
               payload={"kind": "ledger_bundle", "format": "json", "output_sha256": b["bundle_signature"]["digest"], "blocks": len(b["blocks"])})
        return b

    @r.post("/anchor")
    def anchor(db=Depends(get_db), user=Depends(priv_dep)):
        if not (allowed(user["role"], "write") or allowed(user["role"], "manage_users")):
            raise HTTPException(403, "Insufficient permissions")
        blk = anchor_now(_actor(user))
        if blk is None:
            return {"created": False, "message": "Nothing to anchor: every block is already covered."}
        return {"created": True, "index": blk.idx, "block_hash": blk.block_hash, "merkle_root": _loads(blk.payload, {}).get("merkle_root")}

    @r.post("/rotate-key")
    def rotate(db=Depends(get_db), user=Depends(require_any_perm("manage_users"))):
        try:
            out = rotate_key(_actor(user))
        except LedgerError as exc:
            raise HTTPException(409, str(exc))
        ctx.audit(db, user["sub"], "integrity.key_rotated", out)
        return out

    @r.post("/baseline")
    def baseline(db=Depends(get_db), user=Depends(require_any_perm("manage_users"))):
        """Accept every not-yet-recorded row as the starting point (for data created by scripts or imported before the
        ledger existed). This is an explicit, logged decision: it cannot tell whether such rows were legitimate."""
        rows = unrecorded_rows(db, None)
        idx = record("baseline_accepted", actor=_actor(user), payload={"rows": [], "rows_accepted": len(rows), "note": "administrator accepted unrecorded rows as baseline"}, rows=rows) if rows else None
        ctx.audit(db, user["sub"], "integrity.baseline_accepted", {"rows": len(rows), "block": idx})
        return {"rows_recorded": len(rows), "block": idx}

    @r.get("/exports")
    def exports(case_number: str | None = None, actor: str | None = Query(None, max_length=100), kind: str | None = Query(None, max_length=60),
                watermark: str | None = Query(None, max_length=40), sha256: str | None = Query(None, max_length=64),
                limit: int = Query(100, ge=1, le=500), db=Depends(get_db), user=Depends(priv_dep)):
        cn = _scope_case(db, user, (case_number or "").strip())
        with SessionLocal() as s:
            q = s.query(LedgerBlock).filter(LedgerBlock.event_type.in_(EXPORT_TYPES + ("export_anomaly", "backup_created")))
            q = q.filter(LedgerBlock.case_number == cn) if cn else _scope_filter(q, db, user)
            if actor:
                q = q.filter(LedgerBlock.actor == actor)
            if kind:
                q = q.filter(LedgerBlock.ref == kind)
            rows = q.order_by(LedgerBlock.idx.desc()).limit(2000 if (watermark or sha256) else limit).all()
            out = []
            for b in rows:
                p = _loads(b.payload, {})
                if watermark and p.get("watermark_id") != watermark.strip():
                    continue
                if sha256 and sha256.lower() not in (p.get("output_sha256"), p.get("sha256"), p.get("source_sha256")):
                    continue
                out.append({"index": b.idx, "ts": b.ts, "event_type": b.event_type, "actor": b.actor, "case_number": b.case_number, "kind": p.get("kind") or b.ref,
                            "format": p.get("format"), "bytes": p.get("bytes"), "watermark_id": p.get("watermark_id"), "output_sha256": p.get("output_sha256") or p.get("sha256"),
                            "ledger_ref": f"L-{b.idx}"})
            n, minutes = export_threshold()
            since = (datetime.now(timezone.utc) - timedelta(minutes=minutes)).strftime("%Y-%m-%dT%H:%M:%S.%fZ")
            per_user = dict(s.query(LedgerBlock.actor, func.count(LedgerBlock.idx)).filter(LedgerBlock.event_type.in_(EXPORT_TYPES), LedgerBlock.ts >= since)
                            .group_by(LedgerBlock.actor).all())
            return {"exports": out[:limit], "recent_by_user": per_user, "threshold": n, "window_minutes": minutes,
                    "note": "Each export carries a visible watermark (who, when, ledger reference). Given a leaked copy, search by its watermark id or file hash."}

    @r.get("/schema")
    def schema(user=Depends(read_dep)):
        return {"concepts": {"block": "One recorded event: who, what, when, plus fingerprints (never file contents or personal data).",
                             "hash": "A fixed-length fingerprint of data. Change one character and the fingerprint changes completely.",
                             "chain": "Each block stores the previous block's hash, so altering, deleting or re-ordering any block breaks every later link.",
                             "anchor": "Every N blocks the ledger records one combined fingerprint (Merkle root) of the batch, so a single event can be proven with a short proof.",
                             "signature": "The server signs each block with a private key that never leaves the server; anyone with the public key can check it."},
                "verify_offline": "python scripts/verify_ledger.py bundle.json [--expect-head HASH] [--trust-key FINGERPRINT]"}

    return r


def on_startup(ctx):
    """Subscribe to the event bus and make sure the ledger has its genesis block (plus a baseline of existing data)."""
    hk = ctx.hooks
    hk.on("document_extracted", _h_document_extracted)
    hk.on("document_reviewed", _h_document_reviewed)
    hk.on("document_reprocessed", _h_document_reprocessed)
    hk.on("table_ingested", _h_table_ingested)
    hk.on("evidence_uploaded", _h_evidence_uploaded)
    hk.on("evidence_downloaded", _h_evidence_downloaded)
    hk.on("evidence_verified", _h_evidence_verified)
    hk.on("case_created", _h_case_created)
    hk.on("entity_created", _h_entity_created)
    hk.on("relationship_created", _h_relationship_created)
    hk.on("export_generated", _h_export_generated)
    for ev in ("social_imported", "social_item_flagged", "surveillance_entry_added", "surveillance_entry_amended", "surveillance_report_generated"):
        hk.on(ev, _h_generic(ev))
    for ev in ("surveillance_operation_created", "surveillance_operation_closed"):
        hk.on(ev, _h_generic(ev))
    for ev in ("backup_created", "backup_downloaded", "role_changed", "user_created", "user_deleted", "mfa_enabled", "mfa_disabled", "login_failed_burst"):
        hk.on(ev, _h_simple(ev))
    hk.on("demo_enabled", _h_demo_enabled)
    hk.on("demo_disabled", _h_demo_disabled)
    hk.on("case_purged", _h_case_purged)
    try:
        with SessionLocal() as s:
            empty = _head(s) is None
        if empty:
            with SessionLocal() as db:
                sweep(db, None, "system", "baseline_accepted", "startup",
                      {"note": "ledger started: rows that already existed were recorded as the baseline (their history before this point is not covered)"})
            if _head_missing():
                record("ledger_started", actor="system", payload={"version": VERSION})
    except Exception:  # noqa: BLE001
        log.exception("ledger: start-up baseline failed")


def _head_missing() -> bool:
    with SessionLocal() as s:
        return _head(s) is None
