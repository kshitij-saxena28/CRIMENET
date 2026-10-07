"""Governance & oversight: dashboard, redaction, hold/purge, backup/restore, access history, legal packs, import, inbox."""
import io
import json
import os
import re
import subprocess
import sys
import uuid
from pathlib import Path

import pytest

from backend.app.features import governance as G

ROOT = Path(__file__).resolve().parents[2]


def _cn(prefix="GOV"):
    return f"{prefix}-{uuid.uuid4().hex[:8].upper()}"


def _case(client, h, cn=None, title="Governance test"):
    cn = cn or _cn()
    assert client.post("/cases", json={"case_number": cn, "title": title}, headers=h).status_code == 200
    return cn


def _ent(client, h, cn, eid, name, typ, **attrs):
    r = client.post("/entities", json={"external_id": eid, "name": name, "entity_type": typ, "attributes": attrs, "case_number": cn}, headers=h)
    assert r.status_code == 200, r.text


def _rel(client, h, cn, a, b, rel="CALLED"):
    r = client.post("/relationships", json={"source_id": a, "target_id": b, "relation_type": rel, "case_number": cn, "verification_state": "verified"}, headers=h)
    assert r.status_code == 200, r.text


def _seed_case(client, h):
    cn = _case(client, h)
    p = uuid.uuid4().hex[:6]
    _ent(client, h, cn, f"P1-{p}", "Rahul Sharma", "PERSON", note="Rahul Sharma called 9876543210 from Aadhaar 2345 6789 0123")
    _ent(client, h, cn, f"P2-{p}", "Priya Verma", "PERSON")
    _ent(client, h, cn, "9876543210" + p[:2], "9876543210", "PHONE")           # id embeds the phone number
    _ent(client, h, cn, f"A-{p}", "ACCT-445566", "ACCOUNT", account_number="445566778899")
    _ent(client, h, cn, f"V-{p}", "UP14 AB 4521", "VEHICLE")
    _rel(client, h, cn, f"P1-{p}", "9876543210" + p[:2])
    _rel(client, h, cn, f"P2-{p}", "9876543210" + p[:2])
    _rel(client, h, cn, f"P1-{p}", f"P2-{p}", "ASSOCIATED_WITH")
    assert client.post("/events", json={"event_id": f"E-{p}", "event_type": "COMMUNICATION", "entity_id": f"P1-{p}", "related_entity_id": f"P2-{p}",
                                        "event_time": "2026-03-01T10:00:00", "case_number": cn, "metadata": {"remark": "call from +91 98765 43210"}}, headers=h).status_code == 200
    return cn, p


# ------------------------------------------------------------------------------------------ redaction engine
def test_text_redactor_patterns():
    txt = ("Aadhaar 2345 6789 0123 and 234567890123, PAN ABCDE1234F, phone +91 98765 43210 / 9876543210 / 098765-43210, "
           "mail a.b@example.co.in, plate UP14 AB 4521 and MH-12-AB-1234, acct ACCT-445566 and 445566778899.")
    out = G.redact_text(txt, mode="mask")
    for raw in ("2345 6789 0123", "234567890123", "ABCDE1234F", "98765 43210", "9876543210", "a.b@example.co.in", "UP14 AB 4521", "MH-12-AB-1234", "445566778899"):
        assert raw not in out, raw
    assert out.count("XXXX") >= 6 and "0123" in out            # masks keep the last four characters
    rem = G.redact_text(txt, mode="remove")
    assert rem.count("[REMOVED]") >= 10
    # dates, short numbers and case numbers are left alone
    plain = "Case CASE-2026-101 dated 12/08/2026, 42 items, section 379"
    assert G.redact_text(plain, mode="remove") == plain


def test_pseudonyms_are_consistent_and_salted():
    text = "call 9876543210 then +91 98765 43210 then 9123456789"
    a = G.redact_text(text, mode="pseudonymise")
    labels = re.findall(r"PHONE-\d+", a)
    assert labels[0] == labels[1] and labels[0] != labels[2]      # same number, same pseudonym; different number differs


def test_redacted_export_preserves_link_structure(client, login):
    h = login("admin")
    cn, p = _seed_case(client, h)
    body = {"case_number": cn, "format": "json", "mode": "pseudonymise",
            "redact": {"phones": True, "accounts": True, "emails": True, "vehicles": True, "national_ids": True, "names": True, "addresses": False}}
    r = client.post("/governance/redacted-export", json=body, headers=h)
    assert r.status_code == 200, r.text
    raw = r.text
    for secret in ("9876543210", "98765 43210", "Rahul Sharma", "Priya Verma", "445566778899", "UP14 AB 4521", "2345 6789 0123"):
        assert secret not in raw, secret
    data = r.json()
    ids = {e["external_id"] for e in data["entities"]}
    assert len(data["entities"]) == 5
    assert all(rel["source"] in ids and rel["target"] in ids for rel in data["relationships"])      # no dangling links
    assert len(data["relationships"]) == 3
    # the phone node is still the shared neighbour of the two people
    phone = next(e for e in data["entities"] if e["type"] == "PHONE")
    assert phone["name"].startswith("PHONE-")
    people = [e for e in data["entities"] if e["type"] == "PERSON"]
    assert {x["name"] for x in people} == {"PERSON-001", "PERSON-002"}
    touching = [rel for rel in data["relationships"] if phone["external_id"] in (rel["source"], rel["target"])]
    assert len(touching) == 2
    ev = data["events"][0]
    assert ev["entity_id"] in ids and "98765" not in json.dumps(ev)
    log_ = data["redaction_log"]
    assert log_["replacements_by_category"]["phones"] >= 2 and log_["mode"] == "pseudonymise"
    # audit entry with counts, and no PII in it
    audit = client.get("/audit", params={"case_number": cn}, headers=h).json()["events"]
    entry = [a for a in audit if a["action"] == "governance.redacted_export"]
    assert entry and "9876543210" not in entry[0]["details"]


def test_redacted_export_formats_and_modes(client, login):
    h = login("admin")
    cn, _ = _seed_case(client, h)
    base = {"case_number": cn, "redact": {"phones": True, "names": True}}
    m = client.post("/governance/redacted-export", json={**base, "mode": "mask", "format": "json"}, headers=h).json()
    assert any(e["name"].startswith("R*** S***") for e in m["entities"])
    rm = client.post("/governance/redacted-export", json={**base, "mode": "remove", "format": "json"}, headers=h).text
    assert "Rahul" not in rm and "[REMOVED]" in rm
    c = client.post("/governance/redacted-export", json={**base, "mode": "pseudonymise", "format": "csv"}, headers=h)
    assert c.status_code == 200 and "text/csv" in c.headers["content-type"] and "redaction_log" in c.text and "9876543210" not in c.text
    d = client.post("/governance/redacted-export", json={**base, "mode": "pseudonymise", "format": "docx"}, headers=h)
    from docx import Document
    doc = Document(io.BytesIO(d.content))
    assert any("Redaction log" in p.text for p in doc.paragraphs)
    pv = client.post("/governance/redacted-export", json={**base, "preview": True}, headers=h).json()
    assert pv["preview"] and pv["redaction_log"]["replacements_by_category"]["names"] >= 2


def test_export_validation_and_authorization(client, login, make_user):
    h = login("admin")
    cn, _ = _seed_case(client, h)
    ok = {"case_number": cn, "redact": {"phones": True}}
    assert client.post("/governance/redacted-export", json={**ok, "format": "exe"}, headers=h).status_code == 422
    assert client.post("/governance/redacted-export", json={**ok, "redact": {k: False for k in ("phones", "accounts", "emails", "vehicles", "national_ids")}}, headers=h).status_code == 400
    assert client.post("/governance/redacted-export", json={**ok, "redact": {"free_text_patterns": ["re:(a+)+$"]}}, headers=h).status_code == 400
    assert client.post("/governance/redacted-export", json=ok).status_code == 401
    _, _, sub = make_user("auditor")
    assert client.post("/governance/redacted-export", json=ok, headers=sub).status_code == 403
    _, _, inv = make_user("supervisor")                      # restricted case they cannot see
    assert client.post("/governance/redacted-export", json=ok, headers=inv).status_code == 404
    ft = client.post("/governance/redacted-export", json={**ok, "redact": {"free_text_patterns": ["Shipra Mall"]}}, headers=h)
    assert ft.status_code == 200


# ------------------------------------------------------------------------------------------ hold / retention / purge
def _upload(client, h, cn, name="note.txt", data=None):
    r = client.post("/evidence/upload", files={"file": (name, data or f"evidence {uuid.uuid4()}".encode())}, data={"case_number": cn}, headers=h)
    assert r.status_code == 200, r.text
    return r.json()[0]


def test_hold_blocks_purge_and_guards(client, login, make_user):
    h = login("admin")
    cn = _case(client, h)
    ev = _upload(client, h, cn)
    path = G.A._abs_path(G.A.EVIDENCE_DIR / (ev["evidence_id"] + ".txt"))
    assert path.exists()
    # only admins may hold / purge
    _, _, sup = make_user("supervisor")
    assert client.put(f"/governance/cases/{cn}/hold", json={"legal_hold": True, "reason": "court order"}, headers=sup).status_code == 403
    assert client.post(f"/governance/cases/{cn}/purge", json={"confirm_case_number": cn, "reason": "retention expired long ago", "dry_run": True}, headers=sup).status_code == 403
    assert client.put(f"/governance/cases/{cn}/hold", json={"legal_hold": True, "reason": "x"}, headers=h).status_code == 422
    assert client.put(f"/governance/cases/{cn}/hold", json={"legal_hold": True, "reason": "Court order 12/2026"}, headers=h).status_code == 200
    body = {"confirm_case_number": cn, "reason": "Retention period elapsed, approved by SP", "dry_run": False}
    assert client.post(f"/governance/cases/{cn}/purge", json=body, headers=h).status_code == 409     # hold blocks
    assert path.exists()
    rows = client.get("/governance/retention", headers=h).json()["cases"]
    assert next(x for x in rows if x["case_number"] == cn)["legal_hold"] is True
    assert client.put(f"/governance/cases/{cn}/hold", json={"legal_hold": False, "reason": "Court released"}, headers=h).status_code == 200
    # confirmation mismatch and short reason
    assert client.post(f"/governance/cases/{cn}/purge", json={**body, "confirm_case_number": cn.lower()}, headers=h).status_code == 400
    assert client.post(f"/governance/cases/{cn}/purge", json={**body, "reason": "too short"}, headers=h).status_code == 422
    assert path.exists() and client.get("/cases", headers=h).json()
    actions = [a["action"] for a in client.get("/audit", params={"case_number": cn}, headers=h).json()["events"]]
    assert "governance.legal_hold_placed" in actions and "governance.legal_hold_released" in actions


def test_retention_settings_and_flags(client, login):
    h = login("admin")
    cn = _case(client, h)
    assert client.put(f"/governance/cases/{cn}/retention", json={"retention_days": 0}, headers=h).status_code == 422
    assert client.put(f"/governance/cases/{cn}/retention", json={"retention_days": 30, "classification": "top"}, headers=h).status_code == 400
    assert client.put(f"/governance/cases/{cn}/retention", json={"retention_days": 30, "classification": "confidential"}, headers=h).json()["classification"] == "CONFIDENTIAL"
    row = next(x for x in client.get("/governance/retention", headers=h).json()["cases"] if x["case_number"] == cn)
    assert row["retention_days"] == 30 and row["past_retention"] is False and row["retention_due"]
    from backend.app.db.database import SessionLocal
    from backend.app.models.models import Case
    from datetime import datetime, timedelta
    with SessionLocal() as db:
        db.query(Case).filter_by(case_number=cn).update({"created_at": datetime.utcnow() - timedelta(days=90)}); db.commit()
    row = next(x for x in client.get("/governance/retention", headers=h).json()["cases"] if x["case_number"] == cn)
    assert row["past_retention"] and row["purge_review_candidate"]


def test_purge_dry_run_then_real_keeps_audit_chain(client, login):
    h = login("admin")
    cn, p = _seed_case(client, h)
    other = _case(client, h)
    ev = _upload(client, h, cn, data=b"to be purged " + uuid.uuid4().hex.encode())
    keep = _upload(client, h, other, data=b"unrelated case file " + uuid.uuid4().hex.encode())
    client.post("/notes", json={"case_number": cn, "body": "a note"}, headers=h)
    f = G.A.EVIDENCE_DIR / (ev["evidence_id"] + ".txt"); kf = G.A.EVIDENCE_DIR / (keep["evidence_id"] + ".txt")
    assert f.exists() and kf.exists()
    body = {"confirm_case_number": cn, "reason": "Retention expired; approved by SP (ref 55/2026)"}
    dry = client.post(f"/governance/cases/{cn}/purge", json=body, headers=h)                  # dry_run defaults to true
    assert dry.status_code == 200 and dry.json()["dry_run"] is True
    s = dry.json()["summary"]
    assert s["files_deleted"] == 1 and s["entities_deleted"] == 5 and s["relationships_deleted"] == 3 and s["events_deleted"] == 1
    assert f.exists()                                                                            # dry run deleted nothing
    assert len(client.get("/entities", params={"case_number": cn}, headers=h).json()) == 5
    before = client.get("/audit/verify", headers=h).json()
    assert before["verified"]
    real = client.post(f"/governance/cases/{cn}/purge", json={**body, "dry_run": False}, headers=h)
    assert real.status_code == 200, real.text
    out = real.json()
    assert out["audit_chain_after"]["verified"] is True and out["summary"]["files_deleted"] == 1
    assert not f.exists() and kf.exists()                                                        # only this case's file removed
    assert client.get(f"/cases/{cn}/workspace", headers=h).status_code == 404
    assert client.get(f"/cases/{other}/workspace", headers=h).status_code == 200
    after = client.get("/audit/verify", headers=h).json()
    assert after["verified"] and after["checked"] > before["checked"]                            # chain grew, still verifies
    events = client.get("/audit", headers=h).json()["events"]
    tomb = next(a for a in events if a["action"] == "governance.case_purged")
    d = json.loads(tomb["details"])
    assert d["tombstone"] and d["purged_files"][0]["sha256"] == ev["sha256"] and d["case_number"] == cn
    assert any(a["action"] == "governance.purge_dry_run" for a in events)


def test_purge_never_touches_files_outside_evidence_dir(client, login, tmp_path):
    h = login("admin")
    cn = _case(client, h)
    victim = tmp_path / "precious.txt"; victim.write_text("do not delete")
    from backend.app.db.database import SessionLocal
    from backend.app.models.models import Evidence
    with SessionLocal() as db:
        db.add(Evidence(evidence_id="EV-OUTSIDE-" + uuid.uuid4().hex[:6], filename="x.txt", sha256="0" * 64, stored_path=str(victim), case_number=cn, uploaded_by="t"))
        db.add(Evidence(evidence_id="EV-TRAV-" + uuid.uuid4().hex[:6], filename="y.txt", sha256="0" * 64, stored_path=str(G.A.EVIDENCE_DIR / ".." / ".." / "precious.txt"), case_number=cn, uploaded_by="t"))
        db.commit()
    r = client.post(f"/governance/cases/{cn}/purge", json={"confirm_case_number": cn, "reason": "cleanup of test data set", "dry_run": False}, headers=h)
    assert r.status_code == 200, r.text
    assert victim.exists() and victim.read_text() == "do not delete"
    assert all(f["action"].startswith("skipped") for f in r.json()["files"])


# ------------------------------------------------------------------------------------------ backup / restore
PASS = "correct horse battery staple"


def test_backup_create_verify_tamper_and_restore(client, login, make_user, tmp_path):
    h = login("admin")
    cn = _case(client, h)
    ev = _upload(client, h, cn, data=b"backup me " + uuid.uuid4().hex.encode())
    _, _, inv = make_user("supervisor")
    assert client.post("/governance/backup", json={"passphrase": PASS}, headers=inv).status_code == 403
    assert client.post("/governance/backup", json={"passphrase": "short"}, headers=h).status_code == 422
    r = client.post("/governance/backup", json={"passphrase": PASS}, headers=h)
    assert r.status_code == 200, r.text
    name = r.json()["name"]
    assert name in [b["name"] for b in client.get("/governance/backups", headers=h).json()["backups"]]
    v = client.post(f"/governance/backups/{name}/verify", json={"passphrase": PASS}, headers=h)
    assert v.status_code == 200 and v.json()["ok"] is True, v.text
    assert all(c["ok"] for c in v.json()["checks"]) and v.json()["files_checked"] >= 2
    assert client.post(f"/governance/backups/{name}/verify", json={"passphrase": "wrong passphrase!!"}, headers=h).status_code == 400
    dl = client.get(f"/governance/backups/{name}/download", headers=h)
    assert dl.status_code == 200 and dl.content.startswith(b"DCNBAK1\n") and b"backup me" not in dl.content     # encrypted
    # tamper with the encrypted file -> detected
    p = G._backup_path(name)
    raw = bytearray(p.read_bytes()); raw[-20] ^= 0xFF
    tampered = p.parent / re.sub(r"-[0-9a-f]{6}\.dcnbak$", "-abcdef.dcnbak", name)
    tampered.write_bytes(bytes(raw))
    try:
        assert client.post(f"/governance/backups/{tampered.name}/verify", json={"passphrase": PASS}, headers=h).status_code == 400
    finally:
        tampered.unlink()
    # inner-manifest tamper: re-encrypt a modified archive with the right passphrase
    import zipfile
    plain = G._decrypt_blob(p.read_bytes(), PASS)
    zin = zipfile.ZipFile(io.BytesIO(plain)); buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zout:
        for n in zin.namelist():
            data = zin.read(n)
            if n.startswith("evidence/"):
                data = b"altered contents"
            zout.writestr(n, data)
    res = G._verify_backup(G._encrypt_blob(buf.getvalue(), PASS), PASS)
    assert res["ok"] is False and res["mismatches"]
    # offline restore into a fresh directory
    target = tmp_path / "restored"
    env = {**os.environ, "DCN_PW": PASS}
    run = subprocess.run([sys.executable, str(ROOT / "scripts" / "restore_backup.py"), str(p), "--target", str(target), "--passphrase-env", "DCN_PW"],
                         capture_output=True, text=True, env=env, cwd=str(ROOT))
    assert run.returncode == 0, run.stderr + run.stdout
    assert "Manifest verification: OK" in run.stdout and "Audit chain: VERIFIED" in run.stdout
    assert (target / "database.sqlite").exists() and (target / "evidence" / (ev["evidence_id"] + ".txt")).exists()
    import sqlite3
    conn = sqlite3.connect(target / "database.sqlite")
    assert conn.execute("SELECT COUNT(*) FROM cases WHERE case_number=?", (cn,)).fetchone()[0] == 1
    assert str(target / "evidence") in conn.execute("SELECT stored_path FROM evidence WHERE evidence_id=?", (ev["evidence_id"],)).fetchone()[0]
    conn.close()
    again = subprocess.run([sys.executable, str(ROOT / "scripts" / "restore_backup.py"), str(p), "--target", str(target), "--passphrase-env", "DCN_PW"],
                           capture_output=True, text=True, env=env, cwd=str(ROOT))
    assert again.returncode != 0 and "refusing to overwrite" in again.stderr                   # never overwrites
    wrong = subprocess.run([sys.executable, str(ROOT / "scripts" / "restore_backup.py"), str(p), "--target", str(tmp_path / "x"), "--passphrase-env", "NOPE"],
                           capture_output=True, text=True, env={**os.environ, "NOPE": "nope-nope-nope"}, cwd=str(ROOT))
    assert wrong.returncode != 0 and not (tmp_path / "x").exists()


def test_backup_name_traversal_rejected(client, login):
    h = login("admin")
    for bad in ("../../etc/passwd", "..%2F..%2Fx.db", "dcn-backup-x.dcnbak", "%2e%2e/secret", "dcn-backup-20260101T000000Z-abcdef.dcnbak/../../x"):
        assert client.get(f"/governance/backups/{bad}/download", headers=h).status_code in (400, 404), bad
        assert client.post(f"/governance/backups/{bad}/verify", json={"passphrase": "x"}, headers=h).status_code in (400, 404, 405), bad
    assert client.get("/governance/backups/dcn-backup-20260101T000000Z-abcdef.dcnbak/download", headers=h).status_code == 404


def test_backup_size_cap(client, login, monkeypatch):
    h = login("admin")
    monkeypatch.setenv("BACKUP_MAX_MB", "1")
    monkeypatch.setattr(G, "backup_max_bytes", lambda: 10)
    assert client.post("/governance/backup", json={"passphrase": PASS}, headers=h).status_code == 413


# ------------------------------------------------------------------------------------------ access & sessions
def test_login_history_and_force_logout(client, login, make_user):
    h = login("admin")
    name, pw, uh = make_user("supervisor")
    for _ in range(3):
        assert client.post("/auth/login", json={"username": name, "password": "wrong-password-1"}).status_code == 401
    assert client.post("/auth/login", json={"username": name, "password": pw}).status_code == 200
    assert client.post("/auth/login", json={"username": name, "password": "wrong-password-2"}).status_code == 401
    hist = client.get("/governance/login-history", params={"username": name, "days": 7}, headers=h).json()
    u = next(x for x in hist["users"] if x["username"] == name)
    assert u["failed_logins"] == 4 and u["logins"] >= 2 and u["max_failed_streak"] == 3 and u["current_failed_streak"] == 1
    assert u["ips"] and any(e["action"] == "auth.login_failed" and e["ip"] for e in hist["events"])
    assert client.get("/governance/login-history", headers=uh).status_code == 403
    assert client.get("/governance/login-history", params={"days": 0}, headers=h).status_code == 422
    # my activity: any authenticated user, only their own actions
    mine = client.get("/governance/my-activity", headers=uh)
    assert mine.status_code == 200 and mine.json()["events"] and all(e["action"] for e in mine.json()["events"])
    assert client.get("/governance/my-activity").status_code == 401
    _, _, sub = make_user("auditor")
    assert client.get("/governance/my-activity", headers=sub).status_code == 200
    # force logout kills the token
    assert client.get("/cases", headers=uh).status_code == 200
    assert client.post(f"/governance/users/{name}/force-logout", headers=uh).status_code == 403
    r = client.post(f"/governance/users/{name}/force-logout", headers=h)
    assert r.status_code == 200 and r.json()["self"] is False
    assert client.get("/cases", headers=uh).status_code == 401
    assert client.post("/auth/login", json={"username": name, "password": pw}).status_code == 200      # can sign in again
    assert client.post("/governance/users/nobody-here/force-logout", headers=h).status_code == 404
    actions = [a["action"] for a in client.get("/audit", headers=h).json()["events"]]
    assert "governance.force_logout" in actions


def test_force_logout_self_warns(client, login):
    name = "selfkick" + uuid.uuid4().hex[:5]
    admin = login("admin")
    assert client.post("/admin/users", json={"username": name, "password": "Selfkick-Passw0rd-9", "role": "admin"}, headers=admin).status_code == 200
    tok = client.post("/auth/login", json={"username": name, "password": "Selfkick-Passw0rd-9"}).json()["access_token"]
    hh = {"Authorization": "Bearer " + tok}
    r = client.post(f"/governance/users/{name}/force-logout", headers=hh)
    assert r.status_code == 200 and r.json()["self"] is True and r.json()["warning"]
    assert client.get("/cases", headers=hh).status_code == 401


# ------------------------------------------------------------------------------------------ legal packs
def test_legal_packs_open_and_are_scoped(client, login, make_user):
    h = login("admin")
    cn = _case(client, h)
    ev = _upload(client, h, cn, data=b"section 63 evidence " + uuid.uuid4().hex.encode())
    client.get(f"/evidence/{ev['evidence_id']}/verify", params={"case_number": cn}, headers=h)
    client.get(f"/evidence/{ev['evidence_id']}/download", params={"case_number": cn}, headers=h)
    from docx import Document
    r = client.get("/governance/legal/bsa63-certificate", params={"evidence_id": ev["evidence_id"], "case_number": cn, "format": "docx"}, headers=h)
    assert r.status_code == 200
    doc = Document(io.BytesIO(r.content))
    text = "\n".join(p.text for p in doc.paragraphs) + "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    assert "TEMPLATE AID" in text and ev["sha256"] in text and "PART A" in text and "PART B" in text and ev["evidence_id"] in text and "SHA-256" in text
    pdf = client.get("/governance/legal/bsa63-certificate", params={"evidence_id": ev["evidence_id"], "case_number": cn, "format": "pdf"}, headers=h)
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    idx = client.get("/governance/legal/evidence-index", params={"case_number": cn}, headers=h)
    d2 = Document(io.BytesIO(idx.content))
    cells = [c.text for t in d2.tables for row in t.rows for c in row.cells]
    assert ev["evidence_id"] in cells and ev["sha256"] in cells and "evidence.downloaded" in cells and "evidence.uploaded" in cells
    assert client.get("/governance/legal/evidence-index", params={"case_number": cn, "format": "pdf"}, headers=h).content.startswith(b"%PDF")
    assert client.get("/governance/legal/document-index", params={"case_number": cn}, headers=h).status_code == 200
    assert client.get("/governance/legal/document-index", params={"case_number": cn, "format": "pdf"}, headers=h).content.startswith(b"%PDF")
    # evidence of another case / unknown evidence / role checks
    other = _case(client, h)
    assert client.get("/governance/legal/bsa63-certificate", params={"evidence_id": ev["evidence_id"], "case_number": other}, headers=h).status_code == 404
    _, _, sub = make_user("auditor")
    assert client.get("/governance/legal/evidence-index", params={"case_number": cn}, headers=sub).status_code == 403
    _, _, inv = make_user("supervisor")
    assert client.get("/governance/legal/evidence-index", params={"case_number": cn}, headers=inv).status_code == 404
    assert client.get("/governance/legal/bsa63-certificate", params={"evidence_id": ev["evidence_id"], "case_number": cn, "format": "rtf"}, headers=h).status_code == 422
    actions = [a["action"] for a in client.get("/audit", params={"case_number": cn}, headers=h).json()["events"]]
    assert "legal.bsa63_generated" in actions and "legal.evidence_index_generated" in actions


# ------------------------------------------------------------------------------------------ import connectors
CSV = ("FIR No,Police Station,Complainant Name,Date of Occurrence,Sections,Accused Name,Mobile No,Brief Facts,Remarks\n"
       "0101/2026,Indirapuram,Asha Rao,05/03/2026,\"379, 34\",Unknown,9812345678,Bike stolen near market,x\n"
       "0102/2026,Kavi Nagar,Vikram Das,06/03/2026,420,Sohan,9822233344,Online fraud of Rs 5000,y\n"
       ",,,,,,,,\n"
       ",Kavi Nagar,No Number,,,,,,\n").encode()


def test_import_preview_maps_cctns_headers():
    m = G.suggest_mapping(["FIR No", "Police Station", "Complainant Name", "Date of Occurrence", "Sections", "Accused Name", "Mobile No", "Weird Col", "Complainant Name (Mr/Ms)"])
    assert m["FIR No"]["field"] == "fir_number" and m["Police Station"]["field"] == "police_station"
    assert m["Complainant Name"]["field"] == "complainant" and m["Date of Occurrence"]["field"] == "occurrence_date"
    assert m["Sections"]["field"] == "sections" and m["Accused Name"]["field"] == "accused" and m["Mobile No"]["field"] == "phones"
    assert "Weird Col" not in m and "Complainant Name (Mr/Ms)" not in m          # a field is mapped once
    assert G.suggest_mapping(["Fir Nmber"])["Fir Nmber"]["method"] == "fuzzy"


def test_import_preview_commit_pending_review_and_profiles(client, login, make_user):
    h = login("admin")
    cn = _case(client, h)
    _, _, sup = make_user("investigator")
    assert client.post("/governance/import/preview", files={"file": ("r.csv", CSV)}, headers=sup).status_code == 403
    assert client.post("/governance/import/preview", files={"file": ("r.exe", b"MZ")}, headers=h).status_code == 400
    pv = client.post("/governance/import/preview", files={"file": ("register.csv", CSV, "text/csv")}, headers=h)
    assert pv.status_code == 200, pv.text
    pj = pv.json()
    assert pj["rows"] == 4 and len(pj["preview_rows"]) == 4
    sug = {c["name"]: c["suggested"]["field"] for c in pj["columns"] if c["suggested"]}
    assert sug["FIR No"] == "fir_number" and sug["Mobile No"] == "phones" and "Remarks" not in sug
    mapping = {v: k for k, v in sug.items()}
    bad = client.post("/governance/import/commit", json={"upload_token": pj["upload_token"], "case_number": cn, "mapping": {"nonsense": "FIR No"}}, headers=h)
    assert bad.status_code == 400
    nofir = client.post("/governance/import/commit", json={"upload_token": pj["upload_token"], "case_number": cn, "mapping": {"complainant": "Complainant Name"}}, headers=h)
    assert nofir.status_code == 400
    assert client.post("/governance/import/commit", json={"upload_token": "does-not-exist", "case_number": cn, "mapping": mapping}, headers=h).status_code == 404
    # someone else cannot use my upload
    _, _, sup2 = make_user("supervisor")
    assert client.post("/governance/import/commit", json={"upload_token": pj["upload_token"], "case_number": cn, "mapping": mapping}, headers=sup2).status_code in (403, 404)
    r = client.post("/governance/import/commit", json={"upload_token": pj["upload_token"], "case_number": cn, "mapping": mapping, "profile_name": "cctns-basic"}, headers=h)
    assert r.status_code == 200, r.text
    out = r.json()
    assert out["imported"] == 2 and out["skipped_count"] == 1 and out["status"] == "Pending Review"
    docs = client.get("/documents", params={"case_number": cn}, headers=h).json()["documents"]
    assert len(docs) == 2 and all(d["status"] == "Pending Review" for d in docs)          # never auto-verified
    d = next(x for x in docs if x["structured"]["fir_number"] == "0101/2026")
    assert d["structured"]["police_station"] == "Indirapuram" and d["structured"]["complainant"] == "Asha Rao"
    assert d["structured"]["sections"] == ["379", "34"] and d["structured"]["import"]["row"] == 2
    assert client.get("/entities", params={"case_number": cn}, headers=h).json() == []          # nothing enters the graph before review
    again = client.post("/governance/import/commit", json={"upload_token": pj["upload_token"], "case_number": cn, "mapping": mapping}, headers=h).json()
    assert again["imported"] == 0 and again["duplicates"] == 2
    # profile CRUD
    pr = client.post("/governance/import/profiles", json={"name": "cctns-" + uuid.uuid4().hex[:4], "mapping": mapping}, headers=h)
    assert pr.status_code == 200
    pid = pr.json()["id"]
    assert client.post("/governance/import/profiles", json={"name": pr.json()["name"], "mapping": mapping}, headers=h).status_code == 409
    assert client.post("/governance/import/profiles", json={"name": "bad-one", "mapping": {"zzz": "A"}}, headers=h).status_code == 400
    assert any(p["id"] == pid for p in client.get("/governance/import/profiles", headers=h).json()["profiles"])
    assert client.put(f"/governance/import/profiles/{pid}", json={"name": pr.json()["name"] + "-v2", "mapping": {"fir_number": "FIR No"}}, headers=h).status_code == 200
    assert client.delete(f"/governance/import/profiles/{pid}", headers=h).status_code == 200
    assert client.delete(f"/governance/import/profiles/{pid}", headers=h).status_code == 404
    audit = [a["action"] for a in client.get("/audit", params={"case_number": cn}, headers=h).json()["events"]]
    assert "import.committed" in audit


def test_import_commit_is_case_scoped(client, login, make_user):
    h = login("admin")
    cn = _case(client, h)
    pj = client.post("/governance/import/preview", files={"file": ("register.csv", CSV)}, headers=h).json()
    _, _, inv = make_user("supervisor")
    pv = client.post("/governance/import/preview", files={"file": ("register.csv", CSV)}, headers=inv).json()
    r = client.post("/governance/import/commit", json={"upload_token": pv["upload_token"], "case_number": cn, "mapping": {"fir_number": "FIR No"}}, headers=inv)
    assert r.status_code == 404                                                            # cannot import into a case they cannot see
    assert client.get("/documents", params={"case_number": cn}, headers=h).json()["documents"] == []


# ------------------------------------------------------------------------------------------ folder inbox
def test_inbox_scan_ingests_then_moves(client, login, make_user, tmp_path, monkeypatch):
    monkeypatch.setenv("INBOX_DIR", str(tmp_path / "inbox"))
    h = login("admin")
    cn = _case(client, h)
    root = tmp_path / "inbox"
    (root / cn).mkdir(parents=True)
    fir = (ROOT / "sample_firs" / "FIR_INTERLINK_SAMPLE_01.txt").read_bytes()
    (root / cn / "fir1.txt").write_bytes(fir)
    (root / cn / "bad.exe").write_bytes(b"MZ....")
    (root / cn / "table.csv").write_bytes(b"source,target,relation,confidence\nP-IN1,P-IN2,KNOWS,0.9\n")
    (root / "NO-SUCH-CASE").mkdir(); (root / "NO-SUCH-CASE" / "a.txt").write_bytes(b"hello")
    (root / "..hidden..").mkdir()
    outside = tmp_path / "outside.txt"; outside.write_bytes(fir)
    os.symlink(outside, root / cn / "link.txt")
    os.symlink(tmp_path, root / "SYMDIR-CASE")
    _, _, inv = make_user("supervisor")
    assert client.post("/governance/inbox/scan-now", headers=inv).status_code == 403
    r = client.post("/governance/inbox/scan-now", headers=h)
    assert r.status_code == 200, r.text
    res = r.json()
    assert len(res["processed"]) == 2 and len(res["failed"]) == 1 and res["skipped_unknown_case"] == ["NO-SUCH-CASE"]
    assert not (root / cn / "fir1.txt").exists() and list((root / "_processed" / cn).glob("*fir1.txt"))
    failed = list((root / "_failed" / cn).glob("*bad.exe"))
    assert failed and "not allowed" in Path(str(failed[0]) + ".reason.txt").read_text()
    assert (root / cn / "link.txt").is_symlink() and outside.exists()                       # symlinks untouched
    assert (root / "NO-SUCH-CASE" / "a.txt").exists()
    docs = client.get("/documents", params={"case_number": cn}, headers=h).json()["documents"]
    assert len(docs) == 1 and docs[0]["status"] == "Pending Review"                        # documents wait for a human
    assert any(rel["source"] == "P-IN1" for rel in client.get("/relationships", params={"case_number": cn}, headers=h).json())
    st = client.get("/governance/inbox/status", headers=h).json()
    assert st["totals"]["processed"] >= 2 and st["last_scan_at"] and st["recent"]
    assert st["directory"] and client.get("/governance/inbox/status", headers=inv).json()["directory"] is None
    actions = [a["action"] for a in client.get("/audit", params={"case_number": cn}, headers=h).json()["events"]]
    assert "inbox.ingested" in actions and "inbox.failed" in actions
    # a second pass finds nothing new
    assert client.post("/governance/inbox/scan-now", headers=h).json()["processed"] == []


def test_inbox_respects_min_age(client, login, tmp_path, monkeypatch):
    monkeypatch.setenv("INBOX_DIR", str(tmp_path / "inbox2"))
    h = login("admin")
    cn = _case(client, h)
    (tmp_path / "inbox2" / cn).mkdir(parents=True)
    (tmp_path / "inbox2" / cn / "fresh.txt").write_bytes(b"Case note that just arrived " + uuid.uuid4().hex.encode())
    res = client.post("/governance/inbox/scan-now", json={"min_age_seconds": 3600}, headers=h).json()
    assert res["processed"] == [] and res["pending_recent"] == 1


# ------------------------------------------------------------------------------------------ dashboard
def test_dashboard_is_derived_from_real_rows(client, login, make_user):
    h = login("admin")
    cn = _case(client, h)
    before = client.get("/governance/dashboard", params={"days": 30}, headers=h).json()
    client.post("/notes", json={"case_number": cn, "body": "n1"}, headers=h)
    client.post("/tasks", json={"case_number": cn, "title": "t", "assignee": "admin"}, headers=h)
    ev = _upload(client, h, cn)
    fir = (ROOT / "sample_firs" / "FIR_INTERLINK_SAMPLE_01.txt").read_bytes()
    doc = client.post("/documents/extract", files={"file": ("fir.txt", fir)}, data={"case_number": cn}, headers=h).json()
    d = client.get("/governance/dashboard", params={"days": 30}, headers=h)
    assert d.status_code == 200, d.text
    j = d.json()
    assert j["cases"]["total"] >= 1 and j["cases"]["by_status"]["Active"] >= 1 and j["scope"]["all_cases"] is True
    assert j["backlog"]["unverified_documents"] == before["backlog"]["unverified_documents"] + 1
    assert j["evidence"]["total"] >= 2 and j["evidence"]["by_integrity"].get("verified", 0) >= 1
    assert j["ingestion_weekly"] and sum(w["documents"] for w in j["ingestion_weekly"]) >= 1
    off = next(o for o in j["officers"] if o["username"] == "admin")
    assert off["notes"] >= 1 and off["tasks_open"] >= 1 and off["cases_created"] >= 1
    counts = [t["audited_actions"] for t in j["top_cases"]]
    assert counts == sorted(counts, reverse=True) and all(c > 0 for c in counts)
    assert "available" in j["workflow"]
    # turnaround becomes non-null after a review
    assert client.post(f"/documents/{doc['document_id']}/review", json={"status": "Verified"}, headers=h).status_code == 200
    j2 = client.get("/governance/dashboard", headers=h).json()
    assert j2["verification"]["median_hours"] is not None and j2["verification"]["reviewed_in_window"] >= 1
    assert next(o for o in j2["officers"] if o["username"] == "admin")["documents_verified"] >= 1
    # roles: supervisor allowed but scoped; investigator / auditor / anonymous refused
    _, _, sup = make_user("supervisor")
    js = client.get("/governance/dashboard", headers=sup)
    assert js.status_code == 200 and js.json()["scope"]["all_cases"] is False
    assert js.json()["cases"]["total"] < j2["cases"]["total"]       # the admin's restricted cases are not visible to the supervisor
    _, _, inv = make_user("investigator"); _, _, sub = make_user("auditor")
    assert client.get("/governance/dashboard", headers=inv).status_code == 403
    assert client.get("/governance/dashboard").status_code == 401
    assert client.get("/governance/dashboard", params={"days": 5000}, headers=h).status_code == 422
