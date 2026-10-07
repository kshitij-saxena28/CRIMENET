"""Surveillance reports: authority, append-only log, hash chain and tamper detection, movement analysis, reports, permissions."""
import io
import json
import uuid

import pytest

from backend.app.db.database import SessionLocal

AUTH = {"authority_text": "Permission granted by the Deputy Commissioner under the applicable provision", "authority_ref": "DCP/SURV/2026/041"}


def _case(client, login, tag="SV"):
    cn = f"CASE-{tag}-{uuid.uuid4().hex[:6].upper()}"
    sup, admin = login("supervisor"), login("admin")
    assert client.post("/cases", json={"case_number": cn, "title": "surveillance test"}, headers=sup).status_code == 200
    client.patch(f"/admin/cases/{cn}/visibility", json={"visibility": "shared"}, headers=admin)
    return cn, sup


def _op(client, hdr, cn, **kw):
    body = {"case_number": cn, "codename": "TEST", "subject_text": "Subject Alpha (synthetic)", "objective": "Establish movements of the subject", **AUTH,
            "start_at": "2026-06-01T00:00:00", "end_at": "2026-06-30T23:59:00", "timezone": "IST", "supervising_officer": "Insp. Rao", "team": ["SI Mehta"], **kw}
    return client.post("/surveillance/operations", json=body, headers=hdr)


def _entry(client, hdr, op, at="2026-06-02T09:00:00", place="Sector 18 Market, Noida", text="Subject seen buying fruit at a stall", **kw):
    body = {"observed_at": at, "timezone": "IST", "location_text": place, "observer": "SI Mehta", "observation": text, **kw}
    return client.post(f"/surveillance/{op}/entries", json=body, headers=hdr)


@pytest.fixture()
def world(client, login):
    cn, sup = _case(client, login)
    r = _op(client, sup, cn)
    assert r.status_code == 200, r.text
    return {"cn": cn, "sup": sup, "op": r.json()["operation"]["op_number"]}


# ------------------------------------------------------------------ operations and authority
def test_operation_requires_authority_text_and_reference(client, login):
    cn, sup = _case(client, login)
    for missing in ("authority_text", "authority_ref", "objective", "supervising_officer"):
        body = {"case_number": cn, "objective": "Establish movements", **AUTH, "start_at": "2026-06-01T00:00:00", "supervising_officer": "Insp. Rao", "subject_text": "x"}
        body.pop(missing)
        assert client.post("/surveillance/operations", json=body, headers=sup).status_code == 422, missing
    body = {"case_number": cn, "objective": "Establish movements", "authority_text": "", "authority_ref": "", "start_at": "2026-06-01T00:00:00", "supervising_officer": "Insp. Rao", "subject_text": "x"}
    assert client.post("/surveillance/operations", json=body, headers=sup).status_code == 422


def test_operation_view_and_listing(client, world):
    d = client.get(f"/surveillance/{world['op']}", headers=world["sup"]).json()
    assert d["operation"]["authority_ref"] == AUTH["authority_ref"] and d["operation"]["status"] == "active"
    assert d["integrity"]["status"] == "verified" and d["entries"] == [] and set("ABCDEF") == set(d["scales"]["source"]) and set("123456") == set(d["scales"]["information"])
    lst = client.get("/surveillance/operations", params={"case_number": world["cn"]}, headers=world["sup"]).json()
    assert [o["op_number"] for o in lst["operations"]] == [world["op"]] and lst["limits"]


def test_entities_and_evidence_links_are_validated(client, login):
    cn, sup = _case(client, login)
    assert client.post("/entities", json={"external_id": f"P-{cn}", "name": "Ravi Khatri", "entity_type": "PERSON", "case_number": cn}, headers=sup).status_code == 200
    op = _op(client, sup, cn, subject_entity_id=f"P-{cn}").json()["operation"]
    assert op["subject"] == "Ravi Khatri"
    assert _op(client, sup, cn, subject_entity_id="NOPE").status_code == 400
    ok = _entry(client, sup, op["op_number"], persons=[{"name": "Ravi", "role": "associate", "entity_id": f"P-{cn}"}])
    assert ok.status_code == 200, ok.text
    assert _entry(client, sup, op["op_number"], persons=[{"name": "X", "role": "associate", "entity_id": "NOPE"}]).status_code == 400
    assert _entry(client, sup, op["op_number"], evidence_ids=["EV-DOES-NOT-EXIST"]).status_code == 400


# ------------------------------------------------------------------ log, ratings, vehicles
def test_entry_normalises_vehicles_and_validates_ratings(client, world):
    op, sup = world["op"], world["sup"]
    r = _entry(client, sup, op, vehicles=[{"reg": "dl 3c ab 1234", "description": "white i20"}], source_rating="B", info_rating="2", lat=28.57, lon=77.32)
    assert r.status_code == 200, r.text
    e = client.get(f"/surveillance/{op}", headers=sup).json()["entries"][0]
    assert e["vehicles"][0]["reg"] == "DL3CAB1234" and e["rating"] == "B2" and e["seq"] == 1 and e["lat"] == 28.57
    assert _entry(client, sup, op, source_rating="Z").status_code in (400, 422)
    assert _entry(client, sup, op, info_rating="9").status_code in (400, 422)
    assert _entry(client, sup, op, observation="x").status_code == 422
    unknown = _entry(client, sup, op, place="Somewhere unlisted Lane 4").json()
    assert unknown  # a place the gazetteer does not know is still recorded, just not placed on the map


def test_entry_outside_authorised_period_is_flagged_not_hidden(client, world):
    op, sup = world["op"], world["sup"]
    assert _entry(client, sup, op, at="2026-08-01T10:00:00").status_code == 200
    e = client.get(f"/surveillance/{op}", headers=sup).json()["entries"][0]
    assert e["outside_authorised_period"] is True


# ------------------------------------------------------------------ chain and tamper evidence
def _fill(client, world, n=4):
    for i in range(n):
        assert _entry(client, world["sup"], world["op"], at=f"2026-06-0{i + 1}T09:00:00", text=f"Observation number {i} at the market").status_code == 200


def test_chain_verifies_and_links_each_entry_to_the_previous(client, world):
    _fill(client, world)
    d = client.get(f"/surveillance/{world['op']}", headers=world["sup"]).json()
    es = d["entries"]
    assert [e["seq"] for e in es] == [1, 2, 3, 4]
    assert es[0]["prev_hash"] == d["operation"]["header_hash"]
    assert all(es[i]["prev_hash"] == es[i - 1]["entry_hash"] for i in range(1, 4))
    assert d["operation"]["head_hash"] == es[-1]["entry_hash"] and d["operation"]["entry_count"] == 4
    v = client.get(f"/surveillance/{world['op']}/verify", headers=world["sup"]).json()
    assert v["status"] == "verified" and v["entries_checked"] == 4 and not v["problems"]


def test_editing_a_stored_entry_breaks_the_chain(client, world):
    from backend.app.features.surveillance import SurvEntry
    _fill(client, world)
    with SessionLocal() as db:
        row = db.query(SurvEntry).filter_by(op_number=world["op"], seq=2).one()
        row.observation = "Subject was never seen at the market"  # someone edits the database directly
        db.commit()
    v = client.get(f"/surveillance/{world['op']}/verify", headers=world["sup"]).json()
    assert v["status"] == "broken"
    assert any(p["problem"] == "content_changed" and p.get("seq") == 2 for p in v["problems"]), v["problems"]
    assert client.get(f"/surveillance/{world['op']}", headers=world["sup"]).json()["integrity"]["status"] == "broken"


def test_deleting_or_reordering_entries_is_detected(client, world):
    from backend.app.features.surveillance import SurvEntry
    _fill(client, world)
    with SessionLocal() as db:
        db.query(SurvEntry).filter_by(op_number=world["op"], seq=3).delete()
        db.commit()
    v = client.get(f"/surveillance/{world['op']}/verify", headers=world["sup"]).json()
    kinds = {p["problem"] for p in v["problems"]}
    assert v["status"] == "broken" and kinds & {"sequence_gap", "prev_hash_mismatch", "count_mismatch", "head_mismatch"}, v["problems"]


def test_changing_the_operation_header_is_detected(client, world):
    from backend.app.features.surveillance import SurvOperation
    _fill(client, world, 2)
    with SessionLocal() as db:
        o = db.query(SurvOperation).filter_by(op_number=world["op"]).one()
        o.authority_ref = "FORGED/1"
        db.commit()
    v = client.get(f"/surveillance/{world['op']}/verify", headers=world["sup"]).json()
    assert v["status"] == "broken" and any(p["problem"] == "operation_header_changed" for p in v["problems"])


def test_there_is_no_edit_or_delete_endpoint_and_amendment_appends_a_linked_version(client, world):
    op, sup = world["op"], world["sup"]
    _fill(client, world, 2)
    first = client.get(f"/surveillance/{op}", headers=sup).json()["entries"][0]
    routes = client.get("/openapi.json").json()["paths"]
    assert not [p for p, m in routes.items() if p.startswith("/surveillance") and set(m) & {"put", "patch", "delete"}]
    assert client.put(f"/surveillance/{op}/entries/{first['uid']}", json={}, headers=sup).status_code in (404, 405)
    assert client.delete(f"/surveillance/{op}/entries/{first['uid']}", headers=sup).status_code in (404, 405)
    no_reason = client.post(f"/surveillance/{op}/entries/{first['uid']}/amend", json={"observation": "Corrected text of the observation"}, headers=sup)
    assert no_reason.status_code == 422
    r = client.post(f"/surveillance/{op}/entries/{first['uid']}/amend", json={"reason": "Wrong vehicle number noted", "vehicles": [{"reg": "UP16BH5678"}]}, headers=sup)
    assert r.status_code == 200, r.text
    d = client.get(f"/surveillance/{op}", headers=sup).json()
    es = d["entries"]
    assert len(es) == 3 and es[-1]["supersedes_uid"] == first["uid"] and es[-1]["version"] == 2 and es[-1]["amendment_reason"] == "Wrong vehicle number noted"
    orig = next(e for e in es if e["uid"] == first["uid"])
    assert orig["observation"] == first["observation"] and orig["entry_hash"] == first["entry_hash"] and orig["current"] is False
    assert d["integrity"]["status"] == "verified" and d["operation"]["entry_count"] == 3


def test_closed_operation_accepts_no_entries_and_close_is_supervisor_only(client, world, login):
    op, sup = world["op"], world["sup"]
    _fill(client, world, 2)
    inv = login("investigator")
    assert client.post(f"/surveillance/{op}/close", json={"remarks": "Objective met, closing"}, headers=inv).status_code == 403
    r = client.post(f"/surveillance/{op}/close", json={"remarks": "Objective met, closing"}, headers=sup)
    assert r.status_code == 200 and r.json()["closure_hash"]
    assert _entry(client, sup, op).status_code in (400, 409)
    first = client.get(f"/surveillance/{op}", headers=sup).json()["entries"][0]
    assert client.post(f"/surveillance/{op}/entries/{first['uid']}/amend", json={"reason": "late correction here", "observer": "SI X"}, headers=sup).status_code in (400, 409)
    assert client.post(f"/surveillance/{op}/close", json={"remarks": "Closing again for testing"}, headers=sup).status_code == 409
    d = client.get(f"/surveillance/{op}", headers=sup).json()
    assert d["operation"]["status"] == "closed" and d["can_add"] is False and d["integrity"]["status"] == "verified"


def test_concurrent_looking_appends_keep_the_chain_intact(client, world):
    import concurrent.futures as cf
    with cf.ThreadPoolExecutor(6) as pool:
        res = list(pool.map(lambda i: _entry(client, world["sup"], world["op"], at=f"2026-06-{i + 1:02d}T10:00:00", text=f"Parallel observation {i} noted").status_code, range(8)))
    assert res == [200] * 8
    v = client.get(f"/surveillance/{world['op']}/verify", headers=world["sup"]).json()
    assert v["status"] == "verified" and v["entries_checked"] == 8


# ------------------------------------------------------------------ movement analysis
def test_sample_operation_and_movement_analysis(client, login):
    cn, sup = _case(client, login)
    r = client.post("/surveillance/sample", json={"case_number": cn}, headers=sup)
    assert r.status_code == 200, r.text
    op = r.json()["operation"]["op_number"]
    d = client.get(f"/surveillance/{op}", headers=sup).json()
    assert d["integrity"]["status"] == "verified" and d["integrity"]["audit_anchor"]["mismatched"] == 0 and "SYNTHETIC" in d["operation"]["authority_text"]
    m = client.get(f"/surveillance/{op}/movement", headers=login("investigator")).json()
    assert m["points"] and m["dwell_locations"] and m["legs"] and m["meetings"] and m["vehicles"]
    assert any("separate days" in s for s in m["pattern_of_life"]) and m["caveats"]
    assert any(v["seen"] >= 2 for v in m["vehicles"])
    assert client.get(f"/surveillance/{op}/movement", params={"tz": "UTC", "window_minutes": 5}, headers=sup).status_code == 200


def test_movement_dwell_and_copresence_on_planted_entries(client, world):
    op, sup = world["op"], world["sup"]
    for at, place, lat, lon, per in [("2026-06-02T09:00:00", "Cafe One", 28.6315, 77.2167, []), ("2026-06-02T09:40:00", "Cafe One", 28.6316, 77.2168, [{"name": "Rakesh Kumar", "role": "associate"}]),
                                     ("2026-06-03T09:05:00", "Cafe One", 28.6315, 77.2167, []), ("2026-06-03T18:00:00", "Depot", 28.4595, 77.0266, [])]:
        assert _entry(client, sup, op, at=at, place=place, lat=lat, lon=lon, persons=per, vehicles=[{"reg": "DL3CAB1234"}]).status_code == 200
    m = client.get(f"/surveillance/{op}/movement", headers=sup).json()
    cafe = next(p for p in m["dwell_locations"] if "Cafe One" in p["label"])
    assert cafe["distinct_days"] == 2 and cafe["observed_minutes"] >= 40
    assert any(x["persons"] and "Rakesh Kumar" in " ".join(x["persons"]) for x in m["meetings"])
    assert m["vehicles"][0]["reg"] == "DL3CAB1234" and m["vehicles"][0]["seen"] == 4


# ------------------------------------------------------------------ reports
def _sample(client, login):
    cn, sup = _case(client, login)
    op = client.post("/surveillance/sample", json={"case_number": cn}, headers=sup).json()["operation"]["op_number"]
    return cn, sup, op


def test_reports_are_valid_and_carry_the_integrity_footer(client, login):
    cn, sup, op = _sample(client, login)
    head = client.get(f"/surveillance/{op}", headers=sup).json()["operation"]["head_hash"]
    pdf = client.get(f"/surveillance/{op}/report", params={"format": "pdf"}, headers=sup)
    assert pdf.status_code == 200 and pdf.headers["content-type"].startswith("application/pdf") and pdf.content[:5] == b"%PDF-" and len(pdf.content) > 5000
    import pymupdf
    doc = pymupdf.open(stream=pdf.content, filetype="pdf")
    text = "".join(p.get_text() for p in doc)
    assert doc.page_count >= 2 and head[:16] in text and "DEMO-0001" in text and "Sign" in text
    docx = client.get(f"/surveillance/{op}/report", params={"format": "docx"}, headers=sup)
    assert docx.status_code == 200 and docx.content[:2] == b"PK"
    import docx as pydocx
    dd = pydocx.Document(io.BytesIO(docx.content))
    dtext = "\n".join(p.text for p in dd.paragraphs) + "\n".join(c.text for t in dd.tables for row in t.rows for c in row.cells)
    assert head in dtext and "SYNTHETIC" in dtext
    html = client.get(f"/surveillance/{op}/report", params={"format": "html"}, headers=sup)
    assert html.status_code == 200 and "<html" in html.text.lower() and head in html.text and "<script" not in html.text.lower()
    js = client.get(f"/surveillance/{op}/report", params={"format": "json"}, headers=sup)
    assert js.status_code == 200
    body = js.json()
    assert body["integrity"]["head_hash"] == head and body["integrity"]["entries"] == 13 and body["integrity"]["generated_by"] == "supervisor"
    assert body["integrity"]["report_sha256"] and body["limitations"] and body["signoff"]
    assert client.get(f"/surveillance/{op}/report", params={"format": "exe"}, headers=sup).status_code == 422


def test_report_hash_can_be_recomputed_and_file_header_matches(client, login):
    import hashlib
    from ai_engine.surveillance import chain as ch
    cn, sup, op = _sample(client, login)
    a = client.get(f"/surveillance/{op}/report", params={"format": "html"}, headers=sup)
    assert a.headers["x-report-file-sha256"] == hashlib.sha256(a.content).hexdigest()
    r = client.get(f"/surveillance/{op}/report", params={"format": "json"}, headers=sup)
    body = r.json()
    claimed = body["integrity"].pop("report_sha256")
    assert claimed == r.headers["x-report-sha256"] == hashlib.sha256(ch.canonical(body).encode("utf-8")).hexdigest()


def test_report_escapes_html_in_entries(client, world):
    _entry(client, world["sup"], world["op"], text="<script>alert(1)</script> subject seen near the gate")
    html = client.get(f"/surveillance/{world['op']}/report", params={"format": "html"}, headers=world["sup"]).text
    assert "<script>alert(1)" not in html and "&lt;script&gt;" in html


def test_report_marks_broken_chain(client, world):
    from backend.app.features.surveillance import SurvEntry
    _fill(client, world, 2)
    with SessionLocal() as db:
        row = db.query(SurvEntry).filter_by(op_number=world["op"], seq=1).one()
        row.observer = "Someone else"
        db.commit()
    body = client.get(f"/surveillance/{world['op']}/report", params={"format": "json"}, headers=world["sup"]).json()
    assert body["integrity"]["chain_status"] == "broken"


def test_report_generation_is_audited(client, login):
    cn, sup, op = _sample(client, login)
    client.get(f"/surveillance/{op}/report", params={"format": "html"}, headers=sup)
    log = client.get("/audit", headers=login("admin")).json()
    rows = log if isinstance(log, list) else log.get("events") or log.get("items") or []
    assert any(op in json.dumps(r) and "surveillance" in json.dumps(r) for r in rows[:200])


# ------------------------------------------------------------------ permissions, masking, isolation
def test_investigator_may_read_analyse_report_but_not_write(client, login):
    cn, sup, op = _sample(client, login)
    inv = login("investigator")
    assert client.get(f"/surveillance/{op}", headers=inv).status_code == 200
    assert client.get(f"/surveillance/{op}/verify", headers=inv).status_code == 200
    assert client.get(f"/surveillance/{op}/movement", headers=inv).status_code == 200
    assert client.get(f"/surveillance/{op}/report", params={"format": "html"}, headers=inv).status_code == 200
    assert _op(client, inv, cn).status_code == 403
    assert _entry(client, inv, op).status_code == 403
    assert client.post("/surveillance/sample", json={"case_number": cn}, headers=inv).status_code == 403
    first = client.get(f"/surveillance/{op}", headers=inv).json()["entries"][0]
    assert client.post(f"/surveillance/{op}/entries/{first['uid']}/amend", json={"reason": "no permission here", "observer": "x"}, headers=inv).status_code == 403


def test_auditor_get_masked_views_and_no_analysis_or_reports(client, login):
    cn, sup, op = _sample(client, login)
    for who in ("auditor",):
        h = login(who)
        d = client.get(f"/surveillance/{op}", headers=h)
        assert d.status_code == 200 and d.json()["masked"] is True
        es = d.json()["entries"]
        assert all(e["observation"] == "[hidden for your role]" and e["lat"] is None for e in es)
        assert "DL3CAB1234" not in d.text and "Associate B" not in d.text
        assert client.get(f"/surveillance/{op}/movement", headers=h).status_code == 403
        assert client.get(f"/surveillance/{op}/report", params={"format": "json"}, headers=h).status_code == 403
        assert _entry(client, h, op).status_code == 403


def test_cross_case_isolation_and_no_existence_oracle(client, login, make_user):
    cn, sup, op = _sample(client, login)
    other, _ = _case(client, login)
    assert client.get("/surveillance/operations", params={"case_number": other}, headers=sup).json()["operations"] == []
    admin = login("admin")
    priv = f"CASE-SVPRIV-{uuid.uuid4().hex[:6].upper()}"
    client.post("/cases", json={"case_number": priv, "title": "private"}, headers=admin)
    client.patch(f"/admin/cases/{priv}/visibility", json={"visibility": "private"}, headers=admin)
    pop = _op(client, admin, priv).json()["operation"]["op_number"]
    _, _, outsider = make_user("investigator")
    a = client.get(f"/surveillance/{pop}", headers=outsider)
    b = client.get("/surveillance/OP-9999-9999", headers=outsider)
    assert a.status_code == b.status_code == 404 and a.json() == b.json()
    assert client.get(f"/surveillance/{pop}/report", params={"format": "json"}, headers=outsider).status_code == 404
    assert client.get("/surveillance/operations", params={"case_number": priv}, headers=outsider).status_code == 404


def test_demo_account_is_confined_to_demo_cases_and_purge_removes_rows(client, login):
    from backend.app.features.surveillance import SurvEntry, SurvOperation
    demo, admin, sup = login("demo"), login("admin"), login("supervisor")
    assert client.post("/demo/enable", headers=demo).status_code == 200
    dc, dop = f"CASE-DEMOSV-{uuid.uuid4().hex[:5].upper()}", ""
    try:
        real, _ = _case(client, login, "REALSV")
        assert client.get("/surveillance/operations", params={"case_number": real}, headers=demo).status_code == 404
        assert _op(client, demo, real).status_code == 404
        assert client.post("/workbench/new", json={"case_number": dc, "title": "demo surveillance"}, headers=demo).status_code == 200
        r = client.post("/surveillance/sample", json={"case_number": dc}, headers=demo)
        assert r.status_code == 200, r.text
        dop = r.json()["operation"]["op_number"]
        assert client.get(f"/surveillance/{dop}", headers=sup).status_code == 404
    finally:
        client.post("/demo/disable", headers=admin)
    with SessionLocal() as db:
        assert db.query(SurvOperation).filter_by(case_number=dc).count() == 0
        assert dop and db.query(SurvEntry).filter_by(op_number=dop).count() == 0
