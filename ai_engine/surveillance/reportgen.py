"""Formal surveillance report: one content model (a list of blocks) rendered as HTML, DOCX and PDF.

The block list is built from the report ``data`` dict assembled by the API (see backend/app/features/surveillance.py). All text is escaped on the
way out. The PDF is made with PyMuPDF (already a dependency); the DOCX with python-docx.
"""
from __future__ import annotations

import html as _html
import io
from typing import Any

E = _html.escape


def _t(v: Any) -> str:
    return "" if v is None else str(v)


def build_blocks(d: dict) -> list[tuple]:
    h = d["header"]
    b: list[tuple] = [("title", d["title"]), ("subtitle", f"Case {h['case_number']}  |  Operation {h['operation_number']}"),
                      ("banner", d.get("classification", "RESTRICTED: for investigation and court use. Contains personal data."))]
    b.append(("kv", [("Case", f"{h['case_number']} {h.get('case_title', '')}".strip()), ("Operation number", h["operation_number"] + (f" ({h['codename']})" if h.get("codename") else "")),
                     ("Subject", h["subject"]), ("Objective", h["objective"]), ("Authority / permission", h["authority_text"]), ("Authority reference", h["authority_ref"]),
                     ("Period", h["period"]), ("Status", h["status"]), ("Supervising officer", h["supervising_officer"]), ("Team", h["team"]),
                     ("Prepared by", h["prepared_by"]), ("Generated at", h["generated_at"]), ("Times shown in", h["timezone"])]))
    s = d["summary"]
    b.append(("h1", "1. Summary"))
    b.append(("p", s["text"]))
    b.append(("kv", [("Current log entries", s["entries_current"]), ("Amendments (superseded versions kept)", s["amendments"]), ("Observation period", s["observed_period"]),
                     ("Observers", s["observers"]), ("Distinct locations", s["locations"]), ("Vehicles recorded", s["vehicles"]), ("Persons recorded", s["persons"]),
                     ("Entries outside the authorised period", s["outside_period"]), ("Source / information ratings", s["ratings"])]))
    b.append(("h1", "2. Chronological log"))
    b.append(("p", "Entries are shown in the order recorded. Each is part of a hash chain: see the integrity statement at the end. Ratings: source reliability A-F, information credibility 1-6 (scales in section 10)."))
    b.append(("table", ["No.", "Date / time", "Location", "Observer", "Observation", "Vehicles", "Persons", "Rating", "Note"],
              [[r["seq"], r["when"], r["location"], r["observer"], r["observation"], r["vehicles"], r["persons"], r["rating"], r["note"]] for r in d["log"]], [4, 12, 12, 8, 30, 8, 12, 5, 9]))
    b.append(("h1", "3. Amendments"))
    if d["amendments"]:
        b.append(("p", "Entries are never edited or deleted. An amendment is a new entry that supersedes an earlier one; the earlier entry remains in the log."))
        b.append(("table", ["Entry no.", "Superseded by", "Recorded", "By", "Reason"], [[a["seq"], a["superseded_by"], a["when"], a["by"], a["reason"]] for a in d["amendments"]], [8, 10, 18, 14, 50]))
    else:
        b.append(("p", "No entry has been amended."))
    b.append(("h1", "4. Subjects and associates"))
    b.append(("table", ["Name / description", "Role", "Recorded entries", "First seen", "Last seen", "Linked case entity"],
              [[p["name"], p["role"], p["entries"], p["first"], p["last"], p["entity"]] for p in d["subjects"]], [26, 10, 12, 16, 16, 20]) if d["subjects"] else ("p", "No persons were recorded in the log."))
    b.append(("h1", "5. Vehicles"))
    b.append(("table", ["Registration", "Format", "Times seen", "Places", "Seen with", "First / last"],
              [[v["reg"], "valid pattern" if v["valid_format"] else "check spelling", v["seen"], v["places"], v["seen_with"], v["span"]] for v in d["vehicles"]], [14, 10, 8, 26, 22, 20]) if d["vehicles"] else ("p", "No vehicles were recorded in the log."))
    b.append(("h1", "6. Locations and movement"))
    m = d["movement"]
    if m["pattern_of_life"]:
        b.append(("bullets", m["pattern_of_life"]))
    else:
        b.append(("p", "There were not enough located observations to describe a movement pattern."))
    if m["dwell"]:
        b.append(("table", ["Location", "Observations", "Days", "Observed time", "Typical hours", "Position"],
                  [[x["label"], x["observations"], x["distinct_days"], x["observed"], x["hours"], x["precision"]] for x in m["dwell"]], [30, 10, 8, 14, 14, 14]))
    if m["meetings"]:
        b.append(("p", "Recorded co-presence (persons logged together in time and place):"))
        b.append(("bullets", m["meetings"]))
    b.append(("p", m["explain"]))
    b.append(("h1", "7. Linked social-media findings"))
    so = d["social"]
    b.append(("p", so["note"]))
    if so["accounts"]:
        b.append(("table", ["Account", "Linked to", "Match", "Status"], [[a["account"], a["entity"], a["match"], a["status"]] for a in so["accounts"]], [30, 30, 15, 25]))
    if so["flags"]:
        b.append(("table", ["Account", "Category", "Matched wording", "Review"], [[f["account"], f["category"], f["phrase"], f["review"]] for f in so["flags"]], [25, 22, 33, 20]))
    b.append(("h1", "8. Evidence references"))
    b.append(("table", ["Evidence ID", "File", "SHA-256 (as stored)", "Integrity", "Log entries"], [[e["evidence_id"], e["filename"], e["sha256"], e["integrity"], e["entries"]] for e in d["evidence"]], [16, 22, 34, 12, 10])
             if d["evidence"] else ("p", "No evidence items are referenced by the log."))
    b.append(("h1", "9. Limitations and caveats"))
    b.append(("bullets", d["limitations"]))
    b.append(("h1", "10. Rating scales used"))
    b.append(("table", ["Code", "Meaning"], [[k, v] for k, v in d["scales"]], [8, 92]))
    b.append(("h1", "11. Sign-off"))
    b.append(("signoff", d["signoff"]))
    i = d["integrity"]
    b.append(("h1", "12. Integrity statement"))
    b.append(("kv", [("Log chain check", i["chain_text"]), ("Entries in chain", i["entries"]), ("Latest entry hash (chain head)", i["head_hash"]), ("Operation header hash", i["header_hash"]),
                     ("Closing record hash", i["closure_hash"] or "not closed"), ("Generated by", i["generated_by"]), ("Generated at", i["generated_at"]), ("Report content SHA-256", i["report_sha256"])]))
    b.append(("p", i["explain"]))
    return b


CSS = """body{font-family:Arial,Helvetica,sans-serif;font-size:9.5pt;color:#111;line-height:1.35}
h1{font-size:12.5pt;margin:16pt 0 5pt;border-bottom:1px solid #444;padding-bottom:2pt}
.title{font-size:19pt;font-weight:bold;margin:0}.subtitle{font-size:11pt;color:#333;margin:2pt 0 8pt}
.banner{border:1px solid #900;color:#900;padding:3pt 6pt;font-size:8.5pt;font-weight:bold;margin-bottom:8pt}
table{border-collapse:collapse;width:100%;margin:4pt 0 8pt}th,td{border:1px solid #888;padding:2.5pt 4pt;vertical-align:top;font-size:8.3pt;text-align:left;word-wrap:break-word}
th{background:#e6e6e6}td.k{background:#f3f3f3}p{margin:3pt 0 6pt}ul{margin:3pt 0 6pt 14pt;padding:0}li{margin:1.5pt 0}
.sig td{height:34pt}.mono{font-family:Courier,monospace;font-size:8pt}"""


def _table_html(head, rows, widths=None):
    w = widths or [100 // max(1, len(head))] * len(head)
    out = ["<table><thead><tr>" + "".join(f'<th width="{w[i]}%">{E(x)}</th>' for i, x in enumerate(head)) + "</tr></thead><tbody>"]
    for r in rows:
        out.append("<tr>" + "".join(f"<td>{E(_t(c))}</td>" for c in r) + "</tr>")
    out.append("</tbody></table>")
    return "".join(out)


def render_html(d: dict, full_document: bool = True) -> str:
    parts = []
    for blk in build_blocks(d):
        k = blk[0]
        if k == "title":
            parts.append(f'<p class="title">{E(blk[1])}</p>')
        elif k == "subtitle":
            parts.append(f'<p class="subtitle">{E(blk[1])}</p>')
        elif k == "banner":
            parts.append(f'<div class="banner">{E(blk[1])}</div>')
        elif k == "h1":
            parts.append(f"<h1>{E(blk[1])}</h1>")
        elif k == "p":
            parts.append(f"<p>{E(_t(blk[1]))}</p>")
        elif k == "bullets":
            parts.append("<ul>" + "".join(f"<li>{E(_t(x))}</li>" for x in blk[1]) + "</ul>")
        elif k == "kv":
            parts.append("<table><tbody>" + "".join(f'<tr><td width="28%" class="k"><b>{E(a)}</b></td><td width="72%">{E(_t(v))}</td></tr>' for a, v in blk[1]) + "</tbody></table>")
        elif k == "table":
            parts.append(_table_html(blk[1], blk[2], blk[3] if len(blk) > 3 else None))
        elif k == "signoff":
            rows = "".join(f'<tr><td width="40%" class="k"><b>{E(a)}</b></td><td width="60%">&nbsp;<br>&nbsp;<br>Signature: ______________________________ &nbsp;&nbsp; Date: ______________</td></tr>' for a in blk[1])
            parts.append(f'<table class="sig"><tbody>{rows}</tbody></table>')
    body = "\n".join(parts)
    if not full_document:
        return body
    return ('<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">'
            f"<title>{E(d['title'])} {E(d['header']['operation_number'])}</title><style>{CSS}</style></head><body>{body}"
            f'<p class="mono">Report content SHA-256: {E(d["integrity"]["report_sha256"])}</p></body></html>')


def render_docx(d: dict) -> bytes:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.shared import Cm, Pt, RGBColor
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(1.6))
    st = doc.styles["Normal"]
    st.font.name = "Arial"
    st.font.size = Pt(9.5)

    def grid(head, rows):
        t = doc.add_table(rows=1, cols=len(head))
        t.style = "Table Grid"
        for i, x in enumerate(head):
            c = t.rows[0].cells[i]
            c.text = ""
            r = c.paragraphs[0].add_run(str(x))
            r.bold = True
            r.font.size = Pt(8.5)
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = ""
                cells[i].paragraphs[0].add_run(_t(v)).font.size = Pt(8.5)
        doc.add_paragraph()

    for blk in build_blocks(d):
        k = blk[0]
        if k == "title":
            doc.add_heading(blk[1], 0)
        elif k == "subtitle":
            doc.add_paragraph(blk[1])
        elif k == "banner":
            r = doc.add_paragraph().add_run(blk[1])
            r.bold = True
            r.font.color.rgb = RGBColor(0x99, 0, 0)
        elif k == "h1":
            doc.add_heading(blk[1], 1)
        elif k == "p":
            doc.add_paragraph(_t(blk[1]))
        elif k == "bullets":
            for x in blk[1]:
                doc.add_paragraph(_t(x), style="List Bullet")
        elif k == "kv":
            t = doc.add_table(rows=0, cols=2)
            t.style = "Table Grid"
            for a, v in blk[1]:
                cells = t.add_row().cells
                cells[0].text = ""
                cells[0].paragraphs[0].add_run(a).bold = True
                cells[1].text = _t(v)
            doc.add_paragraph()
        elif k == "table":
            grid(blk[1], blk[2])
        elif k == "signoff":
            t = doc.add_table(rows=0, cols=4)
            t.style = "Table Grid"
            for a in blk[1]:
                cells = t.add_row().cells
                cells[0].text, cells[2].text = a, "Date"
            doc.add_paragraph()
    for s in doc.sections:
        s.footer.paragraphs[0].text = (f"Operation {d['header']['operation_number']}  |  chain head {d['integrity']['head_hash'][:16]}...  |  "
                                       f"report SHA-256 {d['integrity']['report_sha256'][:24]}...  |  generated {d['integrity']['generated_at']}")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def render_pdf(d: dict) -> bytes:
    import pymupdf as fitz
    html = render_html(d, full_document=False)
    story = fitz.Story(html=html, user_css=CSS)
    buf = io.BytesIO()
    writer = fitz.DocumentWriter(buf)
    media = fitz.paper_rect("a4-l")
    where = media + (34, 34, -34, -44)
    more = 1
    while more:
        dev = writer.begin_page(media)
        more, _ = story.place(where)
        story.draw(dev)
        writer.end_page()
    writer.close()
    doc = fitz.open(stream=buf.getvalue(), filetype="pdf")
    n = doc.page_count
    foot = (f"Operation {d['header']['operation_number']} | chain head {d['integrity']['head_hash'][:16]}... | report SHA-256 {d['integrity']['report_sha256'][:32]}... | "
            f"generated {d['integrity']['generated_at']}")
    for i, page in enumerate(doc):
        page.insert_text((34, media.height - 24), foot, fontsize=6.5, fontname="helv", color=(0.3, 0.3, 0.3))
        page.insert_text((media.width - 90, media.height - 24), f"Page {i + 1} of {n}", fontsize=7, fontname="helv", color=(0.3, 0.3, 0.3))
    doc.set_metadata({"title": f"{d['title']} {d['header']['operation_number']}", "author": d["integrity"]["generated_by"], "subject": f"Case {d['header']['case_number']}",
                      "keywords": f"sha256:{d['integrity']['report_sha256']}"})
    out = doc.tobytes(garbage=3, deflate=True)
    doc.close()
    return out
